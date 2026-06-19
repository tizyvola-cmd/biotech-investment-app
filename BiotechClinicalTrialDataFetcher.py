import time
import random
import os
import requests
import urllib3.exceptions
import pandas as pd
import json
import re
from datetime import datetime, timedelta

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

DATA_DIR          = "data"
CACHE_DIR         = os.path.join(DATA_DIR, "cache")
CLINICAL_CACHE_DIR = os.path.join(CACHE_DIR, "clinicaltrials")
OPENFDA_CACHE_DIR  = os.path.join(CACHE_DIR, "openfda")

os.makedirs(DATA_DIR,            exist_ok=True)
os.makedirs(CLINICAL_CACHE_DIR,  exist_ok=True)
os.makedirs(OPENFDA_CACHE_DIR,   exist_ok=True)

CACHE_TTL_DAYS      = 14
CLINICALTRIALS_URL  = "https://clinicaltrials.gov/api/v2/studies"
OPENFDA_URL         = "https://api.fda.gov/drug/label.json"


# ── Helpers generici ───────────────────────────────────────────────────────────

def normalize(s):
    if not s or not isinstance(s, str):
        return ""
    return " ".join(s.lower().replace(",", " ").replace(".", " ").split())


def safe_filename(text):
    text = normalize(text)
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_") or "unknown"


def cache_is_valid(path, ttl_days=CACHE_TTL_DAYS):
    if not os.path.exists(path):
        return False
    mtime = datetime.fromtimestamp(os.path.getmtime(path))
    return datetime.now() - mtime <= timedelta(days=ttl_days)


def load_json_cache(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def save_json_cache(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def get_clinical_cache_path(company_name):
    return os.path.join(CLINICAL_CACHE_DIR, f"{safe_filename(company_name)}.json")


def get_openfda_cache_path(company_name):
    return os.path.join(OPENFDA_CACHE_DIR, f"{safe_filename(company_name)}.json")


# Evita ore di attesa se DNS/rete è giù: un solo tentativo + circuit breaker sul processo.
_DNS_UNREACHABLE = False
_DNS_UNREACHABLE_BANNER_PRINTED = False


def _is_dns_or_resolution_failure(exc: BaseException) -> bool:
    if isinstance(exc, urllib3.exceptions.NameResolutionError):
        return True
    if isinstance(exc, urllib3.exceptions.MaxRetryError):
        r = getattr(exc, "reason", None)
        if isinstance(r, urllib3.exceptions.NameResolutionError):
            return True
        if isinstance(r, BaseException) and _is_dns_or_resolution_failure(r):
            return True
    if isinstance(exc, requests.exceptions.ConnectionError):
        c = exc.__cause__
        if c is not None and _is_dns_or_resolution_failure(c):
            return True
        msg = str(exc).lower()
        if "getaddrinfo failed" in msg or "failed to resolve" in msg:
            return True
        if "name or service not known" in msg or "nodename nor servname" in msg:
            return True
    if isinstance(exc, OSError):
        if getattr(exc, "winerror", None) == 11001:
            return True
        if exc.errno in (11001, 8):
            return True
    return False


def get_json(url, params=None, headers=None, retries=5, base_wait=1.0, timeout=20):
    global _DNS_UNREACHABLE, _DNS_UNREACHABLE_BANNER_PRINTED
    if _DNS_UNREACHABLE:
        return None

    headers = headers or {"User-Agent": "Mozilla/5.0"}
    retryable = {408, 429, 500, 502, 503, 504}

    for attempt in range(retries):
        try:
            r = requests.get(url, params=params, headers=headers, timeout=timeout)
            if r.status_code in retryable:
                wait = base_wait * (attempt + 1) + random.random()
                print(f"[RETRY {attempt+1}/{retries}] {r.url} — {r.status_code} — attendo {wait:.1f}s")
                time.sleep(wait)
                continue
            if r.status_code == 404:
                return None
            if r.status_code in {400, 401, 403}:
                print(f"[HTTP ERROR] {r.status_code} — {r.url}")
                return None
            r.raise_for_status()
            return r.json()
        except requests.RequestException as e:
            if _is_dns_or_resolution_failure(e):
                _DNS_UNREACHABLE = True
                if not _DNS_UNREACHABLE_BANNER_PRINTED:
                    print(
                        "[clinical/API] DNS o risoluzione nomi non riuscita "
                        f"(es. clinicaltrials.gov / api.fda.gov): {e}\n"
                        "  → Nessun ulteriore tentativo di rete in questo run (circuit breaker). "
                        "Verifica rete/VPN/DNS e rilancia l'orchestrator."
                    )
                    _DNS_UNREACHABLE_BANNER_PRINTED = True
                return None
            wait = base_wait * (attempt + 1) + random.random()
            print(f"[RETRY {attempt+1}/{retries}] {url} — {e} — attendo {wait:.1f}s")
            time.sleep(wait)
        except ValueError as e:
            print(f"[HTTP ERROR] JSON non valido — {e}")
            return None

    return None


# ── Fetch ClinicalTrials.gov ──────────────────────────────────────────────────

def fetch_clinicaltrials(company_name, limit=20):
    cache_path = get_clinical_cache_path(company_name)
    if cache_is_valid(cache_path):
        cached = load_json_cache(cache_path)
        if cached is not None:
            return cached

    data = get_json(CLINICALTRIALS_URL,
                    params={"query.term": company_name, "pageSize": limit},
                    base_wait=1.0, timeout=20)
    if not data:
        save_json_cache(cache_path, [])
        return []

    rows = []
    for study in data.get("studies", []) or []:
        protocol   = study.get("protocolSection", {}) or {}
        ident      = protocol.get("identificationModule",       {}) or {}
        status     = protocol.get("statusModule",               {}) or {}
        sponsor    = protocol.get("sponsorCollaboratorsModule", {}) or {}
        design     = protocol.get("designModule",               {}) or {}
        conditions = protocol.get("conditionsModule",           {}) or {}
        arms       = protocol.get("armsInterventionsModule",    {}) or {}

        lead_sponsor    = sponsor.get("leadSponsor",      {}) or {}
        resp_party      = sponsor.get("responsibleParty", {}) or {}
        collab_list     = sponsor.get("collaborators",    []) or []

        rows.append({
            "source":                  "clinicaltrials",
            "query_company":           company_name,
            "nct_id":                  ident.get("nctId",        ""),
            "brief_title":             ident.get("briefTitle",   ""),
            "official_title":          ident.get("officialTitle",""),
            "overall_status":          status.get("overallStatus", ""),
            "study_type":              design.get("studyType",   ""),
            "phase":                   " | ".join(design.get("phases", []) or []),
            "conditions":              " | ".join(conditions.get("conditions", []) or []),
            "interventions":           " | ".join(
                x.get("name", "") for x in (arms.get("interventions", []) or [])
                if isinstance(x, dict)
            ),
            # ── Sponsor / responsabile ────────────────────────────────────────
            "lead_sponsor":            lead_sponsor.get("name",  ""),
            "lead_sponsor_class":      lead_sponsor.get("class", ""),
            # Responsible party: tipo + organizzazione/affiliazione
            "responsible_party_type":  resp_party.get("responsiblePartyType", ""),
            "responsible_party_org":   resp_party.get("investigatorAffiliation", ""),
            # Collaboratori: nomi separati da " | "
            "collaborators":           " | ".join(
                c.get("name", "") for c in collab_list if isinstance(c, dict) and c.get("name")
            ),
            # ── Date ─────────────────────────────────────────────────────────
            "last_update_posted_date": status.get("lastUpdatePostDateStruct", {}).get("date", ""),
            "start_date":              status.get("startDateStruct", {}).get("date", ""),
            "primary_completion_date": status.get("primaryCompletionDateStruct", {}).get("date", ""),
            "completion_date":         status.get("completionDateStruct", {}).get("date", ""),
        })

    save_json_cache(cache_path, rows)
    return rows


# ── Fetch OpenFDA ─────────────────────────────────────────────────────────────

def build_openfda_queries(company_name):
    norm  = normalize(company_name)
    parts = norm.split()
    queries, seen = [], set()
    for q in ([norm] + ([" ".join(parts[:2])] if len(parts) >= 2 else []) + parts[:1]):
        q = q.strip()
        if q and q not in seen:
            seen.add(q)
            queries.append(q)
    return queries


def fetch_openfda(company_name, limit=20):
    cache_path = get_openfda_cache_path(company_name)
    if cache_is_valid(cache_path):
        cached = load_json_cache(cache_path)
        if cached is not None:
            return cached

    all_results = []
    for q in build_openfda_queries(company_name):
        query = (f'openfda.manufacturer_name:"{q}" OR openfda.brand_name:"{q}" '
                 f'OR openfda.generic_name:"{q}" OR openfda.substance_name:"{q}"')
        data = get_json(OPENFDA_URL, params={"search": query, "limit": limit},
                        base_wait=1.0, timeout=20)
        if data and "results" in data:
            all_results.extend(data["results"])
            break

    def jl(v):
        if isinstance(v, list): return " | ".join(str(x) for x in v if x is not None)
        return "" if v is None else str(v)

    rows = []
    for item in all_results:
        openfda = item.get("openfda", {}) or {}
        rows.append({
            "source":                  "openfda",
            "query_company":           company_name,
            "brand_name":              jl(openfda.get("brand_name",      [])),
            "generic_name":            jl(openfda.get("generic_name",    [])),
            "manufacturer_name":       jl(openfda.get("manufacturer_name",[])),
            "substance_name":          jl(openfda.get("substance_name",  [])),
            "product_type":            item.get("product_type", ""),
            "purpose":                 jl(item.get("purpose",    [])),
            "indications_and_usage":   jl(item.get("indications_and_usage", [])),
            "warnings":                jl(item.get("warnings",   [])),
            "dosage_and_administration": jl(item.get("dosage_and_administration", [])),
            "effective_time":          item.get("effective_time", ""),
            "version":                 item.get("version",      ""),
        })

    save_json_cache(cache_path, rows)
    return rows


# ── Sponsor match ──────────────────────────────────────────────────────────────

def _match_one(q: str, candidate: str) -> str:
    """
    Confronta la company normalizzata (q) con un singolo candidato normalizzato.
    Restituisce "Exact", "Partial" o "" (nessun match).
    """
    if not candidate:
        return ""
    if q == candidate:
        return "Exact"
    if q in candidate or candidate in q:
        return "Partial"
    q_words = set(q.split())
    c_words = set(candidate.split())
    common  = q_words & c_words
    if common and len(common) / max(len(q_words), 1) >= 0.5:
        return "Partial"
    return ""


def _sponsor_match(query_company: str,
                   lead_sponsor:        str = "",
                   responsible_party_org: str = "",
                   collaborators:       str = "") -> str:
    """
    Confronta il nome della company con:
      - lead_sponsor
      - responsible_party_org  (investigatorAffiliation = "Information provided by")
      - collaborators           (lista separata da " | ")

    Valori restituiti:
      "Exact"    — corrispondenza esatta con almeno uno dei campi
      "Partial"  — corrispondenza parziale con almeno uno
      "No match" — nessuna corrispondenza trovata
      "N/D"      — tutti i campi sponsor vuoti (es. righe openfda)
    """
    q = normalize(str(query_company))
    if not q:
        return "N/D"

    # Candidati: sponsor, responsible party, ogni singolo collaboratore
    candidates = []
    for raw in [lead_sponsor, responsible_party_org]:
        v = normalize(str(raw or ""))
        if v:
            candidates.append(v)
    for part in str(collaborators or "").split("|"):
        v = normalize(part)
        if v:
            candidates.append(v)

    if not candidates:
        return "N/D"

    best = ""
    for cand in candidates:
        result = _match_one(q, cand)
        if result == "Exact":
            return "Exact"   # migliore possibile, stop immediato
        if result == "Partial":
            best = "Partial"

    return best if best else "No match"


# ── Caricamento dati di riferimento yfinance ───────────────────────────────────

def load_yf_company_tickers() -> dict:
    """Restituisce { nome_company_normalizzato: ticker }."""
    yf_path = os.path.join(DATA_DIR, "yf.json")
    if not os.path.exists(yf_path):
        return {}
    try:
        with open(yf_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return {}
    return {
        normalize(item["companyName"]): str(item["symbol"]).strip().upper()
        for item in data
        if isinstance(item, dict) and item.get("companyName") and item.get("symbol")
    }


def load_companies_from_yf() -> list:
    yf_path = os.path.join(DATA_DIR, "yf.json")
    if not os.path.exists(yf_path):
        print(f"[WARN] File non trovato: {yf_path}")
        return []
    try:
        with open(yf_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        print(f"[WARN] Impossibile leggere {yf_path}: {e}")
        return []
    seen, companies = set(), []
    for item in data:
        value = item.get("companyName") or item.get("symbol")
        if value and value not in seen:
            seen.add(value)
            companies.append(value)
    return companies


# ── Aggiunta colonne match ─────────────────────────────────────────────────────

def add_match_columns(df: pd.DataFrame) -> pd.DataFrame:
    """
    Aggiunge tre colonne:
      ticker        — ticker associato alla company (da yf.json)
      company_match — "match" se il ticker è stato trovato, "non match" altrimenti
      sponsor_match — "Exact" / "Partial" / "No match" / "N/D"
                      confronto tra query_company e:
                        lead_sponsor + responsible_party_org + collaborators
    """
    if df.empty:
        return df

    ticker_map = load_yf_company_tickers()

    if "query_company" in df.columns:
        df["ticker"] = df["query_company"].apply(
            lambda c: ticker_map.get(normalize(str(c)), "")
        )
        df["company_match"] = df["ticker"].apply(
            lambda t: "match" if str(t).strip() else "non match"
        )
    else:
        df["ticker"]        = ""
        df["company_match"] = "non match"

    # sponsor_match: controlla lead_sponsor, responsible_party_org, collaborators
    if "lead_sponsor" in df.columns:
        df["sponsor_match"] = df.apply(
            lambda row: _sponsor_match(
                query_company        = row.get("query_company",        ""),
                lead_sponsor         = row.get("lead_sponsor",         ""),
                responsible_party_org= row.get("responsible_party_org",""),
                collaborators        = row.get("collaborators",        ""),
            ),
            axis=1,
        )
    else:
        df["sponsor_match"] = "N/D"

    return df


# ── Deduplicazione ────────────────────────────────────────────────────────────

def get_record_key(row: dict):
    if row.get("source") == "clinicaltrials":
        return ("clinicaltrials", row.get("query_company",""), row.get("nct_id",""))
    if row.get("source") == "openfda":
        return ("openfda", row.get("query_company",""), row.get("brand_name",""),
                row.get("generic_name",""), row.get("manufacturer_name",""),
                row.get("effective_time",""))
    return (row.get("source",""), row.get("query_company",""))


# ── Export Excel con raggruppamento ───────────────────────────────────────────

def export_grouped_excel(df: pd.DataFrame, excel_path: str,
                         company_col: str = "query_company") -> None:
    if df.empty:
        return

    df = df.sort_values([company_col]).reset_index(drop=True)

    wb = load_workbook(excel_path)
    try:
        ws      = wb.active
        headers = list(df.columns)

        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ws.delete_rows(1, 1)
        ws.append(headers)

        HEADER_FILL = PatternFill(fill_type="solid", fgColor="B7E3F4")
        GROUP_FILL  = PatternFill(fill_type="solid", fgColor="D6EAF8")

        for cell in ws[1]:
            cell.fill      = HEADER_FILL
            cell.font      = Font(bold=True)
            cell.alignment = Alignment(horizontal="left")

        current_company = None
        group_start_row = None

        def _write_group_summary(company_name, group_start, summary_row_num):
            group_df   = df[df[company_col] == company_name]
            ticker_val = str(group_df["ticker"].iloc[0]) if "ticker" in df.columns and not group_df.empty else ""
            match_val  = str(group_df["company_match"].iloc[0]) if "company_match" in df.columns and not group_df.empty else ""
            spon_val   = ""
            if "sponsor_match" in df.columns and not group_df.empty:
                counts = group_df["sponsor_match"].value_counts()
                spon_val = counts.index[0] if not counts.empty else ""

            summary = []
            for col in headers:
                if   col == company_col:     summary.append(company_name)
                elif col == "ticker":        summary.append(ticker_val)
                elif col == "company_match": summary.append(match_val)
                elif col == "sponsor_match": summary.append(spon_val)
                else:                        summary.append("")
            ws.append(summary)
            sr = ws.max_row
            for cell in ws[sr]:
                cell.fill      = GROUP_FILL
                cell.font      = Font(bold=True)
                cell.alignment = Alignment(horizontal="left")
            for r in range(group_start, sr):
                ws.row_dimensions[r].outlineLevel = 1
                ws.row_dimensions[r].hidden       = False

        for _, row in df.iterrows():
            company = str(row.get(company_col, "")).strip()
            if company != current_company:
                if current_company is not None and group_start_row is not None:
                    _write_group_summary(current_company, group_start_row, ws.max_row)
                current_company = company
                group_start_row = ws.max_row + 1
            ws.append([row.get(col, "") for col in headers])

        if current_company is not None and group_start_row is not None:
            _write_group_summary(current_company, group_start_row, ws.max_row)

        ws.sheet_properties.outlinePr.summaryBelow = True
        wb.save(excel_path)
    finally:
        wb.close()


# ── Caricamento output esistente ──────────────────────────────────────────────

def load_existing_output(csv_path: str, excel_path: str) -> pd.DataFrame:
    for path, reader in [(csv_path, pd.read_csv), (excel_path, pd.read_excel)]:
        if os.path.exists(path):
            try:
                return reader(path)
            except Exception as e:
                print(f"[WARN] Impossibile leggere {path}: {e}")
    return pd.DataFrame()


# ── Main ──────────────────────────────────────────────────────────────────────

def search_company(company_name: str) -> pd.DataFrame:
    rows = fetch_clinicaltrials(company_name)
    time.sleep(0.1)
    rows += fetch_openfda(company_name)
    return pd.DataFrame(rows)


def main():
    companies = load_companies_from_yf()
    if not companies:
        print("[WARN] Nessuna azienda trovata in yf.json")
        return

    csv_path   = os.path.abspath(os.path.join(DATA_DIR, "biotech_clinical_openfda.csv"))
    excel_path = os.path.abspath(os.path.join(DATA_DIR, "biotech_clinical_openfda.xlsx"))

    existing_df = load_existing_output(csv_path, excel_path)
    new_dfs     = []

    for company in companies:
        clinical_cache = get_clinical_cache_path(company)
        openfda_cache  = get_openfda_cache_path(company)

        if cache_is_valid(clinical_cache) and cache_is_valid(openfda_cache):
            print(f">> {company} (cache valida) — skip fetch")
            df = pd.DataFrame(
                (load_json_cache(clinical_cache) or []) +
                (load_json_cache(openfda_cache)  or [])
            )
        else:
            print(f"\n>> Processing {company} (nuova fetch)")
            df = search_company(company)

        if not df.empty:
            new_dfs.append(df)
        time.sleep(0.05)

    new_df = pd.concat(new_dfs, ignore_index=True) if new_dfs else pd.DataFrame()

    if not existing_df.empty and not new_df.empty:
        final_df = pd.concat([existing_df, new_df], ignore_index=True)
    elif not existing_df.empty:
        final_df = existing_df.copy()
    else:
        final_df = new_df.copy()

    if not final_df.empty:
        final_df["_key"] = final_df.apply(lambda r: str(get_record_key(r.to_dict())), axis=1)
        final_df = final_df.drop_duplicates(subset=["_key"], keep="last").drop(columns=["_key"])
        final_df = add_match_columns(final_df)

    if not new_dfs:
        print(">> Nessun nuovo dato: uso file esistente, non riscrivo Excel")
    else:
        final_df.to_csv(csv_path,   index=False)
        # Excel supporta max 32767 caratteri per cella: tronca in modo esplicito
        # per evitare warning impliciti e mantenere il processo stabile.
        excel_df = final_df.copy()
        _excel_max = 32767
        _trunc_count = 0
        for _col in excel_df.columns:
            _series = excel_df[_col]
            _as_str = _series.astype(str)
            _mask = _series.notna() & (_as_str.str.len() > _excel_max)
            if _mask.any():
                _trunc_count += int(_mask.sum())
                excel_df.loc[_mask, _col] = _as_str[_mask].str.slice(0, _excel_max)
        if _trunc_count:
            print(f"[WARN] Celle Excel troncate a {_excel_max} caratteri: {_trunc_count}")
        excel_df.to_excel(excel_path, index=False, engine="openpyxl")
        export_grouped_excel(final_df, excel_path, company_col="query_company")
        print(f"\nDone. Salvato in: {csv_path}")
        print(f"Done. Salvato in: {excel_path}")


if __name__ == "__main__":
    main()
