import re
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation


def dump_dv(path: str) -> None:
    with ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8", errors="replace")
    m = re.search(r"<dataValidations[^>]*>.*?</dataValidations>", xml, re.DOTALL)
    print("---", path)
    print(m.group(0) if m else "NO DV")


wb = Workbook()
ws = wb.active
dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
dv.sqref = "C7:C20"
ws.add_data_validation(dv)
wb.save("_tA.xlsx")
dump_dv("_tA.xlsx")

wb2 = Workbook()
ws2 = wb2.active
ws2["AB1"] = 0
ws2["AB2"] = 1
dv2 = DataValidation(type="list", formula1="=$AB$1:$AB$2", allow_blank=True)
dv2.showDropDown = False
dv2.sqref = "C7:C20"
ws2.add_data_validation(dv2)
wb2.save("_tB.xlsx")
dump_dv("_tB.xlsx")

# production-style coalesce
from simulation_grafici_sheet import (
    _apply_list_data_validation,
    _grafici_reset_pending_validations,
    _grafici_setup_yesno_dv_column,
)

wb3 = Workbook()
ws3 = wb3.active
ws3.title = "Grafici"
_grafici_reset_pending_validations(ws3)
_grafici_setup_yesno_dv_column(ws3)
refs = [f"C{r}" for r in range(7, 22)]
_apply_list_data_validation(ws3, refs, '"0,1"', error_msg="0 o 1")
wb3.save("_tC.xlsx")
dump_dv("_tC.xlsx")
