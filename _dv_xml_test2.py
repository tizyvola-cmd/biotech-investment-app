import re
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as gcl
from simulation_grafici_sheet import (
    _apply_list_data_validation,
    _coalesce_excel_cell_refs,
)

refs = []
for r in range(7, 22):
    for c in range(3, 11):
        refs.append(f"{gcl(c)}{r}")
sq = _coalesce_excel_cell_refs(refs)
print("sqref", sq, "len", len(sq))

wb = Workbook()
ws = wb.active
ws["AB1"] = 0
ws["AB2"] = 1
_apply_list_data_validation(ws, refs, "=$AB$1:$AB$2", error_msg="0 o 1")
wb.save("_tD.xlsx")

with ZipFile("_tD.xlsx") as z:
    xml = z.read("xl/worksheets/sheet1.xml").decode()
m = re.search(r"<dataValidations.*?</dataValidations>", xml, re.DOTALL)
print(m.group(0) if m else "NO")
