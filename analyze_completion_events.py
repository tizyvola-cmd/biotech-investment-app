#!/usr/bin/env python3
"""
analyze_completion_events.py
────────────────────────────
Event-study: come i titoli biotech si comportano intorno alla
completion date degli studi clinici negli ultimi 10 anni.

Finestre: da -12 settimane a +12 settimane (punti settimanali).
Rendimento calcolato sempre relativo al prezzo nel giorno evento (t=0).

Uso standalone : python analyze_completion_events.py
Uso da modulo  : from analyze_completion_events import analyze
"""

import os, json
import pandas as pd
import numpy as np
import yfinance as yf
from datetime import date, timedelta
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

# ── Paths ─────────────────────────────────────────────────────────────────────
DATA_DIR      = "data"
CLINICAL_XLSX = os.path.join(DATA_DIR, "biotech_clinical_openfda.xlsx")
OUTPUT_XLSX   = os.path.join(DATA_DIR, "completion_event_study.xlsx")
PRICE_CACHE   = os.path.join(DATA_DIR, "price_history_cache.json")

# ── Parametri ─────────────────────────────────────────────────────────────────
LOOKBACK_YEARS  = 10
SNAP_TOLERANCE  = 5   # giorni di tolleranza per trovare un giorno di borsa

# Finestre settimanali: -12w … 0 … +12w
WEEKS      = list(range(-12, 13))          # 25 punti
WEEKS_DAYS = [w * 7 for w in WEEKS]       # in giorni

# ── Candidate column names ────────────────────────────────────────────────────
DATE_CANDIDATES   = ["primary_completion_date", "completion_date",
                     "PrimaryCompletionDate",   "CompletionDate"]
TICKER_CANDIDATES = ["ticker", "symbol"]
PHASE_CANDIDATES  = ["phase", "Phase", "study_phase"]
TITLE_CANDIDATES  = ["brief_title", "BriefTitle", "study_title", "title"]
NCT_CANDIDATES    = ["nct_id", "NCTId", "nct_number"]
STATUS_CANDIDATES = ["overall_status", "OverallStatus", "status"]

# ── Stili (usati solo in standalone) ─────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
ZEBRA_FILL   = PatternFill("solid", fgColor="EEF2F7")
FILL_GREEN   = PatternFill("solid", fgColor="C6EFCE")
FILL_RED     = PatternFill("solid", fgColor="FFC7CE")
FONT_GREEN   = Font(color="375623", bold=True)
FONT_RED     = Font(color="9C0006", bold=True)
FMT_PCT      = '+0.00"%";-0.00"%";"-"'
FMT_PRICE    = "#,##0.00"


def find_col(candidates, df):
    return next((c for c in candidates if c in df.columns), None)


def week_ret_col(w: int) -> str:
    return "ret_0w_%" if w == 0 else f"ret_{w:+d}w_%"


def week_label(w: int) -> str:
    return "Evento" if w == 0 else f"{w:+d}w"


# ── Price cache ───────────────────────────────────────────────────────────────
def load_price_cache(path=PRICE_CACHE) -> dict:
    if os.path.exists(path):
        with open(path, "r") as f:
            return json.load(f)
    return {}


def save_price_cache(cache: dict, path=PRICE_CACHE):
    with open(path, "w") as f:
        json.dump(cache, f)


def get_price_series(ticker: str, cache: dict) -> pd.Series:
    if ticker in cache:
        raw = cache[ticker]
    else:
        print(f"  → Download storico {ticker} ...")
        try:
            hist = yf.download(ticker, period="14y", progress=False, auto_adjust=True)
            if hist.empty:
                cache[ticker] = {}
                return pd.Series(dtype=float)
            raw = {str(k)[:10]: float(v)
                   for k, v in hist["Close"].dropna().items()}
            cache[ticker] = raw
        except Exception as exc:
            print(f"  ✗ {ticker}: {exc}")
            cache[ticker] = {}
            return pd.Series(dtype=float)

    if not raw:
        return pd.Series(dtype=float)
    s = pd.Series(raw)
    s.index = pd.to_datetime(s.index)
    return s.sort_index()


def nearest_price(series: pd.Series, target_date, max_offset: int = SNAP_TOLERANCE):
    if series.empty:
        return None
    target = pd.Timestamp(target_date)
    for offset in range(0, max_offset + 1):
        for delta in ([0] if offset == 0 else [offset, -offset]):
            d = target + pd.Timedelta(days=delta)
            if d in series.index:
                return float(series[d])
    return None


# ── Core: calcola eventi ──────────────────────────────────────────────────────
def _build_events(df_ev: pd.DataFrame, price_data: dict,
                  phase_col, title_col, nct_col, status_col) -> pd.DataFrame:
    records = []
    total   = len(df_ev)

    for i, (_, row) in enumerate(df_ev.iterrows(), 1):
        ticker  = str(row["_ticker"])
        ev_date = row["_date"]
        series  = price_data.get(ticker, pd.Series(dtype=float))

        ev_price = nearest_price(series, ev_date)

        rec = {
            "ticker":         ticker,
            "nct_id":         row.get(nct_col,    "") if nct_col    else "",
            "brief_title":    row.get(title_col,  "") if title_col  else "",
            "phase":          row.get(phase_col,  "") if phase_col  else "",
            "overall_status": row.get(status_col, "") if status_col else "",
            "completion_date": str(ev_date),
            "price_at_event": ev_price,
        }

        for w, days in zip(WEEKS, WEEKS_DAYS):
            target = ev_date + timedelta(days=days)
            p = nearest_price(series, target)
            col = week_ret_col(w)
            if w == 0:
                rec[col] = 0.0
            elif ev_price and p:
                rec[col] = round((p / ev_price - 1) * 100, 4)
            else:
                rec[col] = None

        records.append(rec)
        if i % 50 == 0 or i == total:
            print(f"  Elaborati {i}/{total} eventi ...")

    return pd.DataFrame(records)


# ── Core: summary settimanale ─────────────────────────────────────────────────
def _build_chart_data(events_df: pd.DataFrame) -> pd.DataFrame:
    """
    Ritorna un DataFrame con una riga per settimana:
    window_label | window_week | avg_ret_% | median_ret_% | pct_positive | count
    """
    rows = []
    for w in WEEKS:
        col = week_ret_col(w)
        if col not in events_df.columns:
            continue
        vals = events_df[col].dropna()
        if w == 0:
            rows.append({
                "window_label":  "Evento",
                "window_week":   0,
                "avg_ret_%":     0.0,
                "median_ret_%":  0.0,
                "pct_positive":  None,
                "count":         int(len(vals)),
            })
        else:
            rows.append({
                "window_label":  week_label(w),
                "window_week":   w,
                "avg_ret_%":     round(float(vals.mean()), 4) if len(vals) else None,
                "median_ret_%":  round(float(vals.median()), 4) if len(vals) else None,
                "pct_positive":  round(float((vals > 0).mean() * 100), 1) if len(vals) else None,
                "count":         int(len(vals)),
            })
    return pd.DataFrame(rows)


# ── Public API (usata dall'orchestratore) ─────────────────────────────────────
def analyze(clinical_df: pd.DataFrame,
            cache_path: str = PRICE_CACHE) -> tuple:
    """
    Esegue l'event study sul clinical_df già caricato.
    Ritorna (events_df, chart_df).
    events_df : una riga per evento clinico
    chart_df  : una riga per finestra settimanale (-12w … +12w)
    """
    date_col   = find_col(DATE_CANDIDATES,   clinical_df)
    ticker_col = find_col(TICKER_CANDIDATES, clinical_df)
    phase_col  = find_col(PHASE_CANDIDATES,  clinical_df)
    title_col  = find_col(TITLE_CANDIDATES,  clinical_df)
    nct_col    = find_col(NCT_CANDIDATES,    clinical_df)
    status_col = find_col(STATUS_CANDIDATES, clinical_df)

    if not date_col or not ticker_col:
        print("[EventStudy] Colonne obbligatorie mancanti — analisi saltata.")
        return pd.DataFrame(), pd.DataFrame()

    df = clinical_df.copy()
    df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
    cutoff = pd.Timestamp(date.today() - timedelta(days=365 * LOOKBACK_YEARS))
    today  = pd.Timestamp(date.today())
    mask   = (df[date_col] >= cutoff) & (df[date_col] < today)
    df_ev  = df[mask].dropna(subset=[date_col]).copy()
    df_ev["_ticker"] = df_ev[ticker_col].astype(str).str.strip().str.upper()
    df_ev["_date"]   = df_ev[date_col].dt.date

    if df_ev.empty:
        print("[EventStudy] Nessun evento negli ultimi 10 anni.")
        return pd.DataFrame(), pd.DataFrame()

    print(f"[EventStudy] {len(df_ev)} eventi | {df_ev['_ticker'].nunique()} ticker")

    tickers    = sorted(df_ev["_ticker"].dropna().unique())
    cache      = load_price_cache(cache_path)
    price_data = {t: get_price_series(t, cache) for t in tickers}
    save_price_cache(cache, cache_path)

    events_df = _build_events(df_ev, price_data, phase_col, title_col, nct_col, status_col)
    chart_df  = _build_chart_data(events_df)

    return events_df, chart_df


# ── Standalone output Excel ───────────────────────────────────────────────────
def _write_standalone(events_df: pd.DataFrame, chart_df: pd.DataFrame):
    ret_cols = [week_ret_col(w) for w in WEEKS if w != 0]

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        events_df.to_excel(writer, sheet_name="Events",   index=False)
        chart_df.to_excel(writer,  sheet_name="Chart_Data", index=False)

        wb = writer.book
        _style_events_sheet(wb["Events"], events_df, ret_cols)

    print(f"\n✓ Output: {OUTPUT_XLSX}")
    print(f"  {len(events_df)} eventi | ticker con dati: "
          f"{events_df[events_df['price_at_event'].notna()]['ticker'].nunique()}")


def _style_events_sheet(ws, df, ret_cols):
    cols    = list(df.columns)
    nrows   = len(df)
    col_map = {c: i + 1 for i, c in enumerate(cols)}

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        lbl = col_name
        w = 40 if "brief_title" in lbl else (14 if ("ret_" in lbl or "price" in lbl) else max(len(lbl)+2, 12))
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    for ro in range(nrows):
        ri = ro + 2
        for ci, cn in enumerate(cols, start=1):
            cell = ws.cell(row=ri, column=ci)
            cell.alignment = ALIGN_CENTER
            if cn in ret_cols:
                cell.number_format = FMT_PCT
                try:
                    val = float(cell.value) if cell.value is not None else None
                except (TypeError, ValueError):
                    val = None
                if val is not None:
                    cell.fill = FILL_GREEN if val >= 0 else FILL_RED
                    cell.font = FONT_GREEN if val >= 0 else FONT_RED
            elif ro % 2 == 1:
                cell.fill = ZEBRA_FILL

    last = nrows + 1
    for cn in ret_cols:
        if cn not in col_map:
            continue
        cl = get_column_letter(col_map[cn])
        ws.conditional_formatting.add(
            f"{cl}2:{cl}{last}",
            ColorScaleRule(start_type="min", start_color="FF4C4C",
                           mid_type="percentile", mid_value=50, mid_color="FFE680",
                           end_type="max", end_color="5CDB5C"),
        )


def run():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(CLINICAL_XLSX):
        print(f"[ERROR] File non trovato: {CLINICAL_XLSX}")
        return
    df_clin = pd.read_excel(CLINICAL_XLSX)
    events_df, chart_df = analyze(df_clin)
    if not events_df.empty:
        _write_standalone(events_df, chart_df)


if __name__ == "__main__":
    run()
