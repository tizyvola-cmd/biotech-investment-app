#!/usr/bin/env python3
"""
Audit predizione (curva ricalibrata) vs storico ai nodi T−60, T−30, T−10, T−7.

Il foglio **Simulation** espone solo le colonne «Δ% vs Pred−60» (predizione).
Il confronto con lo **storico reale** è sul foglio **Accuracy** (blocchi Pred / Storico % / Δ% Pred−Stor).

Output in data/:
  simulation_curve_audit.csv
  simulation_curve_audit_summary.csv

Uso:
  py -3 audit_simulation_curves.py
  .venv_disabled_20260519\\Scripts\\python.exe audit_simulation_curves.py
  py -3 audit_simulation_curves.py --filter HURA
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

NODES = (-60, -30, -10, -7)
MATCH_TIGHT_PCT = 0.5
MATCH_LOOSE_PCT = 5.0

# Intestazioni esatte (riga 3) come in data_orchestrator._write_accuracy_simulation_sheet
PRED_HEADERS = {
    -60: "Δ% vs Pred−60\nPred\n−60",
    -30: "Δ% vs Pred−60\nPred\n−30",
    -10: "Δ% vs Pred−60\nPred\n−10",
    -7: "Δ% vs Pred−60\nPred\n−7",
}
HIST_HEADERS = {
    -60: "Storico %\nvs T−60\nT−60",
    -30: "Storico %\nvs T−60\nT−30",
    -10: "Storico %\nvs T−60\nT−10",
    -7: "Storico %\nvs T−60\nT−7",
}
DELTA_HEADERS = {
    -60: "Δ%\nPred−Stor\nT−60",
    -30: "Δ%\nPred−Stor\nT−30",
    -10: "Δ%\nPred−Stor\nT−10",
    -7: "Δ%\nPred−Stor\nT−7",
}

SIM_PRED_HEADERS = PRED_HEADERS  # stesse etichette su Simulation


def _project_root() -> Path:
    return Path(__file__).resolve().parent


def _pick_xlsx(explicit: str | None) -> Path | None:
    root = _project_root()
    candidates: list[Path] = []
    if explicit:
        candidates.append(Path(explicit))
    data = root / "data"
    candidates.extend(
        [
            data / "biotech_orchestrated_output.xlsx",
            data / "biotech_orchestrated_output_write_fallback.xlsx",
        ]
    )
    if data.is_dir():
        candidates.extend(sorted(data.glob("biotech_orchestrated_output_*.xlsx"), reverse=True))
    for p in candidates:
        if p.is_file():
            return p
    return None


def _find_accuracy_sheet(names: list[str]) -> str | None:
    for name in names:
        if name.strip().lower() == "accuracy":
            return name
    for name in names:
        low = name.lower()
        if "accuratezza simulation" in low:
            return name
    for name in names:
        low = name.lower()
        if low == "accuracy" or (low.startswith("accuracy") and "tempor" not in low and "modello" not in low):
            return name
    return None


def _num(x) -> float | None:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    if isinstance(x, str) and not str(x).strip():
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    if v == int(v) and abs(v) > 2:
        pass
    return v


def _pct_to_display(v: float | None) -> float | None:
    """Celle Excel: spesso frazione (0.05 = 5%) o già in punti %."""
    if v is None:
        return None
    if abs(v) <= 1.5:
        return v * 100.0
    return v


def _status(pred: float | None, hist: float | None) -> str:
    if pred is None and hist is None:
        return "missing"
    if pred is None:
        return "solo_storico"
    if hist is None:
        return "solo_pred"
    pe = pred - hist
    ape = abs(pe)
    if ape < MATCH_TIGHT_PCT:
        return "allineato"
    if ape < MATCH_LOOSE_PCT:
        return "vicino"
    return "scostamento"


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).replace("\r\n", "\n").strip() for c in df.columns]
    return df


def _load_sheet_fast(xlsx: Path, sheet: str) -> pd.DataFrame:
    """Intestazione riga 3 Excel → header=2 (pandas); solo lettura dati."""
    df = pd.read_excel(xlsx, sheet_name=sheet, header=2, engine="openpyxl")
    return _normalize_columns(df)


def _filter_data_rows(df: pd.DataFrame, ticker_col: str = "Ticker") -> pd.DataFrame:
    if ticker_col not in df.columns:
        return df.iloc[0:0]
    s = df[ticker_col].astype(str).str.strip()
    ok = s.ne("") & s.ne("nan") & ~s.str.startswith("Metriche", na=False)
    ok &= ~s.str.contains("CD passate", case=False, na=False)
    ok &= ~s.str.startswith("TOTALE", na=False)
    return df.loc[ok].copy()


def _read_accuracy_rows(xlsx: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet = _find_accuracy_sheet(wb.sheetnames)
    wb.close()
    if not sheet:
        raise RuntimeError("Foglio Accuracy non trovato nel workbook.")

    missing = [PRED_HEADERS[n] for n in NODES]
    df = _load_sheet_fast(xlsx, sheet)
    for h in missing:
        if h not in df.columns:
            raise RuntimeError(
                f"Colonna Pred mancante su «{sheet}»: {h!r}. Rigenera il workbook."
            )

    df = _filter_data_rows(df)
    if df.empty:
        return pd.DataFrame()

    records: list[dict] = []
    for idx, row in df.iterrows():
        ts = str(row.get("Ticker", "")).strip()
        company = row.get("Società")
        cd = row.get("CD")
        label = f"{ts} | {company} | CD={cd}"
        excel_row = int(idx) + 4 if isinstance(idx, (int, float)) else None
        for node in NODES:
            pred = _pct_to_display(_num(row.get(PRED_HEADERS[node])))
            hist = _pct_to_display(_num(row.get(HIST_HEADERS[node])))
            delta = _pct_to_display(_num(row.get(DELTA_HEADERS[node])))
            calc_delta = (pred - hist) if pred is not None and hist is not None else None
            records.append(
                {
                    "sheet": sheet,
                    "row_excel": excel_row,
                    "ticker": ts,
                    "company": company,
                    "completion_date": cd,
                    "label": label,
                    "node": node,
                    "pred_pct": pred,
                    "storico_pct": hist,
                    "delta_excel": delta,
                    "delta_calc_pred_minus_stor": calc_delta,
                    "status": _status(pred, hist),
                }
            )
    return pd.DataFrame(records)


def _read_simulation_pred(xlsx: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet = next((n for n in wb.sheetnames if n.strip().lower() == "simulation"), None)
    wb.close()
    if not sheet:
        return pd.DataFrame()

    df = _load_sheet_fast(xlsx, sheet)
    if SIM_PRED_HEADERS[-60] not in df.columns:
        return pd.DataFrame()
    df = _filter_data_rows(df)
    if df.empty:
        return pd.DataFrame()

    records: list[dict] = []
    for idx, row in df.iterrows():
        ts = str(row.get("Ticker", "")).strip()
        for node in NODES:
            records.append(
                {
                    "ticker": ts,
                    "company": row.get("Società"),
                    "completion_date": row.get("Completion Date"),
                    "node": node,
                    "sim_pred_pct": _pct_to_display(_num(row.get(SIM_PRED_HEADERS[node]))),
                    "sim_row": int(idx) + 4 if isinstance(idx, (int, float)) else None,
                }
            )
    return pd.DataFrame(records)


def build_summary(detail: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for label, grp in detail.groupby("label", dropna=False):
        n_align = (grp["status"] == "allineato").sum()
        n_close = (grp["status"] == "vicino").sum()
        n_div = (grp["status"] == "scostamento").sum()
        both = grp[grp["pred_pct"].notna() & grp["storico_pct"].notna()]
        diff = (both["pred_pct"] - both["storico_pct"]).abs()
        max_d = diff.max() if len(diff) else None
        mean_d = diff.mean() if len(diff) else None
        verdict = "ok"
        if n_div > 0:
            verdict = "scostamenti"
        elif len(both) and n_align + n_close < len(both):
            verdict = "parziale"
        rows.append(
            {
                "ticker": grp["ticker"].iloc[0],
                "company": grp["company"].iloc[0],
                "completion_date": grp["completion_date"].iloc[0],
                "label": label,
                "nodes": len(grp),
                "allineato": int(n_align),
                "vicino": int(n_close),
                "scostamento": int(n_div),
                "max_abs_pred_minus_stor_pct": max_d,
                "mean_abs_pred_minus_stor_pct": mean_d,
                "verdict": verdict,
            }
        )
    return pd.DataFrame(rows).sort_values(
        ["scostamento", "max_abs_pred_minus_stor_pct"],
        ascending=[False, False],
        na_position="last",
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="Workbook orchestrator")
    ap.add_argument("--filter", help="Regex su ticker/company (es. HURA|TuHURA)")
    args = ap.parse_args()

    xlsx = _pick_xlsx(args.xlsx)
    if xlsx is None:
        print("ERRORE: nessun biotech_orchestrated_output.xlsx in data/", file=sys.stderr)
        return 1

    print(f"Workbook: {xlsx}")
    print("Confronto: foglio Accuracy (Pred % vs Storico % vs T-60)")
    print(f"Nodi: {NODES}\n")

    detail = _read_accuracy_rows(xlsx)
    if detail.empty:
        print("Nessuna riga dati su Accuracy.")
        return 2

    sim = _read_simulation_pred(xlsx)
    if not sim.empty:
        for _df in (detail, sim):
            _df["completion_date"] = pd.to_datetime(
                _df["completion_date"], errors="coerce"
            ).dt.strftime("%Y-%m-%d")
        detail = detail.merge(
            sim,
            on=["ticker", "company", "completion_date", "node"],
            how="left",
            suffixes=("", "_sim"),
        )
        detail["sim_vs_acc_pred"] = detail.apply(
            lambda r: (
                abs(r["pred_pct"] - r["sim_pred_pct"])
                if pd.notna(r.get("pred_pct"))
                and pd.notna(r.get("sim_pred_pct"))
                else None
            ),
            axis=1,
        )

    if args.filter:
        rx = re.compile(args.filter, re.I)
        m = detail["label"].astype(str).str.contains(rx) | detail["ticker"].astype(str).str.contains(rx)
        detail = detail[m]
        if detail.empty:
            print(f"Nessuna riga per filtro: {args.filter}")
            return 2

    summary = build_summary(detail)
    out_dir = _project_root() / "data"
    out_dir.mkdir(parents=True, exist_ok=True)
    p1 = out_dir / "simulation_curve_audit.csv"
    p2 = out_dir / "simulation_curve_audit_summary.csv"
    detail.to_csv(p1, index=False, encoding="utf-8-sig")
    summary.to_csv(p2, index=False, encoding="utf-8-sig")

    n = len(summary)
    print(f"Righe evento (ticker+CD): {n}")
    print(f"  ok: {(summary['verdict'] == 'ok').sum()}")
    print(f"  scostamenti: {(summary['verdict'] == 'scostamenti').sum()}")
    print(f"  parziale: {(summary['verdict'] == 'parziale').sum()}")
    print(f"\nDettaglio: {p1}")
    print(f"Riepilogo: {p2}")

    print("\n--- Top 12 scostamenti (max |Pred − Storico| %) ---")
    top = summary.nlargest(12, "max_abs_pred_minus_stor_pct")[
        ["ticker", "company", "verdict", "scostamento", "max_abs_pred_minus_stor_pct"]
    ]
    print(top.to_string(index=False))

    mask_h = summary["label"].astype(str).str.contains(r"tuhura|hura", case=False, na=False)
    if mask_h.any():
        print("\n--- TuHURA / HURA ---")
        print(summary[mask_h].to_string(index=False))

    print(
        "\nNota: su nodi già osservati (T−60…T−7), Pred≈Storico indica che la curva "
        "ricalibrata usa i close reali. Delta Pred-Stor in Excel dovrebbe coincidere con Pred-Stor."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
