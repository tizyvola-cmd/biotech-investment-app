"""
Dati per Dear PyGui Charts Lab — Simulation (tutte le righe) e Ristretta CD±7.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from orchestrator_io_paths import FINAL_XLSX, PAST_CATALYST_PREDICTIONS_JSON


def _points_from_curve_lists(
    offsets: tuple[int, ...],
    curve_pct: list,
    *,
    tipo: str = "curva",
) -> list[dict]:
    out: list[dict] = []
    for i, off in enumerate(offsets):
        v = curve_pct[i] if i < len(curve_pct) else None
        if v is None or v != v:
            continue
        out.append(
            {
                "offset": int(off),
                "label": f"{int(off):+d}" if int(off) > 0 else str(int(off)),
                "nodo": "standard",
                "tipo": tipo,
                "pct_curva": float(v),
            }
        )
    out.sort(key=lambda p: p["offset"])
    return out


def _series_from_table_points(
    sid: str,
    label: str,
    points: list[dict],
    *,
    kind: str,
    var_horizons: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    out: dict[str, Any] = {
        "id": sid,
        "label": label,
        "kind": kind,
        "points": sorted(points, key=lambda p: (p.get("sort", (0, p.get("offset", 0))))),
    }
    if var_horizons:
        out["var_horizons"] = var_horizons
    return out


def _normalize_pct_cell(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, str) and str(v).strip() in ("", "—", "-", "N/D"):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f != f:
        return None
    if abs(f) <= 1.5:
        return round(f * 100.0, 4)
    return round(f, 4)


def _attach_foglio_pred_to_points(
    pts: list[dict],
    ws,
    row: int,
) -> None:
    """
    ``pct_foglio``: stessi valori delle colonne «Pred ±N» del foglio Simulation
    (lettura Excel), così i grafici desktop coincidono con la tabella snapshot.
    """
    from data_orchestrator import (
        SIM_COL_V4_PRED_LO,
        SIMULATION_PRED_CAL_OFFSETS,
        _seq_curve_pct_at_cal_offset,
        _simulation_sheet_pred_pct,
    )

    _display: list[float | None] = []
    for _i in range(len(SIMULATION_PRED_CAL_OFFSETS)):
        _ci = SIM_COL_V4_PRED_LO + _i
        _display.append(_simulation_sheet_pred_pct(ws.cell(row=row, column=_ci).value))
    _lut = {
        int(SIMULATION_PRED_CAL_OFFSETS[_i]): _display[_i]
        for _i in range(len(SIMULATION_PRED_CAL_OFFSETS))
    }
    for _p in pts:
        if _p.get("nodo") not in (None, "standard"):
            continue
        _off = int(_p["offset"])
        _v = _lut.get(_off)
        if _v is None:
            _v = _seq_curve_pct_at_cal_offset(
                _display, _off, knot_offsets=SIMULATION_PRED_CAL_OFFSETS
            )
        if _v is not None:
            _p["pct_foglio"] = round(float(_v), 4)


def _table8_to_grafici10(
    vals: list | None,
    grafici_offsets: tuple[int, ...],
) -> list:
    """Allinea curve μ a 8 nodi tabella alla griglia grafici (−60…+7, 10 punti)."""
    from ristretta_lab_data import TABLE_OFFSETS

    if not isinstance(vals, list):
        return [None] * len(grafici_offsets)
    lut = {
        int(off): vals[i]
        for i, off in enumerate(TABLE_OFFSETS)
        if i < len(vals)
    }
    return [lut.get(int(x)) for x in grafici_offsets]


def _variation_horizons_from_sim_row(ws, row: int) -> list[dict[str, Any]]:
    """Var. 1g / 1M / 3M / 6M dal foglio Simulation (stesse colonne del grafico Excel)."""
    from data_orchestrator import SIMULATION_36_HEADERS

    _spec = (
        ("Var. 6M %", "6M"),
        ("Var. 3M %", "3M"),
        ("Var. 1M %", "1M"),
        ("Var. Giorn. %", "1d"),
    )
    out: list[dict[str, Any]] = []
    for _hdr, _lab in _spec:
        try:
            _ci = SIMULATION_36_HEADERS.index(_hdr) + 1
        except ValueError:
            out.append({"label": _lab, "pct": None})
            continue
        out.append({"label": _lab, "pct": _normalize_pct_cell(ws.cell(row=row, column=_ci).value)})
    return out


def build_simulation_lab_bundle(
    *,
    xlsx_path: str | Path | None = None,
    past_pred_path: str | Path | None = None,
) -> dict[str, Any]:
    """Tutte le società del foglio Simulation + curve μ di controllo (allineate a Grafici Excel)."""
    from data_orchestrator import (
        SIM_GRAFICI_CURVE_OFFSETS,
        _accuracy_sim_impute_missing_pre_cd_model_pcts,
        _accuracy_sim_merge_live_pred,
        _accuracy_sim_synthesize_interp_nodes_from_post_d_only,
        _grafici_ensure_seq_curve_on_pairwise,
        _overlay_model_pcts_from_simulation_row,
        _ristretta_curve_table_points,
        _sn7_cd7_overlay_seq_curve_from_live,
        _sn7_parse_sheet_date_ddmmyyyy,
    )
    from orchestrator_io_paths import SIM_LIVE_PRED_SNAPSHOT_JSON
    from ristretta_lab_data import build_reference_curves, load_past_pred_map

    import json
    import openpyxl

    _offsets = SIM_GRAFICI_CURVE_OFFSETS
    _today = date.today()
    _xlsx = Path(xlsx_path or FINAL_XLSX)
    pp = load_past_pred_map(past_pred_path)
    _sim_pred_data: dict | None = None
    try:
        _sp = Path(SIM_LIVE_PRED_SNAPSHOT_JSON)
        if _sp.is_file():
            _doc = json.loads(_sp.read_text(encoding="utf-8"))
            _sim_pred_data = _doc.get("rows") if isinstance(_doc, dict) else None
    except Exception:
        _sim_pred_data = None
    series: dict[str, dict] = {}
    _ser_cache: dict = {}
    _fin_df = None
    try:
        from prediction.enrich import load_enrich_dataframes

        _, _fin_df = load_enrich_dataframes(xlsx_path=_xlsx)
    except Exception:
        _fin_df = None

    try:
        refs = build_reference_curves(pp)
        for name, curve in refs.items():
            curve10 = _table8_to_grafici10(curve, _offsets)
            pts = _points_from_curve_lists(_offsets, curve10, tipo="μ storico")
            series[f"ref:{name}"] = _series_from_table_points(
                f"ref:{name}", name, pts, kind="control"
            )
    except Exception:
        pass

    if not _xlsx.is_file():
        return {
            "mode": "simulation",
            "offsets": list(_offsets),
            "series": series,
            "note": f"Workbook assente: {_xlsx}",
        }

    wb = openpyxl.load_workbook(_xlsx, data_only=True)
    try:
        ws = wb["Simulation"]
        _pred_log: set[str] = set()
        _max_r = int(ws.max_row or 0)
        for _rn in range(4, _max_r + 1):
            _tkv = ws.cell(row=_rn, column=1).value
            _tk = str(_tkv or "").strip().upper()
            if not _tk or _tk in ("—", "-", "N/D"):
                continue
            cd = _sn7_parse_sheet_date_ddmmyyyy(ws.cell(row=_rn, column=3).value)
            if cd is None:
                continue
            _sm = str(ws.cell(row=_rn, column=4).value or "").strip().lower()
            if _sm not in ("exact", "partial"):
                continue
            pk = f"{_tk}|{cd.isoformat()}"
            pr = dict(pp.get(pk) or {})
            _accuracy_sim_merge_live_pred({pk: pr}, {}, _sim_pred_data)
            _sn7_cd7_overlay_seq_curve_from_live(
                pr, row_key=pk, ticker_upper=_tk, sim_pred_data=_sim_pred_data
            )
            pw = dict(pr)
            _overlay_model_pcts_from_simulation_row(ws, _rn, pw)
            _grafici_ensure_seq_curve_on_pairwise(
                pw, cd, row_key=pk, today=_today, ser_cache=_ser_cache
            )
            _accuracy_sim_synthesize_interp_nodes_from_post_d_only(pw)
            _accuracy_sim_impute_missing_pre_cd_model_pcts(pw)
            pts = _ristretta_curve_table_points(
                pw,
                cd,
                _today,
                pk=pk,
                pred_log=_pred_log,
                wb=wb,
                offsets=_offsets,
                with_price_usd=True,
                ser_cache=_ser_cache,
                storico_basis="calendar",
                financial_df=_fin_df,
            )
            _attach_foglio_pred_to_points(pts, ws, _rn)
            nm = str(
                pr.get("company_name_full")
                or ws.cell(row=_rn, column=2).value
                or _tk
            ).strip()
            label = f"{_tk}" if not nm else f"{_tk} · {nm[:28]}"
            series[f"co:{pk}"] = _series_from_table_points(
                f"co:{pk}",
                label,
                pts,
                kind="company",
                var_horizons=_variation_horizons_from_sim_row(ws, _rn),
            )
    finally:
        wb.close()

    return {
        "mode": "simulation",
        "offsets": list(_offsets),
        "series": series,
        "note": "",
        "loaded_at": _today.isoformat(),
    }


def build_ristretta_dpg_bundle(
    *,
    past_pred_path: str | Path | None = None,
    window_days: int = 7,
) -> dict[str, Any]:
    """Coorte Ristretta CD±7 — stesso motore del Data Lab Tk."""
    from data_orchestrator import SIM_GRAFICI_CURVE_OFFSETS
    from ristretta_lab_data import build_ristretta_bundle

    b = build_ristretta_bundle(
        past_pred_path=past_pred_path,
        window_days=window_days,
    )
    _offsets = tuple(b.get("offsets") or SIM_GRAFICI_CURVE_OFFSETS)
    series: dict[str, dict] = {}

    for name, curve in (b.get("references") or {}).items():
        curve10 = _table8_to_grafici10(curve, _offsets)
        pts = _points_from_curve_lists(_offsets, curve10, tipo="μ storico")
        series[f"ref:{name}"] = _series_from_table_points(
            f"ref:{name}", name, pts, kind="control"
        )

    _xlsx = Path(FINAL_XLSX)
    wb = None
    if _xlsx.is_file():
        import openpyxl

        wb = openpyxl.load_workbook(_xlsx, data_only=True)
    try:
        from data_orchestrator import (
            _accuracy_sim_impute_missing_pre_cd_model_pcts,
            _accuracy_sim_synthesize_interp_nodes_from_post_d_only,
            _ristretta_curve_table_points,
            _sn7_cd7_overlay_seq_curve_from_live,
        )
        from ristretta_lab_data import load_past_pred_map

        pp = load_past_pred_map(past_pred_path)
        _today = date.today()
        _plog: set[str] = set()
        for c in b.get("companies") or []:
            pr = dict(pp.get(c.key) or {})
            _sn7_cd7_overlay_seq_curve_from_live(
                pr, row_key=c.key, ticker_upper=c.ticker, sim_pred_data=None
            )
            pw = dict(pr)
            _accuracy_sim_impute_missing_pre_cd_model_pcts(pw)
            _accuracy_sim_synthesize_interp_nodes_from_post_d_only(pw)
            pts = _ristretta_curve_table_points(
                pw,
                c.cd,
                _today,
                pk=c.key,
                pred_log=_plog,
                wb=wb,
                offsets=_offsets,
                with_price_usd=True,
            )
            series[f"co:{c.key}"] = _series_from_table_points(
                f"co:{c.key}",
                f"{c.ticker} · {c.company[:24]}",
                pts,
                kind="company",
            )
    finally:
        if wb is not None:
            wb.close()

    return {
        "mode": "ristretta",
        "offsets": list(_offsets),
        "series": series,
        "note": str(b.get("note") or ""),
        "loaded_at": str(b.get("loaded_at") or ""),
        "n_cohort": int(b.get("n_cohort") or 0),
    }
