import os
import json
import random
import time
import argparse
import pandas as pd
import yfinance as yf
from datetime import datetime, timedelta
from tqdm import tqdm
from openpyxl.utils import get_column_letter

from variation_colors import signed_pct_fill_font

DATA_DIR = "data"
YF_JSON = os.path.join(DATA_DIR, "yf.json")
VARIATIONS_JSON = os.path.join(DATA_DIR, "variations.json")
VARIATIONS_XLSX = os.path.join(DATA_DIR, "variations.xlsx")
CACHE_DIR = os.path.join(DATA_DIR, "variations_cache")
CACHE_TTL_HOURS = 24

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(CACHE_DIR, exist_ok=True)


def cache_path(symbol):
    return os.path.join(CACHE_DIR, f"{symbol}.json")


def load_cache(symbol):
    path = cache_path(symbol)
    if not os.path.exists(path):
        return None
    mtime = datetime.fromtimestamp(os.path.getmtime(path))
    if datetime.now() - mtime > timedelta(hours=CACHE_TTL_HOURS):
        return None
    try:
        with open(path, "r") as f:
            return json.load(f)
    except Exception:
        return None


def save_cache(symbol, data):
    with open(cache_path(symbol), "w") as f:
        json.dump(data, f)


def pct_change(old, new):
    try:
        if old is not None and new is not None and old != 0:
            return round(((new - old) / old) * 100, 2)
        return None
    except Exception:
        return None


def fetch_variation_for_symbol(symbol: str, force: bool = False):
    if not force:
        cached = load_cache(symbol)
        if cached is not None:
            return cached

    try:
        ticker = yf.Ticker(symbol)
        hist = ticker.history(period="1y", interval="1d")

        if hist.empty or "Close" not in hist.columns:
            result = {"symbol": symbol, "error": "No historical data"}
            save_cache(symbol, result)
            return result

        hist = hist.dropna(subset=["Close"]).copy()

        if hist.empty:
            result = {"symbol": symbol, "error": "No valid close data"}
            save_cache(symbol, result)
            return result

        last_close = float(hist["Close"].iloc[-1])

        def avg_close_bars(backward_days: int):
            target_idx = len(hist) - 1 - backward_days
            if target_idx < 1 or target_idx >= len(hist) - 1:
                return None
            closes = hist["Close"].iloc[target_idx - 1: target_idx + 2]
            if closes.empty:
                return None
            return round(float(closes.mean()), 4)

        avg_1m = avg_close_bars(21)
        avg_3m = avg_close_bars(63)
        avg_6m = avg_close_bars(126)

        result = {
            "symbol": symbol,
            "last_close": round(last_close, 4),
            "avg_1m_close": avg_1m,
            "avg_3m_close": avg_3m,
            "avg_6m_close": avg_6m,
            "variation_1m_%": pct_change(avg_1m, last_close),
            "variation_3m_%": pct_change(avg_3m, last_close),
            "variation_6m_%": pct_change(avg_6m, last_close),
        }

        save_cache(symbol, result)
        time.sleep(random.uniform(0.2, 0.6))
        return result

    except Exception as e:
        result = {"symbol": symbol, "error": str(e)}
        save_cache(symbol, result)
        return result


def load_symbols_from_yf():
    if not os.path.exists(YF_JSON):
        print(f"[WARN] {YF_JSON} not found — no symbols to process")
        return []
    try:
        with open(YF_JSON, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            symbols = []
            for item in data:
                if isinstance(item, dict):
                    sym = item.get("symbol") or item.get("ticker")
                    if sym:
                        symbols.append(str(sym).strip().upper())
            return symbols
        if isinstance(data, dict):
            return [str(k).strip().upper() for k in data.keys()]
        return []
    except Exception as e:
        print(f"[WARN] Could not read symbols from {YF_JSON}: {e}")
        return []


def color_for_value(value):
    """Variazione %: intensità proporzionale (blu neg / ambra pos)."""
    return signed_pct_fill_font(value, max_abs=50.0)


def save_outputs(records):
    os.makedirs(DATA_DIR, exist_ok=True)

    with open(VARIATIONS_JSON, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)

    df = pd.DataFrame(records)

    preferred_cols = [
        "symbol", "last_close",
        "avg_1m_close", "avg_3m_close", "avg_6m_close",
        "variation_1m_%", "variation_3m_%", "variation_6m_%",
        "error",
    ]
    cols = [c for c in preferred_cols if c in df.columns] + [
        c for c in df.columns if c not in preferred_cols
    ]
    df = df[cols]

    variation_cols = [c for c in ["variation_1m_%", "variation_3m_%", "variation_6m_%"] if c in df.columns]

    with pd.ExcelWriter(VARIATIONS_XLSX, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Variations")
        ws = writer.sheets["Variations"]

        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

        for row_idx in range(2, len(df) + 2):
            for col_idx, col_name in enumerate(df.columns, start=1):
                if col_name not in variation_cols:
                    continue
                cell = ws.cell(row=row_idx, column=col_idx)
                try:
                    value = float(cell.value) if cell.value is not None else None
                except (TypeError, ValueError):
                    value = None
                fill, font = color_for_value(value)
                if fill:
                    cell.fill = fill
                    cell.font = font

    print(f"Saved JSON:  {VARIATIONS_JSON}")
    print(f"Saved Excel: {VARIATIONS_XLSX}")


def load_error_symbols():
    if not os.path.exists(VARIATIONS_JSON):
        return set()
    try:
        with open(VARIATIONS_JSON, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {
            item["symbol"].strip().upper()
            for item in data
            if isinstance(item, dict) and item.get("error") and item.get("symbol")
        }
    except Exception:
        return set()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Fetch stock price variations")
    parser.add_argument(
        "--force", nargs="*", metavar="SYMBOL",
        help="Bypass cache. Uso: --force (tutti) oppure --force AAPL GOOG",
    )
    parser.add_argument(
        "--errors-only", action="store_true",
        help="Riprova solo i simboli con errore nell'ultimo run",
    )
    args = parser.parse_args()

    force_all = args.force is not None and len(args.force) == 0
    force_set = {s.strip().upper() for s in args.force} if args.force else set()

    symbols = load_symbols_from_yf()

    if args.errors_only:
        error_syms = load_error_symbols()
        if not error_syms:
            print("[INFO] Nessun simbolo con errori nell'ultimo run — nulla da fare.")
            symbols = []
        else:
            symbols = [s for s in symbols if s in error_syms]
            print(f"[--errors-only] Retry su {len(symbols)} simboli: {', '.join(sorted(symbols))}")
            force_all = True

    if not symbols:
        print("[WARN] No symbols found — skipping variation fetch")
    else:
        if force_all and not args.errors_only:
            print("Cache bypassed for ALL symbols")
        elif force_set:
            print(f"Cache bypassed for: {', '.join(sorted(force_set))}")

        results = []
        for sym in tqdm(symbols, desc="VARIATIONS (CACHE)"):
            force_this = force_all or (sym in force_set)
            results.append(fetch_variation_for_symbol(sym, force=force_this))

        if args.errors_only:
            try:
                with open(VARIATIONS_JSON, "r", encoding="utf-8") as f:
                    existing = json.load(f)
                retried = {r["symbol"]: r for r in results if "symbol" in r}
                merged = [retried.get(item["symbol"], item) if isinstance(item, dict) else item for item in existing]
                save_outputs(merged)
            except Exception:
                save_outputs(results)
        else:
            save_outputs(results)

        print(f"Ticker processati: {len(results)}")
        print(">>> VARIATIONS COMPLETATO (CACHE ATTIVA) <<<")
