"""
Accuracy metrics for v4 (interpolated Pred curve) vs v5 (q50 fan) on the Accuracy sheet.

Errors compare model % vs T−60 to realized «Storico %» at the same calendar horizon.
"""
from __future__ import annotations

import json
import math
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from prediction.pipeline import SIMULATION_V5_Q50_OFFSETS

# Calendar offsets on the Accuracy display grid (8 columns, aligned to date row).
ACCURACY_V4_DISPLAY_OFFSETS: tuple[int, ...] = (-60, -30, -10, -7, -5, -3, 4, 7)

_STABLE_BAND_PP = 5.0
# Banda «piatto» per hit direzionale in Accuracy (±0.5 pp, come confronto segnale).
SIGN_HIT_FLAT_BAND_PP = 0.5


def horizon_label(offset: int) -> str:
    """Italian-style label for sheet headers (T−60, T+7, …)."""
    if offset > 0:
        return f"T+{offset}"
    return f"T{offset}"


def abs_error_header(model: str, offset: int) -> str:
    """Column header for absolute error in percentage points."""
    m = "v4" if model == "v4" else "v5"
    off = horizon_label(offset)
    return f"|Δ| {m}\n{off}"


def signed_error_header(model: str, offset: int) -> str:
    m = "v4" if model == "v4" else "v5"
    off = horizon_label(offset)
    return f"Err {m}\n{off}"


def float_or_none(val: Any) -> float | None:
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, str):
        s = val.strip().replace("%", "").replace(",", ".")
        if not s or s.lower() in ("—", "-", "n/d", "tbd", "nan"):
            return None
        try:
            f = float(s)
        except ValueError:
            return None
    else:
        try:
            f = float(val)
        except (TypeError, ValueError):
            return None
    if not math.isfinite(f):
        return None
    return f


def signed_error_pp(pred_pct: float | None, actual_pct: float | None) -> float | None:
    """Signed error in percentage points: pred − actual."""
    if pred_pct is None or actual_pct is None:
        return None
    return round(float(pred_pct) - float(actual_pct), 2)


def abs_error_pp(pred_pct: float | None, actual_pct: float | None) -> float | None:
    e = signed_error_pp(pred_pct, actual_pct)
    if e is None:
        return None
    return round(abs(e), 2)


def direction_hit(
    pred_pct: float | None,
    actual_pct: float | None,
    *,
    flat_band_pp: float = _STABLE_BAND_PP,
) -> bool | None:
    if pred_pct is None or actual_pct is None:
        return None
    p, a = float(pred_pct), float(actual_pct)
    band = float(flat_band_pp)
    if abs(a) < band:
        return abs(p) < band
    if p > 0 and a > 0:
        return True
    if p < 0 and a < 0:
        return True
    if abs(p) < band:
        return abs(a) < band
    return False


def sign_hit_pp(pred_pct: float | None, actual_pct: float | None) -> bool | None:
    """Hit direzionale con banda stretta (±0.5 pp) per metriche Accuracy."""
    return direction_hit(pred_pct, actual_pct, flat_band_pp=SIGN_HIT_FLAT_BAND_PP)


def v5_pct_at_offset(v5_offsets: Mapping[Any, Any] | None, offset: int) -> float | None:
    if not isinstance(v5_offsets, dict):
        return None
    for key in (str(offset), offset):
        v = float_or_none(v5_offsets.get(key))
        if v is not None:
            return v
    return None


@dataclass
class HorizonErrors:
    offset: int
    actual_pct: float | None = None
    v4_pred_pct: float | None = None
    v5_pred_pct: float | None = None
    v5_raw_pred_pct: float | None = None
    signed_v4_pp: float | None = None
    signed_v5_pp: float | None = None
    signed_v5_raw_pp: float | None = None
    abs_v4_pp: float | None = None
    abs_v5_pp: float | None = None
    abs_v5_raw_pp: float | None = None
    hit_v4: bool | None = None
    hit_v5: bool | None = None
    hit_v5_raw: bool | None = None


@dataclass
class RowAccuracyRecord:
    row_key: str
    ticker: str
    completion_date: str
    is_past: bool
    horizons: list[HorizonErrors] = field(default_factory=list)

    def mae_v4_by_offset(self) -> dict[int, float]:
        return {
            h.offset: h.abs_v4_pp
            for h in self.horizons
            if h.abs_v4_pp is not None
        }

    def mae_v5_by_offset(self) -> dict[int, float]:
        return {
            h.offset: h.abs_v5_pp
            for h in self.horizons
            if h.abs_v5_pp is not None
        }


def accuracy_v5_raw_metrics_enabled() -> bool:
    """Metriche q50 v5 senza anchor per Accuratezza temporale (default on)."""
    import os

    return os.environ.get("ACCURACY_V5_RAW_METRICS", "1").strip().lower() not in (
        "0",
        "false",
        "no",
        "off",
    )


def accuracy_v5_raw_n_paths() -> int:
    import os

    raw = os.environ.get("ACCURACY_V5_RAW_N_PATHS", "800").strip()
    try:
        n = int(raw)
        return max(200, min(n, 4000))
    except (TypeError, ValueError):
        return 800


def v5_q50_offsets_raw_for_accuracy_row(
    row: dict,
    *,
    harvest_row: dict | None = None,
    financial_df: Any = None,
    ser_cache: dict | None = None,
    past_pred_rows: dict | None = None,
    n_paths: int | None = None,
) -> dict[str, float | None]:
    """
    q50 v5 **senza** ``PRED_V5_ANCHOR_Q50_V4`` e senza sync sulla Pred % display.
    Per riga «v5 raw» su Accuratezza temporale.
    """
    from prediction.pipeline import predict_v5_fan_offsets, predict_v5_q50_offsets

    inp = dict(row)
    _tk = str(inp.get("ticker") or "").strip().upper()
    if not _tk and harvest_row:
        _tk = str(harvest_row.get("ticker") or "").strip().upper()
    _prices: list[float] = []
    if _tk and ser_cache is not None:
        try:
            import data_orchestrator as orch

            _cls = orch._histlib_accuracy_sim_resolve_close_series(_tk, ser_cache)
            if _cls is not None:
                _prices = [
                    float(x)
                    for x in _cls.dropna().tail(400)
                    if x is not None and float(x) > 0
                ]
        except Exception:
            _prices = []
    if len(_prices) < 5:
        _ps = inp.get("price_series")
        if isinstance(_ps, (list, tuple)):
            try:
                _prices = [float(x) for x in _ps if x is not None and float(x) > 0]
            except (TypeError, ValueError):
                pass

    def _pick(*keys, src: dict | None = None):
        for _k in keys:
            for _d in (inp, src or {}, harvest_row or {}):
                if not isinstance(_d, dict):
                    continue
                _v = _d.get(_k)
                if _v is not None and not (
                    isinstance(_v, str) and not str(_v).strip()
                ):
                    return _v
        return None

    _v5_in: dict = {
        "ticker": _tk,
        "slope_20d": _pick("slope_20d"),
        "run_up_30d": _pick("run_up_30d"),
        "exc_slope": _pick("exc_slope", "exc_slope_vs_XBI"),
        "exc_slope_vs_XBI": _pick("exc_slope_vs_XBI", "exc_slope"),
        "expected_move_pct": _pick("expected_move_pct"),
        "model_d5_pct": _pick("model_dm5_pct", "model_d5_pct", "d5_pct"),
        "beta": _pick("beta"),
        "liquidity_score": _pick("liquidity_score"),
        "vol_20d": _pick("vol_20d"),
        "direction": _pick("direction"),
        "dir_v4": _pick("dir_v4"),
        "model_dm60_pct": _pick("model_dm60_pct"),
        "model_dm30_pct": _pick("model_dm30_pct"),
        "model_dm10_pct": _pick("model_dm10_pct"),
        "model_dm7_pct": _pick("model_dm7_pct"),
        "model_dm5_pct": _pick("model_dm5_pct"),
        "model_dm3_pct": _pick("model_dm3_pct"),
        "model_d4_pct": _pick("model_d4_pct"),
        "model_d7_pct": _pick("model_d7_pct"),
    }
    if financial_df is not None and not getattr(financial_df, "empty", True) and _tk:
        try:
            import data_orchestrator as orch
            from prediction.financial_liquidity import financial_row_liquidity_keys

            _fr = orch._financial_series_for_ticker(financial_df, _tk)
            if _fr is not None:
                for _lk, _lv in financial_row_liquidity_keys(_fr).items():
                    if _lv is not None and _v5_in.get(_lk) is None:
                        _v5_in[_lk] = _lv
        except Exception:
            pass

    _np = int(n_paths if n_paths is not None else accuracy_v5_raw_n_paths())
    try:
        predict_v5_fan_offsets(
            _v5_in,
            prices=_prices if len(_prices) >= 5 else None,
            past_pred_rows=past_pred_rows,
            n_paths=_np,
            anchor_q50_v4=False,
        )
        _stored = _v5_in.get("v5_q50_offsets")
        if isinstance(_stored, dict) and _stored:
            return dict(_stored)
        return predict_v5_q50_offsets(
            _v5_in,
            prices=_prices if len(_prices) >= 5 else None,
            past_pred_rows=past_pred_rows,
            n_paths=_np,
        )
    except Exception:
        return {str(o): None for o in SIMULATION_V5_Q50_OFFSETS}


def compute_row_errors(
    *,
    row_key: str,
    ticker: str,
    completion_date: date | None,
    is_past: bool,
    v4_pred_pts: Sequence[float | None],
    storico_pts: Sequence[float | None],
    v5_offsets: Mapping[Any, Any] | None,
    v5_raw_offsets: Mapping[Any, Any] | None = None,
) -> RowAccuracyRecord:
    """Build per-horizon errors from sheet row arrays (pred/storico in %)."""
    cd_s = completion_date.isoformat() if completion_date else ""
    horizons: list[HorizonErrors] = []
    for i, off in enumerate(ACCURACY_V4_DISPLAY_OFFSETS):
        actual = storico_pts[i] if i < len(storico_pts) else None
        v4p = v4_pred_pts[i] if i < len(v4_pred_pts) else None
        v5p = v5_pct_at_offset(v5_offsets, off) if off in SIMULATION_V5_Q50_OFFSETS else None
        v5r = (
            v5_pct_at_offset(v5_raw_offsets, off)
            if off in SIMULATION_V5_Q50_OFFSETS
            else None
        )
        h = HorizonErrors(
            offset=off,
            actual_pct=actual,
            v4_pred_pct=v4p,
            v5_pred_pct=v5p,
            v5_raw_pred_pct=v5r,
            signed_v4_pp=signed_error_pp(v4p, actual),
            signed_v5_pp=signed_error_pp(v5p, actual),
            signed_v5_raw_pp=signed_error_pp(v5r, actual),
            abs_v4_pp=abs_error_pp(v4p, actual),
            abs_v5_pp=abs_error_pp(v5p, actual),
            abs_v5_raw_pp=abs_error_pp(v5r, actual),
            hit_v4=sign_hit_pp(v4p, actual),
            hit_v5=sign_hit_pp(v5p, actual),
            hit_v5_raw=sign_hit_pp(v5r, actual),
        )
        horizons.append(h)
    return RowAccuracyRecord(
        row_key=row_key,
        ticker=ticker,
        completion_date=cd_s,
        is_past=is_past,
        horizons=horizons,
    )


def _quarter_key(cd: date) -> str:
    q = (cd.month - 1) // 3 + 1
    return f"{cd.year}-Q{q}"


def aggregate_period_metrics(
    records: Iterable[RowAccuracyRecord],
    *,
    past_only: bool = True,
) -> dict[str, Any]:
    """
    Aggregate MAE / hit rate by calendar quarter (completion_date) and totals.

    MAE per horizon = mean of |Pred − Storico| in pp, computed internally from
    v4 pred vs storico % and v5 q50 offsets vs storico (not from Excel columns).
    Only past catalyst rows contribute when ``past_only`` (default).
    """
    by_q: dict[str, list[RowAccuracyRecord]] = defaultdict(list)
    pool: list[RowAccuracyRecord] = []
    for rec in records:
        if past_only and not rec.is_past:
            continue
        if not rec.completion_date:
            continue
        try:
            cd = date.fromisoformat(rec.completion_date[:10])
        except ValueError:
            continue
        by_q[_quarter_key(cd)].append(rec)
        pool.append(rec)

    def _agg_model(rows: list[RowAccuracyRecord], model: str) -> dict[str, Any]:
        mae: dict[str, float | None] = {}
        mae_n: dict[str, int] = {}
        hit_pct: dict[str, float | None] = {}
        pooled_abs: list[float] = []
        pooled_hits: list[bool] = []
        for off in ACCURACY_V4_DISPLAY_OFFSETS:
            lbl = horizon_label(off)
            if model in ("v5", "v5_raw") and off not in SIMULATION_V5_Q50_OFFSETS:
                mae[lbl] = None
                mae_n[lbl] = 0
                hit_pct[lbl] = None
                continue
            if model == "v4":
                abs_vals = [
                    h.abs_v4_pp
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.abs_v4_pp is not None
                ]
                hits = [
                    h.hit_v4
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.hit_v4 is not None
                ]
            elif model == "v5_raw":
                abs_vals = [
                    h.abs_v5_raw_pp
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.abs_v5_raw_pp is not None
                ]
                hits = [
                    h.hit_v5_raw
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.hit_v5_raw is not None
                ]
            else:
                abs_vals = [
                    h.abs_v5_pp
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.abs_v5_pp is not None
                ]
                hits = [
                    h.hit_v5
                    for r in rows
                    for h in r.horizons
                    if h.offset == off and h.hit_v5 is not None
                ]
            mae[lbl] = round(sum(abs_vals) / len(abs_vals), 2) if abs_vals else None
            mae_n[lbl] = len(abs_vals)
            hit_pct[lbl] = (
                round(100.0 * sum(1 for x in hits if x) / len(hits), 1) if hits else None
            )
            pooled_abs.extend(abs_vals)
            pooled_hits.extend(h for h in hits if h is not None)
        return {
            "mae": mae,
            "mae_n": mae_n,
            "hit_pct": hit_pct,
            "mae_global": round(sum(pooled_abs) / len(pooled_abs), 2) if pooled_abs else None,
            "mae_global_n": len(pooled_abs),
            "hit_pct_global": (
                round(100.0 * sum(1 for x in pooled_hits if x) / len(pooled_hits), 1)
                if pooled_hits
                else None
            ),
            "hit_pct_global_n": len(pooled_hits),
        }

    def _agg_period(rows: list[RowAccuracyRecord]) -> dict[str, Any]:
        return {
            "n_rows": len(rows),
            "v4": _agg_model(rows, "v4"),
            "v5": _agg_model(rows, "v5"),
            "v5_raw": _agg_model(rows, "v5_raw"),
        }

    by_quarter = {qk: _agg_period(rs) for qk, rs in sorted(by_q.items())}
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "past_only": past_only,
        "by_quarter": by_quarter,
        "total": _agg_period(pool),
    }


def run_snapshot_has_metrics(run: dict[str, Any] | None) -> bool:
    """True se il run ha almeno una riga passata con MAE calcolabile."""
    if not isinstance(run, dict):
        return False
    tot = run.get("total")
    if isinstance(tot, dict):
        try:
            if int(tot.get("n_rows") or 0) > 0:
                return True
        except (TypeError, ValueError):
            pass
    periods = run.get("periods")
    if isinstance(periods, dict):
        for blk in periods.values():
            if not isinstance(blk, dict):
                continue
            try:
                if int(blk.get("n_rows") or 0) > 0:
                    return True
            except (TypeError, ValueError):
                continue
    by_q = run.get("by_quarter")
    if isinstance(by_q, dict):
        for blk in by_q.values():
            if isinstance(blk, dict) and int(blk.get("n_rows") or 0) > 0:
                return True
    return False


def legacy_flat_period_to_model_block(flat: Mapping[str, Any]) -> dict[str, Any]:
    """Converte un blocco ``periods`` legacy (chiavi ``mae_v4_T-60``) in schema v4/v5."""
    if not isinstance(flat, dict):
        flat = {}

    def _model_side(model: str) -> dict[str, Any]:
        mae: dict[str, float | None] = {}
        mae_n: dict[str, int] = {}
        pooled_num = 0.0
        pooled_den = 0
        for off in ACCURACY_V4_DISPLAY_OFFSETS:
            if model == "v5" and off not in SIMULATION_V5_Q50_OFFSETS:
                mae[horizon_label(off)] = None
                mae_n[horizon_label(off)] = 0
                continue
            lbl = horizon_label(off)
            mae_key = f"mae_{model}_{lbl}"
            n_key = f"{mae_key}_n"
            mv = float_or_none(flat.get(mae_key))
            try:
                nv = int(flat.get(n_key) or 0)
            except (TypeError, ValueError):
                nv = 0
            mae[lbl] = mv
            mae_n[lbl] = nv
            if mv is not None and nv > 0:
                pooled_num += float(mv) * nv
                pooled_den += nv
        return {
            "mae": mae,
            "mae_n": mae_n,
            "mae_global": round(pooled_num / pooled_den, 2) if pooled_den else None,
            "mae_global_n": pooled_den,
        }

    try:
        n_rows = int(flat.get("n_rows") or 0)
    except (TypeError, ValueError):
        n_rows = 0
    return {"n_rows": n_rows, "v4": _model_side("v4"), "v5": _model_side("v5")}


def aggregate_legacy_periods_to_total(periods: Mapping[str, Any]) -> dict[str, Any]:
    """Pool MAE su tutti i trimestri legacy (pesato per ``mae_*_n``)."""
    if not isinstance(periods, dict) or not periods:
        return {"n_rows": 0, "v4": {"mae": {}, "mae_global": None}, "v5": {"mae": {}, "mae_global": None}}

    def _pool_model(model: str) -> dict[str, Any]:
        mae: dict[str, float | None] = {}
        mae_n: dict[str, int] = {}
        pooled_num = 0.0
        pooled_den = 0
        for off in ACCURACY_V4_DISPLAY_OFFSETS:
            lbl = horizon_label(off)
            if model == "v5" and off not in SIMULATION_V5_Q50_OFFSETS:
                mae[lbl] = None
                mae_n[lbl] = 0
                continue
            num = 0.0
            den = 0
            for blk in periods.values():
                if not isinstance(blk, dict):
                    continue
                mv = float_or_none(blk.get(f"mae_{model}_{lbl}"))
                try:
                    nv = int(blk.get(f"mae_{model}_{lbl}_n") or 0)
                except (TypeError, ValueError):
                    nv = 0
                if mv is not None and nv > 0:
                    num += float(mv) * nv
                    den += nv
            mae[lbl] = round(num / den, 2) if den else None
            mae_n[lbl] = den
            pooled_num += num
            pooled_den += den
        return {
            "mae": mae,
            "mae_n": mae_n,
            "mae_global": round(pooled_num / pooled_den, 2) if pooled_den else None,
            "mae_global_n": pooled_den,
        }

    n_rows = 0
    for blk in periods.values():
        if isinstance(blk, dict):
            try:
                n_rows = max(n_rows, int(blk.get("n_rows") or 0))
            except (TypeError, ValueError):
                pass
    return {"n_rows": n_rows, "v4": _pool_model("v4"), "v5": _pool_model("v5")}


def history_entry_total_block(entry: dict[str, Any]) -> dict[str, Any] | None:
    """Blocco TOTALE per un run nello storico (schema nuovo o legacy)."""
    if not isinstance(entry, dict):
        return None
    tot = entry.get("total")
    if isinstance(tot, dict) and run_snapshot_has_metrics({"total": tot}):
        return tot
    periods = entry.get("periods")
    if isinstance(periods, dict) and periods:
        agg = aggregate_legacy_periods_to_total(periods)
        return agg if run_snapshot_has_metrics({"total": agg}) else None
  # by_quarter only (no total saved): pool quarters
    by_q = entry.get("by_quarter")
    if isinstance(by_q, dict) and by_q:
        legacy_like: dict[str, Any] = {}
        for qk, blk in by_q.items():
            if not isinstance(blk, dict):
                continue
            flat: dict[str, Any] = {"n_rows": blk.get("n_rows")}
            for model in ("v4", "v5"):
                mb = blk.get(model) if isinstance(blk.get(model), dict) else {}
                mae_map = mb.get("mae") if isinstance(mb.get("mae"), dict) else {}
                mae_n_map = mb.get("mae_n") if isinstance(mb.get("mae_n"), dict) else {}
                for off in ACCURACY_V4_DISPLAY_OFFSETS:
                    lbl = horizon_label(off)
                    flat[f"mae_{model}_{lbl}"] = mae_map.get(lbl)
                    flat[f"mae_{model}_{lbl}_n"] = mae_n_map.get(lbl, 0)
            legacy_like[str(qk)] = flat
        agg = aggregate_legacy_periods_to_total(legacy_like)
        return agg if run_snapshot_has_metrics({"total": agg}) else None
    return None


def load_accuracy_summary_document(path: str | Path) -> dict[str, Any]:
    p = Path(path)
    if not p.is_file():
        return {"history": [], "latest": {}}
    try:
        raw = json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {"history": [], "latest": {}}
    return raw if isinstance(raw, dict) else {"history": [], "latest": {}}


def temporal_summary_table_headers() -> tuple[str, ...]:
    """Accuratezza temporale: tutti i MAE, poi tutti gli Hit % (non alternati)."""
    cols = ["Periodo", "Modello", "N righe"]
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        cols.append(f"MAE {horizon_label(off)}")
    cols.append("MAE globale")
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        cols.append(f"Hit % {horizon_label(off)}")
    cols.append("Hit % globale")
    return tuple(cols)


def summary_row_for_model_temporal(
    period_key: str,
    model: str,
    period_block: dict[str, Any],
) -> list[Any]:
    """Riga foglio temporale: blocchi MAE poi Hit %."""
    n_rows = period_block.get("n_rows")
    _model_key = "v5_raw" if model == "v5 raw" else model
    model_block = (
        period_block.get(_model_key)
        if isinstance(period_block.get(_model_key), dict)
        else {}
    )
    mae_map = model_block.get("mae") if isinstance(model_block.get("mae"), dict) else {}
    hit_map = model_block.get("hit_pct") if isinstance(model_block.get("hit_pct"), dict) else {}
    row: list[Any] = [period_key, model, n_rows]
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        lbl = horizon_label(off)
        if model in ("v5", "v5_raw") and off not in SIMULATION_V5_Q50_OFFSETS:
            row.append(None)
        else:
            row.append(mae_map.get(lbl))
    row.append(model_block.get("mae_global"))
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        lbl = horizon_label(off)
        if model in ("v5", "v5_raw") and off not in SIMULATION_V5_Q50_OFFSETS:
            row.append(None)
        else:
            row.append(hit_map.get(lbl))
    row.append(model_block.get("hit_pct_global"))
    return row


def temporal_run_rows(
    run_iso: str, total_block: dict[str, Any]
) -> tuple[list[Any], list[Any], list[Any] | None]:
    """Tre righe per run: v4, v5 (anchor/sync), v5 raw (opz.)."""
    v4 = summary_row_for_model_temporal("TOTALE", "v4", total_block)
    v5 = summary_row_for_model_temporal("TOTALE", "v5", total_block)
    v5_raw_blk = total_block.get("v5_raw") if isinstance(total_block, dict) else None
    v5_raw = None
    if isinstance(v5_raw_blk, dict) and v5_raw_blk.get("mae_global") is not None:
        v5_raw = summary_row_for_model_temporal("TOTALE", "v5 raw", total_block)
    return (
        [run_iso] + v4,
        [run_iso] + v5,
        [run_iso] + v5_raw if v5_raw else None,
    )


def temporal_sheet_headers() -> tuple[str, ...]:
    return ("Run",) + temporal_summary_table_headers()


def temporal_metric_column_kind(header: str) -> str | None:
    """``mae`` | ``hit`` | None per colonne non metriche."""
    h = str(header or "").strip()
    if h.startswith("MAE"):
        return "mae"
    if h.startswith("Hit"):
        return "hit"
    return None


# Accuratezza temporale — scale cromatiche (solo verde Hit %, solo rosso MAE).
_TEMPORAL_HIT_GREEN_LIGHT = (232, 245, 233)  # #E8F5E9
_TEMPORAL_HIT_GREEN_DARK = (27, 94, 32)  # #1B5E20
_TEMPORAL_MAE_RED_LIGHT = (255, 205, 210)  # #FFCDD2
_TEMPORAL_MAE_RED_DARK = (183, 28, 28)  # #B71C1C
_TEMPORAL_HIT_PCT_MIN = 0.0
_TEMPORAL_HIT_PCT_MAX = 100.0
_TEMPORAL_MAE_PP_MIN = 15.0
_TEMPORAL_MAE_PP_MAX = 35.0


def _temporal_lerp_rgb(
    lo: tuple[int, int, int],
    hi: tuple[int, int, int],
    t: float,
) -> tuple[int, int, int]:
    t = max(0.0, min(1.0, float(t)))
    return tuple(int(lo[i] + (hi[i] - lo[i]) * t) for i in range(3))


def _temporal_rgb_hex(rgb: tuple[int, int, int]) -> str:
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def _temporal_font_for_rgb(rgb: tuple[int, int, int], *, dark_hex: str, light_hex: str) -> str:
    lum = 0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]
    return light_hex if lum < 145 else dark_hex


def temporal_hit_fill_and_font(hit_pct: float) -> tuple[Any, Any]:
    """
    Scala **solo verde**: Hit % basso → verde chiaro; Hit % alto → verde scuro (buono).
    """
    from openpyxl.styles import Font, PatternFill

    span = _TEMPORAL_HIT_PCT_MAX - _TEMPORAL_HIT_PCT_MIN
    t = (
        (float(hit_pct) - _TEMPORAL_HIT_PCT_MIN) / span
        if span > 0
        else 0.5
    )
    rgb = _temporal_lerp_rgb(_TEMPORAL_HIT_GREEN_LIGHT, _TEMPORAL_HIT_GREEN_DARK, t)
    fill = PatternFill("solid", fgColor=_temporal_rgb_hex(rgb))
    font = Font(
        bold=True,
        size=10,
        color=_temporal_font_for_rgb(rgb, dark_hex="1B5E20", light_hex="FFFFFF"),
    )
    return fill, font


def temporal_mae_fill_and_font(mae_pp: float) -> tuple[Any, Any]:
    """
    Scala **solo rosso**: MAE basso → rosso chiaro; MAE alto → rosso scuro (peggio).
    """
    from openpyxl.styles import Font, PatternFill

    span = _TEMPORAL_MAE_PP_MAX - _TEMPORAL_MAE_PP_MIN
    t = (
        (float(mae_pp) - _TEMPORAL_MAE_PP_MIN) / span
        if span > 0
        else 0.5
    )
    rgb = _temporal_lerp_rgb(_TEMPORAL_MAE_RED_LIGHT, _TEMPORAL_MAE_RED_DARK, t)
    fill = PatternFill("solid", fgColor=_temporal_rgb_hex(rgb))
    font = Font(
        bold=True,
        size=10,
        color=_temporal_font_for_rgb(rgb, dark_hex="B71C1C", light_hex="FFFFFF"),
    )
    return fill, font


def style_temporal_metric_cell(
    cell: Any,
    header: str,
    value: Any,
    *,
    model: str,
    ref_value: Any = None,
) -> None:
    """
    Accuratezza temporale: colori **assoluti** per metrica (non v4 vs v5).

    - Colonne **Hit %** *T*: scala verde (chiaro → scuro) in proporzione al valore Hit %.
    - Colonne **MAE** *T*: scala rosso (chiaro → scuro) in proporzione al MAE (pp).
    """
    from openpyxl.styles import Alignment

    del model, ref_value  # confronto v4/v5 non usato per il fill

    kind = temporal_metric_column_kind(header)
    if kind is None or value is None:
        return
    try:
        v = float(value)
    except (TypeError, ValueError):
        return
    if not math.isfinite(v):
        return

    if kind == "hit":
        fill, font = temporal_hit_fill_and_font(v)
    elif kind == "mae":
        fill, font = temporal_mae_fill_and_font(v)
    else:
        return

    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="center", vertical="center")


def merge_summary_history(
    existing: dict | None,
    new_run: dict[str, Any],
) -> dict[str, Any]:
    """Append run snapshot; keep prior ``history`` entries."""
    base = dict(existing) if isinstance(existing, dict) else {}
    hist = list(base.get("history") or [])
    if run_snapshot_has_metrics(new_run):
        entry = {
            "run_iso": new_run.get("generated_at"),
            "by_quarter": new_run.get("by_quarter"),
            "total": new_run.get("total"),
        }
        hist.append(entry)
        base["history"] = hist[-120:]
        base["latest"] = new_run
    else:
        print(
            "[Accuracy] Run v4/v5 senza righe passate — storico JSON non aggiornato "
            f"(generated_at={new_run.get('generated_at')!r}).",
            flush=True,
        )
    return base


def save_accuracy_summary_json(path: str | Path, payload: dict[str, Any]) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    existing = (
        json.loads(p.read_text(encoding="utf-8")) if p.is_file() else None
    )
    if not run_snapshot_has_metrics(payload):
        print(
            f"[Accuracy] Skip salvataggio {p.name}: coorte Accuracy passata vuota.",
            flush=True,
        )
        return
    merged = merge_summary_history(existing, payload)
    p.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")


def summary_table_headers() -> tuple[str, ...]:
    """Headers for the temporal summary (v4 row, then v5 row per period)."""
    cols = ["Periodo", "Modello", "N righe"]
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        lbl = horizon_label(off)
        cols.append(f"MAE {lbl}")
        cols.append(f"Hit % {lbl}")
    cols.append("MAE globale")
    cols.append("Hit % globale")
    return tuple(cols)


def summary_row_for_model(
    period_key: str,
    model: str,
    period_block: dict[str, Any],
) -> list[Any]:
    """One sheet row: MAE + Hit % per horizon for v4 or v5 (blank T−5 on v5)."""
    n_rows = period_block.get("n_rows")
    model_block = period_block.get(model) if isinstance(period_block.get(model), dict) else {}
    mae_map = model_block.get("mae") if isinstance(model_block.get("mae"), dict) else {}
    hit_map = model_block.get("hit_pct") if isinstance(model_block.get("hit_pct"), dict) else {}
    row: list[Any] = [period_key, model, n_rows]
    for off in ACCURACY_V4_DISPLAY_OFFSETS:
        lbl = horizon_label(off)
        if model == "v5" and off not in SIMULATION_V5_Q50_OFFSETS:
            row.append(None)
            row.append(None)
        else:
            row.append(mae_map.get(lbl))
            row.append(hit_map.get(lbl))
    row.append(model_block.get("mae_global"))
    row.append(model_block.get("hit_pct_global"))
    return row


def summary_rows_for_period(period_key: str, period_block: dict[str, Any]) -> tuple[list[Any], list[Any]]:
    """Two consecutive rows: v4 then v5 (temporal comparison)."""
    return (
        summary_row_for_model(period_key, "v4", period_block),
        summary_row_for_model(period_key, "v5", period_block),
    )
