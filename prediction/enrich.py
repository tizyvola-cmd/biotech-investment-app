"""Past-prediction accuracy metadata enrichment (lazy orchestrator backend)."""
from __future__ import annotations

import json
import logging
import os
from typing import Any

from orchestrator_io_paths import DATA_DIR, FINAL_JSON

from prediction.errors import log_prediction_error

_log = logging.getLogger(__name__)

CLINICAL_XLSX = os.path.join(DATA_DIR, "biotech_clinical_openfda.xlsx")
CLINICAL_CSV = os.path.join(DATA_DIR, "biotech_clinical_openfda.csv")


def _try_pandas():
    try:
        import pandas as pd
        return pd
    except ImportError:
        return None


def load_enrich_dataframes(
    *,
    workbook_path: str | None = None,
    json_path: str | None = None,
) -> tuple[Any, Any]:
    """
    Load clinical + financial frames for enrich.

    Prefer ``biotech_orchestrated_output.json``; optional Financial sheet from workbook.
    """
    pd = _try_pandas()
    if pd is None:
        return None, None

    clinical_df = None
    financial_df = None

    jpath = json_path or FINAL_JSON
    if os.path.isfile(jpath):
        try:
            payload = json.loads(open(jpath, encoding="utf-8").read())
            if isinstance(payload, dict):
                clin_recs = payload.get("clinical_openfda") or []
                fin_recs = payload.get("financial") or []
                if clin_recs:
                    clinical_df = pd.DataFrame(clin_recs)
                if fin_recs:
                    financial_df = pd.DataFrame(fin_recs)
        except Exception as exc:
            log_prediction_error(f"load_enrich_json:{jpath}", exc)

    if clinical_df is None:
        for path, reader in ((CLINICAL_CSV, pd.read_csv), (CLINICAL_XLSX, pd.read_excel)):
            if os.path.isfile(path):
                try:
                    clinical_df = reader(path, dtype=str)
                    clinical_df.columns = [str(c).strip() for c in clinical_df.columns]
                    break
                except Exception as exc:
                    log_prediction_error(f"load_clinical:{path}", exc)

    if financial_df is None and workbook_path and os.path.isfile(workbook_path):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(workbook_path, read_only=True, data_only=True)
            for sheet in ("Financial", "financial"):
                if sheet in wb.sheetnames:
                    ws = wb[sheet]
                    headers = [
                        str(ws.cell(1, c).value or "").strip()
                        for c in range(1, (ws.max_column or 0) + 1)
                    ]
                    if not any(headers):
                        headers = [
                            str(ws.cell(3, c).value or "").strip()
                            for c in range(1, (ws.max_column or 0) + 1)
                        ]
                        start_row = 4
                    else:
                        start_row = 2
                    rows = []
                    for rn in range(start_row, (ws.max_row or 0) + 1):
                        row = {
                            headers[i - 1]: ws.cell(rn, i).value
                            for i in range(1, len(headers) + 1)
                            if headers[i - 1]
                        }
                        if any(v is not None and str(v).strip() for v in row.values()):
                            rows.append(row)
                    if rows:
                        financial_df = pd.DataFrame(rows)
                    break
            wb.close()
        except Exception as exc:
            log_prediction_error("load_financial_sheet", exc)

    return clinical_df, financial_df


def enrich_past_pred_accuracy_metadata(
    past_pred_data: dict | None,
    *,
    clinical_df=None,
    financial_df=None,
    workbook_path: str | None = None,
    json_path: str | None = None,
) -> bool:
    """
    Enrich rows in-place. Returns True if orchestrator enrich ran.

    Lazy-imports ``data_orchestrator._enrich_past_pred_accuracy_metadata``.
    """
    if not past_pred_data:
        return False
    if clinical_df is None and financial_df is None:
        clinical_df, financial_df = load_enrich_dataframes(
            workbook_path=workbook_path,
            json_path=json_path or FINAL_JSON,
        )
    from data_orchestrator import _enrich_past_pred_accuracy_metadata as _orch_enrich

    _orch_enrich(past_pred_data, clinical_df, financial_df)
    return True
