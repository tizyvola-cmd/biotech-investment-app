"""Lightweight JSON ↔ Excel alignment checks."""
from __future__ import annotations

import os
import re
from datetime import date, datetime

from past_pred_io import load_past_pred_map

_SIM_SHEET = "Simulation"
_KEY_RE = re.compile(r"^([A-Z0-9.\-]+)\|(\d{4}-\d{2}-\d{2})$", re.I)


def _parse_sim_cd(raw) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    s = str(raw).strip()
    if not s:
        return None
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def simulation_keys_from_workbook(workbook_path: str, *, max_rows: int = 5000) -> set[str]:
    """Harvest ``TICKER|YYYY-MM-DD`` keys from Simulation sheet (cols A + C)."""
    if not os.path.isfile(workbook_path):
        return set()
    try:
        from openpyxl import load_workbook
    except ImportError:
        return set()

    keys: set[str] = set()
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if _SIM_SHEET not in wb.sheetnames:
            return keys
        ws = wb[_SIM_SHEET]
        limit = min(int(ws.max_row or 0), max_rows)
        for rn in range(4, limit + 1):
            tk = str(ws.cell(rn, 1).value or "").strip().upper()
            if not tk or tk in ("—", "-", "N/D"):
                continue
            cds = _parse_sim_cd(ws.cell(rn, 3).value)
            if cds:
                keys.add(f"{tk}|{cds}")
    finally:
        wb.close()
    return keys


def json_prediction_keys(json_path: str) -> set[str]:
    rows = load_past_pred_map(json_path)
    keys: set[str] = set()
    for k in rows:
        m = _KEY_RE.match(str(k))
        if m:
            keys.add(f"{m.group(1).upper()}|{m.group(2)}")
    return keys


def check_json_excel_alignment(
    json_path: str,
    workbook_path: str,
    *,
    sample_limit: int = 5000,
) -> dict:
    """
    Compare ticker|CD keys in JSON vs Simulation sheet.

    Returns dict with ``json_only``, ``excel_only``, ``aligned_count``, ``warnings``.
    """
    jkeys = json_prediction_keys(json_path)
    ekeys = simulation_keys_from_workbook(workbook_path, max_rows=sample_limit)
    json_only = sorted(jkeys - ekeys)
    excel_only = sorted(ekeys - jkeys)
    warnings: list[str] = []
    if json_only:
        warnings.append(f"{len(json_only)} keys in JSON but not on Simulation.")
    if excel_only:
        warnings.append(f"{len(excel_only)} Simulation rows missing from JSON.")
    return {
        "json_key_count": len(jkeys),
        "excel_key_count": len(ekeys),
        "aligned_count": len(jkeys & ekeys),
        "json_only": json_only[:50],
        "excel_only": excel_only[:50],
        "json_only_total": len(json_only),
        "excel_only_total": len(excel_only),
        "warnings": warnings,
    }
