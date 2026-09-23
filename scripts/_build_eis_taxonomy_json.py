"""One-shot: build config/eis_event_taxonomy.json from the curated xlsx."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required", file=sys.stderr)
    raise

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\tizyv\Downloads\tassonomia_EIS_eventi.xlsx")
OUT = ROOT / "config" / "eis_event_taxonomy.json"

DIM_MAP = {
    "Clinica": "clinical",
    "Finanziaria": "financial",
    "Societaria": "corporate",
    "Market Access": "market_access",
}


def slug(*parts: str) -> str:
    s = "_".join(parts).lower()
    s = s.replace("'", "").replace("/", "_")
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")[:96]


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    events: list[dict] = []
    for row in wb["Tassonomia eventi"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        dim_label = str(row[0]).strip()
        dim = DIM_MAP[dim_label]
        et = str(row[1]).strip()
        mods_raw = str(row[4] or "")
        events.append(
            {
                "id": slug(dim, et),
                "dimension": dim,
                "dimension_label": dim_label,
                "event_type": et,
                "polarity": str(row[2] or "0").strip(),
                "base_weight": float(row[3]) if row[3] is not None else 0.0,
                "applicable_modifiers": [
                    x.strip() for x in mods_raw.split(";") if x.strip()
                ],
                "typical_example": str(row[5] or "").strip().strip('"'),
                "notes": row[6],
            }
        )

    mods: list[dict] = []
    for row in wb["Modificatori di grandezza"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        applies = str(row[3] or "")
        applies_dims: list[str] = []
        if "tutte" in applies.lower():
            applies_dims = list(DIM_MAP.values())
        else:
            for label, key in DIM_MAP.items():
                if label.lower() in applies.lower():
                    applies_dims.append(key)
        mods.append(
            {
                "id": slug(str(row[0]), str(row[1])),
                "category": row[0],
                "modifier": row[1],
                "multiplier": float(row[2]) if row[2] is not None else 1.0,
                "applies_to": applies_dims,
                "applies_to_label": applies,
            }
        )

    out = {
        "version": 1,
        "source": "tassonomia_EIS_eventi.xlsx",
        "score_scale": {"min": -3.0, "max": 3.0},
        "dimensions": [
            {"key": "clinical", "label": "Clinica"},
            {"key": "financial", "label": "Finanziaria"},
            {"key": "corporate", "label": "Societaria"},
            {"key": "market_access", "label": "Market Access"},
        ],
        "events": events,
        "modifiers": mods,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {OUT} events={len(events)} mods={len(mods)}")


if __name__ == "__main__":
    main()
