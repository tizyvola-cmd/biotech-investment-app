#!/usr/bin/env python3
"""Verifica prezzi acquisto portafoglio: JSON locale vs foglio orchestrato."""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def main() -> int:
    inp_path = ROOT / "data" / "invest_sim_inputs.json"
    xlsx_path = ROOT / "data" / "biotech_orchestrated_output.xlsx"
    if not inp_path.exists():
        print("Manca data/invest_sim_inputs.json")
        return 1

    raw = json.loads(inp_path.read_text(encoding="utf-8"))
    inputs = raw.get("inputs", raw)

    sheet_by_ticker: dict[str, dict] = {}
    if xlsx_path.exists():
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        if "Simulation" in wb.sheetnames:
            ws = wb["Simulation"]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

            def col(*parts: str) -> int | None:
                for i, h in enumerate(headers):
                    if not h:
                        continue
                    hl = str(h).lower()
                    if all(p in hl for p in parts):
                        return i
                return None

            ti = headers.index("Ticker") if "Ticker" in headers else None
            bi = col("prezzo", "acquisto")
            ci = col("prezzo", "corrente")
            pi = col("p&l", "%")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if ti is None:
                    break
                tk = str(row[ti] or "").strip().upper()
                if not tk or "TOTALE" in tk:
                    continue
                sheet_by_ticker[tk] = {
                    "sheet_buy": row[bi] if bi is not None else None,
                    "curr": row[ci] if ci is not None else None,
                    "pnl_pct": row[pi] if pi is not None else None,
                }
        wb.close()

    print(f"{'Ticker':<8} {'local_buy':>10} {'curr':>8} {'ratio':>7}  {'sheet_buy':>10} {'sheet_pnl%':>10}  flag")
    print("-" * 72)
    flags = 0
    for key in sorted(inputs):
        e = inputs[key]
        if not isinstance(e, dict):
            continue
        cap = float(e.get("capital") or 0)
        buy = float(e.get("buyPrice") or 0)
        if cap <= 0 and buy <= 0:
            continue
        tk = key.split("|")[0].upper()
        sh = sheet_by_ticker.get(tk, {})
        curr = sh.get("curr")
        ratio = (buy / curr) if buy and curr else None
        warn = ""
        if ratio and (ratio > 2.5 or ratio < 0.35):
            warn = "BUY/SPOT sospetto"
            flags += 1
        elif sh.get("pnl_pct") and abs(float(sh["pnl_pct"])) > 50 and buy and curr:
            inferred = float(curr) / (1 + float(sh["pnl_pct"]) / 100)
            if abs(inferred - buy) > 1:
                warn = "P&L foglio non allineato a buy locale"
                flags += 1
        print(
            f"{tk:<8} {buy:>10.2f} {curr or 0:>8.2f} "
            f"{ratio or 0:>7.2f}  {sh.get('sheet_buy') or '':>10} "
            f"{sh.get('pnl_pct') or '':>10}  {warn}"
        )
    print(f"\nPosizioni con capitale: {sum(1 for k,v in inputs.items() if isinstance(v,dict) and (v.get('capital') or 0)>0)}")
    print(f"Anomalie: {flags}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
