#!/usr/bin/env python3
"""Debug rapido: stato calibrazione, coorte pred, tempi blocchi orchestrator."""
from __future__ import annotations

import json
import os
import sys
import time
from datetime import date, timedelta
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

os.environ.setdefault("ORCH_SKIP_FETCH", "1")
os.environ.setdefault("ORCH_SKIP_FINANCIAL_ENRICH", "1")
os.environ.setdefault("ORCH_SKIP_RETROSPECTIVE", "1")
os.environ.setdefault("ORCH_SKIP_VARIATIONS_RETRY", "1")
os.environ.setdefault("ORCH_SKIP_SEC_K8", "1")
os.environ.setdefault("ORCH_SKIP_AUTO_DISCOVERY", "1")
os.environ.setdefault("ORCH_SKIP_IMMUTABLE_INDEX", "1")


def _section(title: str) -> None:
    print(f"\n{'=' * 60}\n{title}\n{'=' * 60}")


def main() -> int:
    import data_orchestrator as orch

    _section("1. File e stato calibrazione")
    cal_path = orch._CALIB_STATE_PATH
    print(f"  path: {cal_path}  exists={cal_path.exists()}")
    st = orch._load_calibration_state()
    has_curves = orch._state_has_curve_data(st)
    print(f"  curves presenti: {has_curves}")
    if isinstance(st, dict):
        v4 = (st.get("curves") or {}).get("v4_options") or {}
        for cat in ("success", "failure", "neutral", "control"):
            blk = v4.get(cat) or {}
            n = blk.get("n")
            med5 = (blk.get("median") or {}).get("5")
            print(f"    {cat}: n={n}  median[+5]={med5}")

    _section("2. past_pred + ensure curve (timed)")
    doc = orch._past_pred_disk_load()
    rows = doc.get("rows") if isinstance(doc.get("rows"), dict) else {}
    print(f"  past_pred rows: {len(rows)}")
    t0 = time.perf_counter()
    st2 = orch._ensure_calibration_state_for_empirical(st)
    dt = time.perf_counter() - t0
    print(f"  ensure_calibration_state_for_empirical: {dt:.1f}s")
    print(f"  curves dopo ensure: {orch._state_has_curve_data(st2)}")
    pe = orch._predict_empirical_curve_cat(st2, "v4_options", "success", 5)
    print(f"  sample pred_emp_d5 (success): {pe.get('pred_pct')}")

    _section("3. Coorte Simulation vs Accuracy (merge)")
    t0 = time.perf_counter()
    master_df, clinical_df = orch.merge_by_symbol()
    print(f"  merge_by_symbol: {time.perf_counter() - t0:.1f}s")
    financial_df = orch._build_financial_df_for_orchestrator(master_df)
    if clinical_df.empty:
        clinical_df_rich = clinical_df
    else:
        import pandas as pd
        clinical_df_rich = orch.add_modality_column(clinical_df)
        clinical_df_rich = orch.add_sponsor_match_column(clinical_df_rich)

    today = date.today()
    horizon_sim = orch.SIM_SHEET_DISPLAY_HORIZON_CAL_DAYS
    horizon_acc = orch._acc_sim_catalyst_horizon_days()

    sim_rows: list = []
    if not financial_df.empty:
        try:
            from simulation_core import build_rows_from_df

            sim_rows = build_rows_from_df(
                financial_df,
                catalyst_dates=orch._extract_catalyst_dates(
                    clinical_df_rich if not clinical_df_rich.empty else clinical_df,
                    horizon_calendar_days=horizon_sim,
                ),
            )
        except Exception as e:
            print(f"  build_rows_from_df: {e}")

    sim_future = []
    for r in sim_rows:
        if not isinstance(r, dict):
            continue
        cd = r.get("completion_date")
        if hasattr(cd, "date"):
            cd = cd.date()
        if isinstance(cd, date) and cd >= today:
            sm = str(r.get("sponsor_match") or "").lower()
            if sm in ("exact", "partial"):
                sim_future.append(r)

    sim_past = [r for r in sim_rows if isinstance(r, dict) and r.get("past_catalyst")]
    merged_acc = orch._sim_rows_merged_for_accuracy_sheet(
        clinical_df_rich, financial_df, sim_past, sim_future, log_supplement=False
    )
    acc_future = []
    for r in merged_acc:
        if not isinstance(r, dict):
            continue
        cd = r.get("completion_date")
        if hasattr(cd, "date"):
            cd = cd.date()
        if isinstance(cd, date) and cd >= today:
            sm = str(r.get("sponsor_match") or "").lower()
            if sm in ("exact", "partial"):
                acc_future.append(r)

    print(f"  Simulation horizon: {horizon_sim} gg")
    print(f"  Accuracy supplement horizon: {horizon_acc} gg")
    print(f"  sim_rows total: {len(sim_rows)}  future Exact/Partial: {len(sim_future)}")
    print(f"  merged_acc total: {len(merged_acc)}  future Exact/Partial: {len(acc_future)}")
    print(
        f"  pred live (Accuracy merge): {len(acc_future)} eventi futuri | "
        f"Simulation display: {len(sim_future)} (sponsor non arricchito in questo script)"
    )

    _section("4. _compute_price_predictions (solo _sim_rows, timed)")
    if not sim_future:
        print("  Nessun evento futuro Simulation — skip pred")
    else:
        t0 = time.perf_counter()
        pred = orch._compute_price_predictions(
            sim_future,
            calibration_state=st2,
            ticker_studies=None,
            financial_df=financial_df,
        )
        dt = time.perf_counter() - t0
        print(f"  _compute_price_predictions({len(sim_future)} righe): {dt:.1f}s")
        print(f"  chiavi pred output: {len(pred)}")
        n_emp = sum(
            1 for v in pred.values()
            if isinstance(v, dict) and v.get("pred_emp_d5") is not None
        )
        n_rel = sum(
            1 for v in pred.values()
            if isinstance(v, dict) and v.get("rel_emp_d5") is not None
        )
        print(f"  con pred_emp_d5: {n_emp}/{len(pred)}")
        print(f"  con rel_emp_d5: {n_rel}/{len(pred)}")
        if pred:
            sample = next(iter(pred.values()))
            if isinstance(sample, dict):
                print(
                    f"  esempio: pred_emp_d5={sample.get('pred_emp_d5')} "
                    f"rel_emp_d5={sample.get('rel_emp_d5')} "
                    f"d5_pct={sample.get('d5_pct')}"
                )

    _section("5. Excel Simulation (prime colonne empiriche)")
    xlsx = orch.FINAL_XLSX
    if Path(xlsx).is_file():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            if "Simulation" in wb.sheetnames:
                ws = wb["Simulation"]
                col_emp = orch.SIM_COL_AFFIDABILITA + 7
                col_rel = orch.SIM_COL_DELTA_MOD_EMP
                filled_emp = filled_rel = dash = 0
                for rn in range(4, min(ws.max_row + 1, 204)):
                    v1 = ws.cell(rn, col_emp).value
                    v2 = ws.cell(rn, col_rel).value
                    if v1 in (None, "", "—"):
                        dash += 1
                    else:
                        filled_emp += 1
                    if v2 in (None, "", "—"):
                        pass
                    else:
                        filled_rel += 1
                print(f"  righe dati (max 200): emp={filled_emp} rel={filled_rel} dash_emp~={dash}")
            else:
                print("  foglio Simulation assente")
            wb.close()
        except Exception as ex:
            print(f"  lettura Excel: {ex}")
    else:
        print(f"  Excel non trovato: {xlsx}")

    _section("6. Env utili per velocita")
    for k in (
        "ORCH_FAST_RELUNCH", "ORCH_SKIP_FETCH", "ORCH_SKIP_FINANCIAL_ENRICH",
        "ORCH_SKIP_RETROSPECTIVE", "ORCH_SKIP_SEC_K8", "ORCH_PERF",
        "ACC_SIM_SHEET_CATALYST_HORIZON_CAL_DAYS", "PRED_CURVE_SEQ_CALIB",
        "ORCH_OPTIONS_WORKERS",
    ):
        v = os.environ.get(k, "")
        if v:
            print(f"  {k}={v}")

    print("\nFine debug.\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
