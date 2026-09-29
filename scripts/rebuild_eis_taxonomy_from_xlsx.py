"""One-shot: rebuild config/eis_event_taxonomy.json from Excel v2 + keep guided readout."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\tizyv\Downloads\tassonomia_EIS_eventi_2.xlsx")
OUT = ROOT / "config" / "eis_event_taxonomy.json"

DIM_MAP = {
    "Clinica": "clinical",
    "Finanziaria": "financial",
    "Societaria": "corporate",
    "Market Access": "market_access",
}
MOD_APPLIES = {
    "Clinica": ["clinical"],
    "Finanziaria, Societaria": ["financial", "corporate"],
    "Market Access": ["market_access"],
    "Tutte": ["clinical", "financial", "corporate", "market_access"],
}
STABLE_MOD = {
    ("Fase trial", "Fase 1"): "trial_phase_phase1",
    ("Fase trial", "Fase 2"): "trial_phase_phase2",
    ("Fase trial", "Fase 3 / pivotal"): "trial_phase_phase3",
    ("Fase trial", "Registrational / confirmatory"): "trial_phase_registrational",
    ("Significativita' statistica", "p<0.05 (soglia standard)"): "statistical_significance_p_lt_0_05",
    ("Significativita' statistica", "p<0.01"): "statistical_significance_p_lt_0_01",
    ("Significativita' statistica", "Trend / non significativo"): "statistical_significance_trend_only",
    ("Gerarchia endpoint", "Primario"): "endpoint_type_primary",
    ("Gerarchia endpoint", "Secondario"): "endpoint_type_secondary",
    ("Gerarchia endpoint", "Esplorativo / post-hoc"): "endpoint_type_exploratory",
    ("Aspettativa di mercato", "Atteso / gia' prezzato"): "surprise_vs_expected_expected",
    ("Aspettativa di mercato", "Sorpresa"): "surprise_vs_expected_surprise",
    ("Importo finanziario (% market cap)", "< 5% market cap"): "amount_pct_market_cap_under_5",
    ("Importo finanziario (% market cap)", "5-20% market cap"): "amount_pct_market_cap_5_to_20",
    ("Importo finanziario (% market cap)", "> 20% market cap"): "amount_pct_market_cap_over_20",
    ("Popolazione paziente coperta", "Ristretta / nicchia"): "patient_population_narrow",
    ("Popolazione paziente coperta", "Ampia / mercato primario"): "patient_population_broad",
}


def _slug(cat: str, mod: str) -> str:
    s = f"{cat}_{mod}".lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")[:80]


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Tassonomia eventi"]
    events: list[dict] = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        events.append(
            {
                "id": str(r[0]).strip(),
                "dimension": DIM_MAP[str(r[1]).strip()],
                "dimension_label": str(r[1]).strip(),
                "event_type": str(r[2] or "").strip(),
                "polarity": str(r[3] or "0").strip(),
                "base_weight": float(r[4] or 0),
                "applicable_modifiers": [
                    x.strip() for x in str(r[5] or "").split(";") if x and str(x).strip()
                ]
                if r[5]
                else [],
                "typical_example": (str(r[6]).strip().strip('"') if r[6] else None),
                "notes": (str(r[7]).strip() if r[7] else None),
            }
        )
    events.append(
        {
            "id": "CLIN_READOUT_GUIDED",
            "dimension": "clinical",
            "dimension_label": "Clinica",
            "event_type": "Readout atteso con finestra dichiarata",
            "polarity": "0",
            "base_weight": 0.6,
            "applicable_modifiers": ["Fase trial"],
            "typical_example": "expects to report Phase 3 topline results in 2H 2026",
            "notes": "Catalyst scheduled, not yet occurred — low weight",
        }
    )

    ws2 = wb["Modificatori di grandezza"]
    modifiers: list[dict] = []
    for r in list(ws2.iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        cat, mod, mult, applies = str(r[0]), str(r[1]), float(r[2]), str(r[3] or "")
        mid = STABLE_MOD.get((cat, mod)) or _slug(cat, mod)
        modifiers.append(
            {
                "id": mid,
                "category": cat,
                "modifier": mod,
                "multiplier": mult,
                "applies_to": MOD_APPLIES.get(
                    applies, ["clinical", "financial", "corporate", "market_access"]
                ),
            }
        )

    old = json.loads(OUT.read_text(encoding="utf-8"))
    new_by_key = {(e["dimension"], e["event_type"].lower()): e["id"] for e in events}
    aliases: dict[str, str] = {}
    for e in old["events"]:
        k = (e["dimension"], str(e.get("event_type") or "").lower())
        if k in new_by_key and e["id"] != new_by_key[k]:
            aliases[e["id"]] = new_by_key[k]
    aliases["clinical_readout_attesa_con_finestra_dichiarata"] = "CLIN_READOUT_GUIDED"

    # Old modifier id → new (best-effort by category+multiplier)
    old_mod_aliases: dict[str, str] = {
        "fase_trial_fase_1": "trial_phase_phase1",
        "fase_trial_fase_2": "trial_phase_phase2",
        "fase_trial_fase_3_pivotal": "trial_phase_phase3",
        "fase_trial_registrational_confirmatory": "trial_phase_registrational",
        "significativita_statistica_p_0_05": "statistical_significance_p_lt_0_05",
        "significativita_statistica_p_0_01": "statistical_significance_p_lt_0_01",
        "significativita_statistica_trend_non_significativo": "statistical_significance_trend_only",
        "gerarchia_endpoint_primario": "endpoint_type_primary",
        "gerarchia_endpoint_secondario": "endpoint_type_secondary",
        "gerarchia_endpoint_esplorativo_post_hoc": "endpoint_type_exploratory",
        "aspettativa_di_mercato_atteso_gia_prezzato": "surprise_vs_expected_expected",
        "aspettativa_di_mercato_sorpresa": "surprise_vs_expected_surprise",
        "importo_finanziario_pct_market_cap_under_5": "amount_pct_market_cap_under_5",
        "importo_finanziario_pct_market_cap_5_20": "amount_pct_market_cap_5_to_20",
        "importo_finanziario_pct_market_cap_over_20": "amount_pct_market_cap_over_20",
        "popolazione_paziente_coperta_ristretta_nicchia": "patient_population_narrow",
        "popolazione_paziente_coperta_ampia_mercato_primario": "patient_population_broad",
    }
    aliases.update(old_mod_aliases)

    out = {
        "version": 3,
        "source": "tassonomia_EIS_eventi_2.xlsx",
        "score_scale": {"min": -3.0, "max": 3.0},
        "dimensions": [
            {"key": "clinical", "label": "Clinica"},
            {"key": "financial", "label": "Finanziaria"},
            {"key": "corporate", "label": "Societaria"},
            {"key": "market_access", "label": "Market Access"},
        ],
        "events": events,
        "modifiers": modifiers,
        "id_aliases": aliases,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {OUT} events={len(events)} mods={len(modifiers)} aliases={len(aliases)}")


if __name__ == "__main__":
    main()
