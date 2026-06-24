#!/usr/bin/env python3
"""
Ripara ``data/biotech_orchestrated_output.xlsx`` se Excel non lo apre
(formato non valido / corrotto). Crea un backup prima di sostituire.
"""
from __future__ import annotations

import os
import shutil
import sys
from datetime import datetime

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from orchestrator_io_paths import FINAL_XLSX


def main() -> int:
    src = os.path.abspath(FINAL_XLSX)
    if not os.path.isfile(src):
        print(f"File non trovato: {src}")
        return 1

    from openpyxl import load_workbook
    from data_orchestrator import _sanitize_openpyxl_workbook_string_cells

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = os.path.join(
        os.path.dirname(src),
        f"biotech_orchestrated_output__backup_{stamp}.xlsx",
    )
    tmp = os.path.join(os.path.dirname(src), "_repair_orchestrated_tmp.xlsx")

    print(f"Backup → {bak}")
    shutil.copy2(src, bak)

    print("Caricamento e riscrittura (senza VBA su .xlsx)…")
    wb = load_workbook(src, keep_vba=False)
    n = _sanitize_openpyxl_workbook_string_cells(wb)
    if n:
        print(f"  Celle sanitize: {n}")
    wb.save(tmp)
    wb.close()

    load_workbook(tmp, read_only=True).close()
    os.replace(tmp, src)
    print(f"OK — file riparato: {src} ({os.path.getsize(src):,} byte)")
    print("Apri di nuovo in Excel. Se ancora errore, prova il backup o:")
    print(f"  data\\biotech_orchestrated_output__grafici_staged_*.xlsx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
