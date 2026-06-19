#!/usr/bin/env python3
"""
Aggiunge al workbook il foglio «Pannello refresh» con istruzioni e collegamenti VBS.

Uso (Excel chiuso sul file principale):
  py -3 scripts/write_excel_refresh_panel.py
  py -3 scripts/write_excel_refresh_panel.py --workbook data/biotech_orchestrated_output.xlsx
"""
from __future__ import annotations

import argparse
import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

PANEL_SHEET = "Pannello refresh"
ROWS = (
    ("Refresh giornaliero", "daily", "Simulation + Accuracy. Preserva Prezzo acquisto / Capitale. Macro: BiotechRefreshDaily"),
    ("Solo Accuracy", "accuracy", "Solo foglio Accuracy. Macro: BiotechRefreshAccuracy"),
    ("Solo SEC K-8", "sec_k8", "Foglio filing 8-K. Macro: BiotechRefreshSecK8"),
    ("Domenica (full)", "sunday", "Orchestrator completo 30–90+ min. Macro: BiotechRefreshSunday"),
)


def write_panel(wb_path: str) -> bool:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink

    from orchestrator_io_paths import FINAL_XLSX

    tgt = os.path.abspath(wb_path or FINAL_XLSX)
    if not os.path.isfile(tgt):
        print(f"[Pannello] File non trovato: {tgt}", file=sys.stderr)
        return False

    excel_dir = os.path.join(_ROOT, "excel")
    wb = load_workbook(tgt, read_only=False, keep_vba=False)
    if PANEL_SHEET in wb.sheetnames:
        del wb[PANEL_SHEET]
    ws = wb.create_sheet(PANEL_SHEET, 0)

    hdr_fill = PatternFill("solid", fgColor="1565C0")
    hdr_font = Font(bold=True, color="FFFFFF", size=12)
    ws["A1"] = "Pannello refresh — Biotech Investment"
    ws["A1"].font = hdr_font
    ws["A1"].fill = hdr_fill
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 28

    ws["A2"] = (
        "Assegna un pulsante (Sviluppatore) alla macro indicata oppure doppio clic sul link VBS. "
        "Chiudi questo workbook prima del run se la macro lo richiede."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 36

    headers = ("Azione", "Avvio rapido (VBS)", "Macro Excel (Personal.xlsb)", "Note")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3F2FD")

    _vbs_map = {
        "daily": "Avvia_Refresh_Daily.vbs",
        "accuracy": "Avvia_Refresh_Accuracy.vbs",
        "sec_k8": "Avvia_Refresh_SecK8.vbs",
        "sunday": "Avvia_Refresh_Domenica.vbs",
    }
    for i, (label, profile, note) in enumerate(ROWS, start=5):
        ws.cell(row=i, column=1, value=label)
        vbs = os.path.join(excel_dir, _vbs_map.get(profile, "Avvia_Refresh_Daily.vbs"))
        vbs_abs = os.path.abspath(vbs)
        c2 = ws.cell(row=i, column=2, value="▶ Avvia")
        if os.path.isfile(vbs_abs):
            c2.hyperlink = Hyperlink(ref=c2.coordinate, target=vbs_abs)
            c2.font = Font(color="0563C1", underline="single")
        macro_name = {
            "daily": "BiotechRefreshDaily",
            "accuracy": "BiotechRefreshAccuracy",
            "sec_k8": "BiotechRefreshSecK8",
            "sunday": "BiotechRefreshSunday",
        }.get(profile, "")
        ws.cell(row=i, column=3, value=macro_name)
        ws.cell(row=i, column=4, value=note)
        ws.row_dimensions[i].height = 32

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 52

    ws.cell(row=10, column=1, value="Stato ultimo refresh (giornaliero / accuracy):")
    ws.cell(row=11, column=1, value=os.path.join(_ROOT, "data", "refresh_fast_status.txt"))
    ws.cell(
        row=12,
        column=1,
        value="Macro: Alt+F11 → Importa excel\\BiotechRefreshMacros.bas → pulsante → macro in PERSONAL.XLSB.",
    )
    ws.merge_cells("A12:D12")

    try:
        wb.save(tgt)
        print(f"[Pannello] Foglio «{PANEL_SHEET}» scritto in {tgt}")
        return True
    except PermissionError:
        print(
            "[Pannello] ERRORE: chiudi Excel sul file e rilancia.",
            file=sys.stderr,
        )
        return False


def main() -> int:
    from orchestrator_io_paths import FINAL_XLSX

    ap = argparse.ArgumentParser(description="Foglio Pannello refresh nel workbook.")
    ap.add_argument("--workbook", default=FINAL_XLSX)
    args = ap.parse_args()
    return 0 if write_panel(args.workbook) else 1


if __name__ == "__main__":
    raise SystemExit(main())
