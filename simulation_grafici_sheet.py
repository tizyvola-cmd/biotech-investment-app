"""
Foglio Excel «Grafici»: coorte, tabelle sorgente Δ% (2a–2c), grafici Excel.
I dati % sono sullo stesso foglio (area bassa); il foglio «Simulation — grafici» non è più usato.
"""
from __future__ import annotations

import os
from datetime import date
from typing import Any


def _grafici_fast_enabled() -> bool:
    """GRAFICI_FAST=1: niente refresh Yahoo/HistLib (solo pickle/cache locali)."""
    return os.environ.get("GRAFICI_FAST", "").strip().lower() in (
        "1",
        "true",
        "yes",
        "on",
    )


def _grafici_step(msg: str) -> None:
    print(msg, flush=True)

from prediction.errors import log_prediction_error

from data_orchestrator import (
    FMT_VARIATION,
    GRAFICI_SHEET,
    SIMULATION_36_HEADERS,
    SIMULATION_GRAFICI_SHEET,
    SIM_GRAFICI_CURVE_OFFSETS,
    _SN7_REF_CAL_X,
    _accuracy_sim_enrich_curve_records,
    _accuracy_sim_histlib_backfill_session_closes,
    _accuracy_sim_impute_missing_pre_cd_model_pcts,
    _accuracy_sim_merge_live_pred,
    _accuracy_sim_synthesize_interp_nodes_from_post_d_only,
    _overlay_model_pcts_from_simulation_row,
    _emit_magnitude_bucket_summary_table,
    _emit_posthoc_regression_summary_table,
    _mag_bucket_bundle_active,
    _past_pred_key_from_harvest_row,
    _ristretta_curve_table_points,
    _sim_harvest_rows_from_workbook,
    _grafici_ensure_seq_curve_on_pairwise,
    _grafici_refresh_histlib_from_yfinance,
    _sn7_cd7_overlay_seq_curve_from_live,
    _sn7_emit_pool_company_tables,
    _sn7_emit_sim_grafici_fit_readout_table,
    _sn7_parse_sheet_date_ddmmyyyy,
    _sn7_ref_mu_at_predizione_table_offsets,
    _stub_past_pred_row_from_sim_harvest,
    _sim_grafici_price_usd_from_pct_m60,
)

_SIMULATION_GRAFICI_MAX_SERIES = 255
_GRAF_WIDE_TABLE_START_ROW = 520
_CHART_STACK_GAP = 18
# Colonne O–R: sorgenti grafico Var.% (1g · 1M · 3M · 6M) sulla stessa sheet.
_VAR_CHART_COL0 = 15
_VAR_CHART_LABELS = ("Var.\n1g", "Var.\n1M", "Var.\n3M", "Var.\n6M")
_CO_B_NONE_XL = "(nessuna — solo società A)"
_LIST_COL = 26  # Z — elenco ticker (menu a tendina)
_PANEL_LAB_COL = 19  # S — etichette pannello
_PANEL_VAL_COL = 20  # T — valori / menu
_N_EXPLORER_PCT_SERIES = 10  # A×3 + B×3 + 4 μ
_REF_SPECS: tuple[tuple[str, str], ...] = (
    ("μ Globale primaria", "primary"),
    ("μ Cluster 0", "cluster0"),
    ("μ SuperNova (cl.1)", "cluster1"),
    ("μ controllo negativo", "control"),
    ("μ Post-CD rialzo", "success"),
    ("μ Post-CD ribasso", "failure"),
    ("μ Post-CD neutro", "neutral"),
)
_REF_FALLBACK_NAMES: dict[str, str] = {
    "Globale primaria": "primary",
    "Cluster 0": "cluster0",
    "SuperNova (cl.1)": "cluster1",
    "controllo negativo": "control",
    "Post-CD rialzo": "success",
    "Post-CD ribasso": "failure",
    "Post-CD neutro": "neutral",
}
_PREDIZIONE_GUIDA_SHEET = "Predizione — guida"
# Righe μ su «Predizione — guida» (col. A, senza σ) — fallback se scan assente
_PRED_MU_ROW_NEEDLES: dict[str, tuple[str, ...]] = {
    "μ Globale primaria": ("globale primaria",),
    "μ Cluster 0": (
        "μ · cluster 0",
        "cluster 0 — maggioranza",
        "coorte intera (senza partizione",
    ),
    "μ SuperNova (cl.1)": (
        "cluster 1 — supernova",
        "μ · cluster 1",
        "supernova (minoranza",
    ),
    "μ controllo negativo": ("controllo negativo",),
    "μ Post-CD rialzo": ("μ · post-cd rialzo", "post-cd rialzo"),
    "μ Post-CD ribasso": ("μ · post-cd ribasso", "post-cd ribasso"),
    "μ Post-CD neutro": ("μ · post-cd neutro", "post-cd neutro", "post-cd laterale"),
}


def _pad10(rr: list | None) -> list:
    if not isinstance(rr, list):
        return [None] * 10
    return list(rr[:10]) + [None] * (10 - len(rr))


def _table8_to_ref10(vals: list | None) -> list:
    from ristretta_lab_data import TABLE_OFFSETS

    if not isinstance(vals, list):
        return [None] * len(_SN7_REF_CAL_X)
    lut = {
        int(off): vals[i]
        for i, off in enumerate(TABLE_OFFSETS)
        if i < len(vals)
    }
    return [lut.get(int(x)) for x in _SN7_REF_CAL_X]


def _predizione_mu_menu_label(raw: str) -> str:
    """Etichetta breve per menu grafici (da col. A «Predizione — guida»)."""
    _s = str(raw or "").strip()
    if _s.startswith("μ ·"):
        _s = "μ " + _s[3:].strip()
    elif _s.startswith("μ·"):
        _s = "μ " + _s[2:].strip()
    if "(" in _s:
        _s = _s.split("(", 1)[0].strip()
    return _s[:72] if _s else "μ"


def _scan_predizione_mu_series(wb) -> list[tuple[str, int]]:
    """Tutte le righe μ (non σ) su «Predizione — guida» → (etichetta menu, n° riga)."""
    _out: list[tuple[str, int]] = []
    if wb is None:
        return _out
    _ws = None
    for _sn in wb.sheetnames:
        if str(_sn).strip() == _PREDIZIONE_GUIDA_SHEET:
            _ws = wb[_sn]
            break
    if _ws is None:
        return _out
    for _rn in range(1, int(_ws.max_row or 0) + 1):
        _raw = str(_ws.cell(row=_rn, column=1).value or "").strip()
        if not _raw:
            continue
        _low = _raw.lower()
        if _raw.startswith("σ") or _low.startswith("sigma") or _low.startswith("% oss"):
            continue
        if not (_raw.startswith("μ") or _raw.startswith("\u03bc")):
            continue
        _lab = _predizione_mu_menu_label(_raw)
        if not any(_lab == _x[0] for _x in _out):
            _out.append((_lab, _rn))
    return _out


def _sn7_ref_mu_series_has_signal(vals: list | None) -> bool:
    """True se la serie μ ha almeno un punto numerico non banale."""
    if not isinstance(vals, list):
        return False
    for _v in vals:
        if isinstance(_v, (int, float)) and _v == _v and abs(float(_v)) > 1e-9:
            return True
    return False


def _coalesce_sn7_ref_bundle(bundle: dict[str, list]) -> dict[str, list]:
    """
    Se cluster0 (o cluster1) è vuoto ma «primary» è valorizzato, usa la μ globale.
    Copre workbook «Predizione — guida» generati prima del fix guida_clusters_to_render.
    """
    if not isinstance(bundle, dict):
        return {}
    _out = dict(bundle)
    _primary = _pad10(_out.get("primary"))
    if not _sn7_ref_mu_series_has_signal(_primary):
        return _out
    for _key in ("cluster0", "cluster1"):
        if _sn7_ref_mu_series_has_signal(_out.get(_key)):
            continue
        _out[_key] = list(_primary)
        print(
            f"[Grafici] μ {_key} vuota in riferimento — fallback a μ globale primaria.",
            flush=True,
        )
    return _out


def _sn7_ref_mu_from_predizione_workbook(wb) -> dict[str, list] | None:
    """Legge le μ di riferimento dalla tabella numerica «Predizione — guida»."""
    if wb is None:
        return None
    _rows = _find_predizione_mu_rows(wb)
    if not _rows:
        return None
    _ws = None
    for _sn in wb.sheetnames:
        if str(_sn).strip() == _PREDIZIONE_GUIDA_SHEET:
            _ws = wb[_sn]
            break
    if _ws is None:
        return None
    _key_map = dict(_REF_SPECS)
    _out: dict[str, list] = {}
    _n = len(_SN7_REF_CAL_X)
    for _mlab, _key in _REF_SPECS:
        _rn = int(_rows.get(_mlab) or 0)
        if not _rn:
            continue
        _vals: list[float | None] = []
        for _j in range(_n):
            _v = _ws.cell(row=_rn, column=2 + _j).value
            if isinstance(_v, (int, float)) and _v == _v:
                _vals.append(float(_v))
            else:
                _vals.append(None)
        if any(v is not None for v in _vals):
            _out[_key_map.get(_mlab, _key)] = _vals
    if not _out:
        return None
    return _coalesce_sn7_ref_bundle(_out)


def _resolve_sn7_ref_mu(sn7_ref_mu: dict | None, wb=None) -> dict[str, list]:
    _from_pred = _sn7_ref_mu_from_predizione_workbook(wb)
    if _from_pred:
        return _coalesce_sn7_ref_bundle(_from_pred)
    if isinstance(sn7_ref_mu, dict):
        _has = any(
            isinstance(sn7_ref_mu.get(k), list) and sn7_ref_mu.get(k)
            for k in (
                "primary",
                "cluster0",
                "cluster1",
                "control",
                "success",
                "failure",
            )
        )
        if _has:
            return _coalesce_sn7_ref_bundle(sn7_ref_mu)
    try:
        import json
        from orchestrator_io_paths import DATA_DIR

        _p = DATA_DIR / "model_calibration_state.json"
        if _p.is_file():
            _doc = json.loads(_p.read_text(encoding="utf-8"))
            _cur = _doc.get("current") if isinstance(_doc.get("current"), dict) else _doc
            _cached = None
            if isinstance(_cur, dict):
                _cached = _cur.get("simulation_grafici_sn7_ref_mu")
            if _cached is None:
                _cached = _doc.get("simulation_grafici_sn7_ref_mu")
            if isinstance(_cached, dict) and _cached:
                return _coalesce_sn7_ref_bundle(_cached)
    except Exception as exc:
        log_prediction_error("grafici_sn7_ref_mu:calib_state", exc)
    try:
        from ristretta_lab_data import build_reference_curves, load_past_pred_map

        _refs = build_reference_curves(load_past_pred_map())
        _out: dict[str, list] = {}
        for _lab, _key in _REF_FALLBACK_NAMES.items():
            _curve = _refs.get(_lab)
            if isinstance(_curve, list) and _curve:
                _out[_key] = _table8_to_ref10(_curve)
        if _out:
            return _coalesce_sn7_ref_bundle(_out)
    except Exception as exc:
        log_prediction_error("grafici_sn7_ref_mu:ristretta_lab", exc)
    if isinstance(sn7_ref_mu, dict) and sn7_ref_mu:
        return _coalesce_sn7_ref_bundle(sn7_ref_mu)
    return {}


def _section_gap(ws_g, row: int) -> int:
    """Una riga vuota tra sezioni."""
    return int(row) + 1


def _write_section_title(
    ws_g,
    row: int,
    section_no: str,
    title: str,
    *,
    merge_cols: int = 15,
    fill: str = "E7EEF7",
) -> int:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws_g.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=merge_cols
    )
    _c = ws_g.cell(row=row, column=1, value=f"{section_no} — {title}")
    _c.font = Font(bold=True, size=10, color="1F3864")
    _c.fill = PatternFill("solid", fgColor=fill)
    _c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return row + 1


def _write_points_table(
    ws_g,
    start_row: int,
    section_no: str,
    company_curves: list[dict[str, Any]],
    *,
    pct_fmt: str,
    usd_fmt: str,
    date_fmt: str,
    sn7_ref_mu: dict | None = None,
    offsets: tuple[int, ...] | None = None,
    wb=None,
) -> tuple[int, int, int]:
    """Dettaglio per punto (K-8 inclusi). Ritorna (title_row, hdr_row, next_row)."""
    from openpyxl.styles import Alignment, Font, PatternFill

    _n_pt = sum(len(e.get("points") or []) for e in company_curves)
    _title_row = int(start_row)
    _hdr_r = _write_section_title(
        ws_g,
        start_row,
        section_no,
        f"Dettaglio punti — {len(company_curves)} società, {_n_pt} righe "
        "(poi μ Cluster 0 / cl.1 / post-CD in coda; 8-K da SEC K-8)",
    )
    _data_r = _hdr_r + 1
    _hdrs = (
        "Ticker",
        "CD",
        "Punto",
        "Gg da CD",
        "Data filing\n8-K (SEC)",
        "Pos. tra\nnodi chiave",
        "Gg\nricalib.",
        "Data\nricalib.",
        "Ricalibrazioni\n(cumul. fino a qui)",
        "Nodo",
        "Data cal.\n(CD+gg)",
        "Data mercato\n(raw)",
        "Origine curva",
        "% curva\nvs T−60",
        "% storico",
        "% modello",
        "Prezzo\n($)",
    )
    for _ci, _h in enumerate(_hdrs, start=1):
        _hc = ws_g.cell(row=_hdr_r, column=_ci, value=_h)
        _hc.font = Font(bold=True, size=9, color="FFFFFF")
        _hc.fill = PatternFill("solid", fgColor="1F3864")
        _hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _row = _data_r
    for _ent in company_curves:
        _tk = _ent["ticker"]
        _cd = _ent["cd"]
        _cd_s = _cd.strftime("%d/%m/%Y") if hasattr(_cd, "strftime") else "—"
        for _p in _ent["points"]:
            _dc = _p.get("data_cal")
            _dr = _p.get("data_raw")
            _gg_cd = _p.get("offset")
            if str(_p.get("nodo") or "") == "K-8" and _p.get("offset_filing") is not None:
                _gg_cd = _p.get("offset_filing")
            _vals = (
                _tk,
                _cd_s,
                str(_p.get("label") or "—"),
                _gg_cd,
                _p.get("k8_filing_date"),
                str(_p.get("pos_temporale") or "—"),
                _p.get("gg_ricalib"),
                _p.get("data_ricalib"),
                str(_p.get("ricalib_cumul") or "—"),
                str(_p.get("nodo") or "—"),
                _dc,
                _dr,
                str(_p.get("tipo") or "—"),
                _p.get("pct_curva"),
                _p.get("pct_reale"),
                _p.get("pct_modello"),
                _p.get("price_usd"),
            )
            for _ci, _vv in enumerate(_vals, start=1):
                _c = ws_g.cell(row=_row, column=_ci)
                if _ci in (5, 8, 11, 12) and hasattr(_vv, "strftime"):
                    _c.value = _vv
                    _c.number_format = date_fmt
                elif _ci in (14, 15, 16) and isinstance(_vv, (int, float)) and _vv == _vv:
                    _c.value = float(_vv) / 100.0
                    _c.number_format = pct_fmt
                elif _ci == 17 and isinstance(_vv, (int, float)) and _vv == _vv:
                    _c.value = float(_vv)
                    _c.number_format = usd_fmt
                else:
                    _c.value = _vv if _vv is not None else "—"
            _row += 1
    if offsets and (wb is not None or (isinstance(sn7_ref_mu, dict) and sn7_ref_mu)):
        _row = _write_control_points_after_companies(
            ws_g, _row, sn7_ref_mu or {}, offsets, pct_fmt=pct_fmt, wb=wb
        )
    return _title_row, _hdr_r, _row


def _write_control_points_after_companies(
    ws_g,
    row: int,
    sn7_ref_mu: dict,
    offsets: tuple[int, ...],
    *,
    pct_fmt: str,
    wb=None,
) -> int:
    """Dopo le società: nodi μ di riferimento (Predizione — guida)."""
    return _write_grafici_mu_source_rows(
        ws_g,
        row,
        wb=wb,
        offsets=offsets,
        sn7_ref_mu=sn7_ref_mu,
        num_fmt=pct_fmt,
        detail_layout=True,
    )


def _write_curve_wide_table(
    ws_g,
    start_row: int,
    section_no: str,
    company_curves: list[dict[str, Any]],
    sn7_ref_mu: dict,
    offsets: tuple[int, ...],
    *,
    num_fmt: str,
    value_field: str = "pct_curva",
    include_mu: bool = True,
    table_label: str = "curva ricalibr.",
    value_divisor: float = 100.0,
    universe: dict[str, Any] | None = None,
    price_mode: str | None = None,
    wb=None,
) -> tuple[int, int, int]:
    """Tabelle sorgente per INDEX/MATCH. Ritorna (hdr_row, first_data_row, next_row)."""
    from openpyxl.styles import Alignment, Font, PatternFill

    _hdr_r = _write_section_title(
        ws_g,
        start_row,
        section_no,
        f"Sintesi Δ% vs T−60 — {table_label}",
        fill="E8EAF6",
    )
    _data_r = _hdr_r + 1
    _hdrs = ["Ticker / serie", "CD"] + [
        f"Δ {o:+d}" if int(o) > 0 else f"Δ {int(o)}" for o in offsets
    ]
    for _ci, _h in enumerate(_hdrs, start=1):
        _hc = ws_g.cell(row=_hdr_r, column=_ci, value=_h)
        _hc.font = Font(bold=True, size=9, color="FFFFFF")
        _hc.fill = PatternFill("solid", fgColor="5E35B1")
        _hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _row = _data_r
    for _ent in company_curves:
        _tk = _ent["ticker"]
        _cd = _ent["cd"]
        _cd_s = _cd.strftime("%d/%m/%Y") if hasattr(_cd, "strftime") else "—"
        _std = {
            int(p["offset"]): p
            for p in (_ent.get("points") or [])
            if p.get("nodo") == "standard"
        }
        ws_g.cell(row=_row, column=1, value=_tk)
        ws_g.cell(row=_row, column=2, value=_cd_s)
        _pk = _ent.get("pk")
        _pw = (
            dict(universe[_pk])
            if isinstance(universe, dict) and _pk and _pk in universe
            else {}
        )
        if price_mode:
            _prices = _grafici_price_series_at_offsets(
                list(_ent.get("points") or []), offsets, price_mode, _pw
            )
            for _ji, _px in enumerate(_prices):
                _ce = ws_g.cell(row=_row, column=3 + _ji)
                if _px is not None and _px == _px:
                    _ce.value = float(_px)
                    _ce.number_format = num_fmt
        else:
            for _ji, _off in enumerate(offsets):
                _pv = (_std.get(int(_off)) or {}).get(value_field)
                _ce = ws_g.cell(row=_row, column=3 + _ji)
                if _pv is not None and _pv == _pv:
                    _ce.value = float(_pv) / value_divisor
                    _ce.number_format = num_fmt
        _row += 1
    if include_mu:
        _row = _write_grafici_mu_source_rows(
            ws_g,
            _row,
            wb=wb,
            offsets=offsets,
            sn7_ref_mu=sn7_ref_mu,
            num_fmt=num_fmt,
            value_divisor=value_divisor,
        )
    return _hdr_r, _data_r, _row


def _grafici_price_series_at_offsets(
    pts: list[dict[str, Any]],
    offsets: tuple[int, ...],
    mode: str,
    pw: dict[str, Any],
) -> list[float | None]:
    """
    Serie $ per grafici Excel da ``points`` (+ fallback modello implicito vs T−60).
    ``path`` → ``price_usd`` per offset; ``model`` → ``price_model_usd`` o da ``pct_modello``.
    """
    _by_off: dict[int, float | None] = {}
    _iso = isinstance(pw, dict)
    _mode = str(mode or "").strip().lower()
    for _p in pts or []:
        try:
            _o = int(_p["offset"])
        except (TypeError, ValueError, KeyError):
            continue
        _v = None
        if _mode == "model":
            _v = _p.get("price_model_usd")
            if (
                (_v is None or _v != _v)
                and _iso
                and isinstance(_p.get("pct_modello"), (int, float))
            ):
                _v = _sim_grafici_price_usd_from_pct_m60(pw, _p.get("pct_modello"))
        else:
            _v = _p.get("price_usd")
        if _v is not None and _v == _v:
            try:
                _by_off[_o] = float(_v)
            except (TypeError, ValueError):
                pass
    _out: list[float | None] = []
    for _off in offsets:
        try:
            _oi = int(_off)
        except (TypeError, ValueError):
            _out.append(None)
            continue
        _x = _by_off.get(_oi)
        _out.append(_x if _x is not None and _x == _x else None)
    return _out


def _excel_sheet_ref(ws) -> str:
    _t = str(getattr(ws, "title", "Simulation") or "Simulation")
    if any(c in _t for c in (" ", "—", "-", "'")):
        return f"'{_t}'!"
    return f"{_t}!"


def _write_ticker_list_column(
    ws_g,
    company_curves: list[dict[str, Any]],
) -> str:
    """Colonna Z: ticker + voce «nessuna» per validazione menu."""
    from openpyxl.utils import get_column_letter as _gcl

    for _i, _ent in enumerate(company_curves):
        ws_g.cell(row=1 + _i, column=_LIST_COL, value=str(_ent["ticker"]))
    _end = len(company_curves) + 1
    ws_g.cell(row=_end, column=_LIST_COL, value=_CO_B_NONE_XL)
    _cl = _gcl(_LIST_COL)
    return f"${_cl}$1:${_cl}${_end}"


def _add_list_validation(ws_g, cell_ref: str, list_formula: str) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    _dv = DataValidation(type="list", formula1=list_formula, allow_blank=False)
    _dv.error = "Scegli un valore dall'elenco"
    _dv.errorTitle = "Selezione non valida"
    ws_g.add_data_validation(_dv)
    _dv.add(cell_ref)


def _add_yesno_validation(ws_g, cell_ref: str) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    _dv = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
    ws_g.add_data_validation(_dv)
    _dv.add(cell_ref)


def _write_grafici_explorer_panel(
    ws_g,
    start_row: int,
    company_curves: list[dict[str, Any]],
    list_formula: str,
    *,
    series_list_formula: str | None = None,
) -> dict[str, str]:
    """Pannello menu (col. S–T): Società A/B, serie grafico Δ% (fino a 15 menu), check curve."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as _gcl

    _row = _write_section_title(
        ws_g,
        start_row,
        "5a",
        "Pannello grafici — Società A/B · spunte coorte (Δ% e Var.%)",
        fill="FFF9C4",
    )
    _vc = _gcl(_PANEL_VAL_COL)
    _lc = _gcl(_PANEL_LAB_COL)
    _specs: list[tuple[str, Any]] = [
        ("Società A", company_curves[0]["ticker"] if company_curves else "—"),
        ("Società B (confronto)", _CO_B_NONE_XL),
        ("Mostra curva ricalibr. (1/0)", 1),
        ("Mostra storico dato vero (1/0)", 1),
        ("Mostra modello (1/0)", 1),
        ("Δ% — tutte le società · curva ricalibr. (1/0)", 1),
        ("Δ% — tutte le società · storico (1/0)", 0),
        ("Δ% — tutte le società · modello (1/0)", 0),
        ("Var.% — tutta la coorte (1/0)", 1),
    ]
    _ser_list = series_list_formula or list_formula
    _refs: dict[str, str] = {}
    for _lab, _def in _specs:
        ws_g.cell(row=_row, column=_PANEL_LAB_COL, value=_lab).font = Font(size=9)
        _c = ws_g.cell(row=_row, column=_PANEL_VAL_COL, value=_def)
        _c.alignment = Alignment(horizontal="left")
        _c.fill = PatternFill("solid", fgColor="FFFDE7")
        _key = (
            "co_a"
            if _lab.startswith("Società A")
            else "co_b"
            if _lab.startswith("Società B")
            else "chk_curva"
            if _lab.startswith("Mostra curva")
            else "chk_storico"
            if _lab.startswith("Mostra storico")
            else "chk_modello"
            if _lab.startswith("Mostra modello")
            else "chk_delta_curva"
            if "· curva" in _lab
            else "chk_delta_storico"
            if "· storico" in _lab
            else "chk_delta_modello"
            if "· modello" in _lab
            else "chk_var_coorte"
            if _lab.startswith("Var.%")
            else "chk_other"
        )
        _refs[_key] = f"${_vc}${_row}"
        if _key in ("co_a", "co_b"):
            _add_list_validation(ws_g, f"{_lc}{_row}:{_vc}{_row}", list_formula)
        else:
            _add_yesno_validation(ws_g, f"{_lc}{_row}:{_vc}{_row}")
        _row += 1
    _refs["list_formula"] = list_formula
    return _refs


def _xl_index_match(
    chk: str,
    co: str,
    wide_data0: int,
    wide_last: int,
    col_idx: int,
    *,
    require_co: bool = True,
) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _cl = _gcl(3 + col_idx)
    _co_test = (
        f'AND({co}<>"",{co}<>"{_CO_B_NONE_XL}")'
        if require_co
        else "TRUE"
    )
    return (
        f"=IF({chk}=1,IF({_co_test},IFERROR(INDEX("
        f"{_cl}${wide_data0}:{_cl}${wide_last},"
        f"MATCH({co},$A${wide_data0}:$A${wide_last},0)),NA()),NA()),NA())"
    )


def _xl_mu_cell(chk_mu: str, mu_row: int, col_idx: int) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _cl = _gcl(3 + col_idx)
    return f"=IF({chk_mu}=1,{_cl}${mu_row},NA())"


def _xl_sim_var_cell(
    co: str,
    sim_ref: str,
    sim_first: int,
    sim_last: int,
    col_sim: int,
) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _cl = _gcl(col_sim)
    _ix = (
        f"IFERROR(INDEX({sim_ref}${_cl}${sim_first}:${_cl}${sim_last},"
        f"MATCH({co},{sim_ref}$A${sim_first}:$A${sim_last},0)),NA())"
    )
    return (
        f'=IF(AND({co}<>"",{co}<>"{_CO_B_NONE_XL}",{co}<>"—"),'
        f"IF(ISNUMBER({_ix}),IF(ABS({_ix})<=1.5,{_ix}*100,{_ix}),{_ix}),NA()),NA())"
    )


def _xl_grafici_var_gate_expr(cohort_row: int) -> str:
    """
    Espressione (senza «=») incluso nel grafico Var.%:
    col. G5 = 1 **oppure** ticker = Società A (C2) / B (C3).
    """
    from openpyxl.utils import get_column_letter as _gcl

    _g5 = f"${_gcl(_GRAF_SEL_COL_VAR)}${cohort_row}"
    _tc = f"${_gcl(_GRAF_SEL_COL_TICKER)}${cohort_row}"
    _ca = f"$C${_GRAF_PANEL_CO_A}"
    _cb = f"$C${_GRAF_PANEL_CO_B}"
    _ab = (
        f"OR({_tc}={_ca},AND({_cb}<>\"\",{_cb}<>\"{_CO_B_NONE_XL}\",{_tc}={_cb}))"
    )
    return f"IF(OR({_g5}=1,{_ab}),1,0)"


def _read_sim_var_pct(
    ws_sim,
    ticker: str,
    col_sim: int,
    sim_first: int,
    sim_last: int,
) -> float | None:
    """Legge Var.% da Simulation (stessa logica di ``_var_pct_for_excel``)."""
    _tk = str(ticker or "").strip().upper()
    if not _tk or _tk in ("—", "-", "N/D"):
        return None
    for _rn in range(int(sim_first), int(sim_last) + 1):
        if str(ws_sim.cell(row=_rn, column=1).value or "").strip().upper() == _tk:
            return _var_pct_for_excel(ws_sim.cell(row=_rn, column=int(col_sim)).value)
    return None


def _xl_grafici_var_row_cell(
    cohort_row: int,
    *,
    pct_value: float | None,
) -> str:
    """Valore Var.% con gate G5/menu; literal così il grafico ha dati senza ricalcolo INDEX."""
    _gate = _xl_grafici_var_gate_expr(cohort_row)
    if pct_value is None or pct_value != pct_value:
        return f"=IF({_gate}<>1,NA(),NA())"
    _lit = str(round(float(pct_value), 6)).replace(",", ".")
    return f"=IF({_gate}<>1,NA(),{_lit})"


def _write_section5_explorer_charts(
    ws_g,
    start_row: int,
    *,
    offsets: tuple[int, ...],
    n_std: int,
    company_curves: list[dict[str, Any]],
    panel: dict[str, str],
    wide: dict[str, Any],
    pct_fmt: str,
    usd_fmt: str,
    ws_sim=None,
    c_var_lo: int = 0,
    c_var_hi: int = 0,
    sim_first_row: int = 4,
    sim_last_row: int = 4,
    combo_reg_first: int = 0,
    combo_reg_last: int = 0,
    combo_reg_lab_col: str = "",
) -> tuple[int, int, int, int, int, int, int, int, list[str]]:
    """
  Pannello + serie dinamiche (menu) + sorgenti grafici.
  Ritorna assi, righe dati, anchor, next_row, etichette serie %.
    """
    from openpyxl.styles import Alignment, Font

    _n_co = len(company_curves)
    _sim_ref = _excel_sheet_ref(ws_sim) if ws_sim is not None else ""
    _row = _write_section_title(
        ws_g,
        start_row,
        "5b",
        "Dati grafici (formule INDEX/MATCH dal pannello 5a e tabelle sez. 2)",
        fill="EDE7F6",
    )
    _row += 1
    _axis_delta_row = _row
    _write_x_session_axis_row(ws_g, _axis_delta_row, offsets)
    _row += 1
    ws_g.cell(row=_row, column=1, value="Serie (grafico %)").font = Font(bold=True, size=9)
    _row += 1
    _chart_pct0 = _row
    _w_curva = wide["curva"]
    _w_stor = wide["storico"]
    _w_mod = wide["modello"]
    _w_px = wide["price_path"]
    _w_pm = wide["price_model"]
    from openpyxl.utils import get_column_letter as _gcl

    _pct_labels: list[str] = []
    _chk_pct_specs = (
        ("chk_delta_curva", "curva ricalibr.", _w_curva),
        ("chk_delta_storico", "storico", _w_stor),
        ("chk_delta_modello", "modello", _w_mod),
    )
    for _chk_k, _suffix, _wt in _chk_pct_specs:
        _chk = panel.get(_chk_k) or "0"
        _d0, _d1 = int(_wt["data0"]), int(_wt["last"])
        for _ent in company_curves:
            _tk = str(_ent["ticker"])
            _cr = _row
            ws_g.cell(row=_cr, column=1, value=f"{_tk} · {_suffix}")
            _pct_labels.append(f"{_tk} · {_suffix}")
            _tk_lit = f'"{_tk}"'
        for _ji in range(n_std):
                _inner = _xl_grafici_cross_pct_raw(_tk_lit, _d0, _d1, _ji)
                _inner = _inner[1:] if _inner.startswith("=") else _inner
                ws_g.cell(row=_cr, column=3 + _ji).value = f"=IF({_chk}=1,{_inner},NA())"
                ws_g.cell(row=_cr, column=3 + _ji).number_format = pct_fmt
    _row += 1
    _row += 1
    ws_g.cell(row=_row, column=1, value="— Var.% 1g · 1M · 3M · 6M —").font = Font(
        bold=True, size=9, italic=True
    )
    _row += 1
    _chk_var = panel.get("chk_var_coorte") or "0"
    _c_var_lo_sg = SIMULATION_36_HEADERS.index("Var. Giorn. %") + 1
    _axis_var_row = _row
    _write_x_var_horizon_axis_row(ws_g, _axis_var_row, x_col0=_VAR_CHART_COL0)
    _row += 1
    _chart_var0 = _row
    _n_var = 0
    for _ent in company_curves:
        _tk = str(_ent["ticker"])
        _cr = _row
        ws_g.cell(row=_cr, column=1, value=_tk)
        for _ji in range(_GRAF_VAR_N):
            _vcol = _c_var_lo_sg + _ji
            _inner = _xl_sim_var_cell(f'"{_tk}"', _sim_ref, sim_first_row, sim_last_row, _vcol)
            _inner = _inner[1:] if _inner.startswith("=") else _inner
            ws_g.cell(row=_cr, column=_VAR_CHART_COL0 + _ji).value = (
                f"=IF({_chk_var}=1,{_inner},NA())"
            )
            ws_g.cell(row=_cr, column=_VAR_CHART_COL0 + _ji).number_format = FMT_VARIATION
        _row += 1
        _n_var += 1
    _var_labels = [f"Var {_i + 1}" for _i in range(_n_var)]
    _row = _chart_var0 + _n_var
    ws_g.cell(row=_row, column=1, value="— Prezzo $ path —").font = Font(
        bold=True, size=9, italic=True
    )
    _row += 1
    _chart_price_path0 = _row
    _px_labels: list[str] = []
    for _pfx, _co_k in (("A", "co_a"), ("B", "co_b")):
        _lb = f"{_pfx} · prezzo path"
        _px_labels.append(_lb)
        _cr = _row
        ws_g.cell(row=_cr, column=1, value=_lb)
        _d0, _d1 = int(_w_px["data0"]), int(_w_px["last"])
        for _ji in range(n_std):
            ws_g.cell(row=_cr, column=3 + _ji).value = _xl_index_match(
                panel["chk_curva"], panel[_co_k], _d0, _d1, _ji
            )
            ws_g.cell(row=_cr, column=3 + _ji).number_format = usd_fmt
    _row += 1
    _row += 1
    ws_g.cell(row=_row, column=1, value="— Prezzo $ modello T−60 —").font = Font(
        bold=True, size=9, italic=True
    )
    _row += 1
    _chart_price_model0 = _row
    _pm_labels: list[str] = []
    for _pfx, _co_k in (("A", "co_a"), ("B", "co_b")):
        _lb = f"{_pfx} · prezzo modello"
        _pm_labels.append(_lb)
        _cr = _row
        ws_g.cell(row=_cr, column=1, value=_lb)
        _d0, _d1 = int(_w_pm["data0"]), int(_w_pm["last"])
        for _ji in range(n_std):
            ws_g.cell(row=_cr, column=3 + _ji).value = _xl_index_match(
                panel["chk_modello"], panel[_co_k], _d0, _d1, _ji
            )
            ws_g.cell(row=_cr, column=3 + _ji).number_format = usd_fmt
        _row += 1
    _chart_anchor = _row + 1
    return (
        _axis_delta_row,
        _chart_pct0,
        len(_pct_labels),
        _axis_var_row,
        _chart_var0,
        len(_var_labels),
        _chart_price_path0,
        len(_px_labels),
        _chart_price_model0,
        len(_pm_labels),
        _chart_anchor,
        _row + 1,
        _pct_labels,
    )


_GRAF_PANEL_CO_A = 2
_GRAF_PANEL_CO_B = 3
_GRAF_PANEL_ALL = 4
_GRAF_COMBO_SLOT_COUNT = 15  # cap coorte Simulation
_GRAF_SEL_TITLE = 5
_GRAF_SEL_HDR = 6
_GRAF_SEL_DATA0 = 7
_GRAF_SEL_COL_TICKER = 2
_GRAF_SEL_COL_G1 = 3
_GRAF_SEL_COL_G2 = 4
_GRAF_SEL_COL_G3 = 5
_GRAF_SEL_COL_G4 = 6
_GRAF_SEL_COL_G4_CURVA = 7
_GRAF_SEL_COL_G4_STOR = 8
_GRAF_SEL_COL_G4_MOD = 9
_GRAF_SEL_COL_VAR = 10
_GRAF_SEL_COL_COMBO_LBL = 11
_GRAF_SEL_COL_G4_MU = 12
_GRAF_SEL_COL_REF_INC = 3
_GRAF_COMBO_CTRL_SLOT_COUNT = 6
_G4_COMBO_TIPO_SPECS: tuple[tuple[str, int, str], ...] = (
    ("curva ricalibr.", _GRAF_SEL_COL_G4_CURVA, "curva"),
    ("storico", _GRAF_SEL_COL_G4_STOR, "storico"),
    ("modello", _GRAF_SEL_COL_G4_MOD, "modello"),
)
_GRAFICI_EXCEL_CHART_STYLE = 12
_GRAF_LIST_COL_G = 26
_GRAF_COMBO_LIST_COL = 27
_GRAF_COMBO_REG_COL_LAB = 13
_GRAF_COMBO_REG_COL0 = 14
_GRAF_COMBO_NONE_XL = "(nessuna)"
_GRAF_COMBO_REG_START_ROW = 400
_GRAF_COMBO_WIDE_GAP_ROWS = 12


def _grafici_combo_registry_row_count(
    company_curves: list[dict[str, Any]],
    *,
    pred_linked: bool,
    pred_series_len: int,
) -> int:
    """Righe occupate dal registro menu combinato (sotto ``_GRAF_COMBO_REG_START_ROW``)."""
    _n_co = len(company_curves or [])
    _n_mu = pred_series_len if pred_linked else len(_REF_SPECS)
    return 2 + _n_co * 3 + max(int(_n_mu), 0) + 2


def _grafici_wide_table_start_row(
    company_curves: list[dict[str, Any]],
    *,
    pred_linked: bool,
    pred_series_len: int,
) -> int:
    """Prima riga tabelle 2a–2c: sotto il registro combo, mai sovrapposta."""
    _need = _grafici_combo_registry_row_count(
        company_curves,
        pred_linked=pred_linked,
        pred_series_len=pred_series_len,
    )
    return max(
        _GRAF_WIDE_TABLE_START_ROW,
        _GRAF_COMBO_REG_START_ROW + _need + _GRAF_COMBO_WIDE_GAP_ROWS,
    )


def _grafici_unmerge_at(ws, row: int, column: int) -> None:
    """Rimuove merge che coprono (row, column) se non è l'angolo in alto a sinistra."""
    for _mr in list(ws.merged_cells.ranges):
        if (
            _mr.min_row <= row <= _mr.max_row
            and _mr.min_col <= column <= _mr.max_col
            and not (row == _mr.min_row and column == _mr.min_col)
        ):
            ws.unmerge_cells(str(_mr))


def _grafici_set_cell(ws, row: int, column: int, value=None, **cell_kw):
    """Scrittura sicura: evita ``MergedCell`` read-only su celle non angolo."""
    _grafici_unmerge_at(ws, row, column)
    _c = ws.cell(row=row, column=column)
    if value is not None:
        _c.value = value
    for _k, _v in cell_kw.items():
        setattr(_c, _k, _v)
    return _c

_GRAF_VAR_MENU_COL = 2  # B — menu a tendina società (uno per slot)
_GRAF_VAR_SLOT_COUNT = 12
_GRAF_VAR_COL0 = 5  # E–H: valori Var.% 1g · 1M · 3M · 6M
_GRAF_VAR_N = 4
_GRAF_VAR_AXIS_LABELS = ("1g", "1M", "3M", "6M")
_GRAF_VAR_SERIES_LABELS = (
    "var 1day",
    "var 1 months",
    "var 3 months",
    "var 6 months",
)
_GRAF_VAR_X_TITLE = "Società"
# Asse X grafici %: sessioni da CD (stessa griglia Predizione / SN7).
_GRAFIC_SESSION_X_MIN = -62
_GRAFIC_SESSION_X_MAX = 9
_GRAFIC_SESSION_X_TITLE = "Sessioni vs CD (−60 … +7)"
# Mesh 5×4: 5 linee verticali, 4 orizzontali (griglia larga come riferimento Excel).
_GRAFIC_MESH_X_LINES = 5
_GRAFIC_MESH_Y_LINES = 4
_GRAFIC_SESSION_Y_MIN = -2.0  # −200% (formato % Excel)
_GRAFIC_SESSION_Y_MAX = 2.0  # +200%
_GRAFIC_VAR_X_MIN = -0.15
_GRAFIC_VAR_X_MAX = 3.15
_GRAFIC_VAR_Y_MIN = -200.0
_GRAFIC_VAR_Y_MAX = 200.0
_GRAFIC_USD_Y_MIN = 0.0
_GRAFIC_USD_Y_MAX = 300.0


def _mesh_major_unit(span: float, n_lines: int) -> float:
    return float(span) / max(1, int(n_lines) - 1)


def _pct_number_for_axis(v: Any) -> float | None:
    """Valore in scala % per calcolo asse Y (accetta frazione o %)."""
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f != f:
        return None
    if abs(f) <= 1.5:
        return f * 100.0
    return f


def _pct_series_groups_from_company_curves(
    company_curves: list[dict[str, Any]],
    *fields: str,
) -> list[list[float]]:
    """Un gruppo per società: tutti i punti standard dei campi richiesti."""
    _groups: list[list[float]] = []
    for _ent in company_curves:
        _vals: list[float] = []
        for _p in _ent.get("points") or []:
            if str(_p.get("nodo") or "") != "standard":
                continue
            for _fk in fields:
                _nv = _pct_number_for_axis(_p.get(_fk))
                if _nv is not None:
                    _vals.append(_nv)
        if _vals:
            _groups.append(_vals)
    return _groups


def _mu_pct_series_groups(
    offsets: tuple[int, ...],
    sn7_ref_mu: dict | None,
) -> list[list[float]]:
    _groups: list[list[float]] = []
    if not isinstance(sn7_ref_mu, dict):
        return _groups
    for _mlab, _ in _REF_SPECS:
        _raw = _grafici_mu_values_for_offsets(sn7_ref_mu, _mlab, offsets)
        _vals = []
        for _v in _raw:
            _nv = _pct_number_for_axis(_v)
            if _nv is not None:
                _vals.append(_nv)
        if _vals:
            _groups.append(_vals)
    return _groups


def _var_pct_series_groups(
    company_curves: list[dict[str, Any]],
    ws_sim,
    c_var_lo: int,
) -> list[list[float]]:
    if ws_sim is None:
        return []
    _rows: dict[str, int] = {}
    for _rn in range(4, int(ws_sim.max_row or 0) + 1):
        _tk = str(ws_sim.cell(row=_rn, column=1).value or "").strip().upper()
        if _tk and _tk not in ("—", "-", "N/D"):
            _rows[_tk] = _rn
    _groups: list[list[float]] = []
    for _ent in company_curves:
        _tk = str(_ent.get("ticker") or "").strip().upper()
        _rn = _rows.get(_tk)
        if not _rn:
            continue
        _vals: list[float] = []
        for _ji in range(_GRAF_VAR_N):
            _nv = _pct_number_for_axis(
                ws_sim.cell(row=_rn, column=int(c_var_lo) + _ji).value
            )
            if _nv is not None:
                _vals.append(_nv)
        if _vals:
            _groups.append(_vals)
    return _groups


def _nice_axis_step(span: float) -> float:
    import math

    if span <= 0:
        return 10.0
    _raw = span / max(1, _GRAFIC_MESH_Y_LINES - 1)
    if _raw <= 0:
        return 10.0
    _mag = 10.0 ** math.floor(math.log10(_raw))
    for _m in (1.0, 2.0, 5.0, 10.0):
        _s = _m * _mag
        if _s >= _raw * 0.85:
            return _s
    return _mag * 10.0


def _y_axis_from_all_visible_series(
    groups: list[list[float]],
    *,
    as_excel_fraction: bool = False,
    anchor_zero: bool = False,
) -> tuple[float, float, float]:
    """
    Asse Y da min/max globali su tutte le serie del grafico (+10% sugli estremi).
    Con ``anchor_zero=True`` (grafico Var.% a barre) l'intervallo include sempre 0
    così le barre negative scendono sotto la baseline.
    """
    _all: list[float] = []
    for _vals in groups:
        _all.extend(_vals)
    if not _all:
        _ymin, _ymax, _maj = -200.0, 200.0, 100.0
    else:
        _lo = min(_all)
        _hi = max(_all)
        _ymin = _lo - 0.1 * abs(_lo)
        _ymax = _hi + 0.1 * abs(_hi)
        if anchor_zero:
            _ymin = min(_ymin, 0.0)
            _ymax = max(_ymax, 0.0)
        if _ymin >= _ymax:
            _pad = max(1.0, 0.1 * max(abs(_ymin), abs(_ymax), 1.0))
            _ymin -= _pad
            _ymax += _pad
            if anchor_zero:
                _ymin = min(_ymin, 0.0)
                _ymax = max(_ymax, 0.0)
        _maj = _nice_axis_step(_ymax - _ymin)
    if as_excel_fraction:
        return _ymin / 100.0, _ymax / 100.0, _maj / 100.0
    return _ymin, _ymax, _maj


_GRAFIC_SESSION_X_MAJOR = _mesh_major_unit(
    _GRAFIC_SESSION_X_MAX - _GRAFIC_SESSION_X_MIN, _GRAFIC_MESH_X_LINES
)
_GRAFIC_SESSION_Y_MAJOR = 1.0  # passo 100% (scala frazione −2…+2)
_GRAFIC_VAR_X_MAJOR = _mesh_major_unit(
    _GRAFIC_VAR_X_MAX - _GRAFIC_VAR_X_MIN, _GRAFIC_MESH_X_LINES
)
_GRAFIC_VAR_Y_MAJOR = 100.0  # passo 100% (−200…+200)
_GRAFIC_USD_Y_MAJOR = _mesh_major_unit(
    _GRAFIC_USD_Y_MAX - _GRAFIC_USD_Y_MIN, _GRAFIC_MESH_Y_LINES
)


def _chart_series_title_ref(ws, row: int, col: int = 1):
    """Riferimento cella etichetta serie (Excel valuta la formula in legenda)."""
    from openpyxl.chart import Reference

    return Reference(ws, min_col=col, min_row=row, max_col=col, max_row=row)


def _apply_chart_series_title(ser, ws, row: int, col: int = 1) -> None:
    from openpyxl.chart.data_source import StrRef
    from openpyxl.chart.series import SeriesLabel

    ser.title = SeriesLabel()
    ser.title.strRef = StrRef()
    ser.title.strRef.f = _chart_series_title_ref(ws, row, col)


def _apply_grafici_chart_grid(ch) -> None:
    """Griglia major ampia (no minor) su X e Y."""
    from openpyxl.chart.axis import ChartLines

    for _ax_name in ("x_axis", "y_axis"):
        _ax = getattr(ch, _ax_name, None)
        if _ax is None:
            continue
        try:
            _ax.majorGridlines = ChartLines()
            _ax.minorGridlines = None
        except Exception:
            pass


def _apply_grafici_axis_tick_marks(ch) -> None:
    """Segni di graduazione (tick) visibili su assi X e Y."""
    for _ax_name in ("x_axis", "y_axis"):
        _ax = getattr(ch, _ax_name, None)
        if _ax is None:
            continue
        try:
            _ax.majorTickMark = "out"
            _ax.minorTickMark = None
        except Exception:
            pass


def _apply_grafici_mesh_axes(
    ch,
    *,
    x_min: float,
    x_max: float,
    y_min: float | None = None,
    y_max: float | None = None,
    x_major: float | None = None,
    y_major: float | None = None,
    x_num_fmt: str = "0",
    y_num_fmt: str | None = None,
) -> None:
    """Griglia X fissa; Y auto (``y_min``/``y_max`` None) o manuale se entrambi impostati."""
    try:
        ch.x_axis.scaling.min = float(x_min)
        ch.x_axis.scaling.max = float(x_max)
        ch.x_axis.majorUnit = float(
            x_major
            if x_major is not None
            else _mesh_major_unit(x_max - x_min, _GRAFIC_MESH_X_LINES)
        )
        ch.x_axis.minorUnit = None
        ch.x_axis.number_format = x_num_fmt
        if y_min is not None and y_max is not None:
            ch.y_axis.scaling.min = float(y_min)
            ch.y_axis.scaling.max = float(y_max)
            ch.y_axis.majorUnit = float(
                y_major
                if y_major is not None
                else _mesh_major_unit(float(y_max) - float(y_min), _GRAFIC_MESH_Y_LINES)
            )
            ch.y_axis.minorUnit = None
        if y_num_fmt:
            ch.y_axis.numFmt = y_num_fmt
    except Exception:
        pass


def _style_grafici_excel_chart(ch) -> None:
    """Stile come «Layout 12»: legenda a destra, linee levigate, area grafico ampia."""
    try:
        ch.style = _GRAFICI_EXCEL_CHART_STYLE
        ch.legend.position = "r"
        ch.legend.overlay = False
    except Exception:
        pass
    _apply_grafici_chart_grid(ch)
    _apply_grafici_axis_tick_marks(ch)
    for _ser in getattr(ch, "series", ()) or ():
        try:
            _ser.smooth = True
        except Exception:
            pass


def _default_ticker_slots(
    company_curves: list[dict[str, Any]],
    *,
    n: int = _GRAF_COMBO_SLOT_COUNT,
) -> list[str]:
    _out = [_GRAF_COMBO_NONE_XL] * n
    for _i, _ent in enumerate(company_curves[:n]):
        _out[_i] = str(_ent["ticker"])
    return _out


def _default_slot_labels(
    company_curves: list[dict[str, Any]],
    suffix: str,
    *,
    n: int = _GRAF_COMBO_SLOT_COUNT,
) -> list[str]:
    _out = [_GRAF_COMBO_NONE_XL] * n
    for _i, _ent in enumerate(company_curves[:n]):
        _out[_i] = f"{_ent['ticker']} · {suffix}"
    return _out


def _write_grafici_cohort_selector(
    ws_p,
    company_curves: list[dict[str, Any]],
    pred_scan: list[tuple[str, int]],
    *,
    pred_linked: bool = False,
    combo_list_formula: str = "",
) -> dict[str, Any]:
    """Tabella coorte: spunte G1–G5, μ (G1–G3 / G4), menu curve controllo sul combinato."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as _gcl

    ws_p.merge_cells("A5:L5")
    _t = ws_p.cell(
        row=_GRAF_SEL_TITLE,
        column=1,
        value=(
            "Selezione coorte — spunta le colonne per includere ogni società nei grafici "
            "(G1 curva · G2 storico · G3 modello · G4 combinato · G5 Var.%). "
            "G4 comb. = società nel combinato; colonne curva/stor./mod. = quali curve "
            "(se tutte 0 ma G4=1 → mostra tutte e tre). "
            "Sotto: curve μ (G1–G3 e/o G4) e fino a 6 serie aggiuntive sul combinato (menu)."
        ),
    )
    _t.font = Font(bold=True, size=9, color="1F3864")
    _t.alignment = Alignment(wrap_text=True)
    ws_p.row_dimensions[_GRAF_SEL_TITLE].height = 32

    _hdr = _GRAF_SEL_HDR
    _heads = (
        (1, "#"),
        (_GRAF_SEL_COL_TICKER, "Ticker"),
        (_GRAF_SEL_COL_G1, "G1 % curva"),
        (_GRAF_SEL_COL_G2, "G2 % stor."),
        (_GRAF_SEL_COL_G3, "G3 % mod."),
        (_GRAF_SEL_COL_G4, "G4 comb."),
        (_GRAF_SEL_COL_G4_CURVA, "G4 curva"),
        (_GRAF_SEL_COL_G4_STOR, "G4 stor."),
        (_GRAF_SEL_COL_G4_MOD, "G4 mod."),
        (_GRAF_SEL_COL_VAR, "G5 Var.%"),
    )
    for _col, _lab in _heads:
        ws_p.cell(row=_hdr, column=_col, value=_lab).font = Font(bold=True, size=9)
    _fill = PatternFill("solid", fgColor="E8F5E9")
    _data0 = _GRAF_SEL_DATA0
    for _i, _ent in enumerate(company_curves):
        _r = _data0 + _i
        ws_p.cell(row=_r, column=1, value=_i + 1)
        ws_p.cell(row=_r, column=_GRAF_SEL_COL_TICKER, value=str(_ent["ticker"]))
        for _c in (
            _GRAF_SEL_COL_G1,
            _GRAF_SEL_COL_G2,
            _GRAF_SEL_COL_G3,
        ):
            _g_on = 1 if _i < _GRAF_DEFAULT_G123_ON_N else 0
            _cv = ws_p.cell(row=_r, column=_c, value=_g_on)
            _cv.fill = _fill
            _add_yesno_validation(ws_p, f"{_gcl(_c)}{_r}")
        _g5v = ws_p.cell(row=_r, column=_GRAF_SEL_COL_VAR, value=0)
        _g5v.fill = _fill
        _add_yesno_validation(ws_p, f"{_gcl(_GRAF_SEL_COL_VAR)}{_r}")
        _g4 = ws_p.cell(row=_r, column=_GRAF_SEL_COL_G4, value=0)
        _g4.fill = _fill
        _add_yesno_validation(ws_p, f"{_gcl(_GRAF_SEL_COL_G4)}{_r}")
        for _tc in (
            _GRAF_SEL_COL_G4_CURVA,
            _GRAF_SEL_COL_G4_STOR,
            _GRAF_SEL_COL_G4_MOD,
        ):
            _gt = ws_p.cell(row=_r, column=_tc, value=1)
            _gt.fill = _fill
            _add_yesno_validation(ws_p, f"{_gcl(_tc)}{_r}")

    _sel_last = _data0 + len(company_curves) - 1 if company_curves else _data0 - 1
    _row = _sel_last + 2
    _ref_data0 = 0
    _ref_last = 0
    _mu_entries: list[tuple[str, int | None]] = (
        list(pred_scan) if pred_scan else [(_mlab, None) for _mlab, _ in _REF_SPECS]
    )
    if _mu_entries:
        _mu_title = (
            "Curve μ (Predizione — guida)"
            if pred_linked
            else "Curve μ di riferimento (cache SN7)"
        )
        ws_p.cell(row=_row, column=1, value=_mu_title).font = Font(
            bold=True, size=9, color="1F3864"
        )
        _row += 1
        ws_p.cell(row=_row, column=1, value="Curva").font = Font(bold=True, size=9)
        ws_p.cell(row=_row, column=_GRAF_SEL_COL_REF_INC, value="G1–G3").font = Font(
            bold=True, size=9
        )
        ws_p.cell(row=_row, column=_GRAF_SEL_COL_G4_MU, value="G4 comb.").font = Font(
            bold=True, size=9
        )
        _row += 1
        _ref_data0 = _row
        _mu_on = 1 if pred_linked else 0
        for _mlab, _pr in _mu_entries:
            ws_p.cell(row=_row, column=1, value=_mlab)
            for _gc in (_GRAF_SEL_COL_REF_INC, _GRAF_SEL_COL_G4_MU):
                _cv = ws_p.cell(row=_row, column=_gc, value=_mu_on)
                _cv.fill = _fill
                _add_yesno_validation(ws_p, f"{_gcl(_gc)}{_row}")
            _row += 1
        _ref_last = _row - 1
        _row += 1

    _ctrl_slot_cells: list[str] = []
    if combo_list_formula:
        ws_p.cell(row=_row, column=1, value="G4 — curve aggiuntive (menu)").font = Font(
            bold=True, size=9, color="1F3864"
        )
        _row += 1
        ws_p.cell(row=_row, column=1, value="Slot").font = Font(bold=True, size=9)
        ws_p.cell(row=_row, column=_GRAF_SEL_COL_COMBO_LBL, value="Serie (μ o altro)").font = Font(
            bold=True, size=9
        )
        _row += 1
        _slot_col = _gcl(_GRAF_SEL_COL_COMBO_LBL)
        for _si in range(1, _GRAF_COMBO_CTRL_SLOT_COUNT + 1):
            ws_p.cell(row=_row, column=1, value=f"Controllo {_si}")
            _sc = ws_p.cell(row=_row, column=_GRAF_SEL_COL_COMBO_LBL, value=_GRAF_COMBO_NONE_XL)
            _sc.fill = PatternFill("solid", fgColor="E3F2FD")
            _add_list_validation(ws_p, f"{_slot_col}{_row}", combo_list_formula)
            _ctrl_slot_cells.append(f"${_slot_col}${_row}")
            _row += 1
        _row += 1

    return {
        "sel_data0": _data0,
        "sel_last": _sel_last,
        "ref_data0": _ref_data0,
        "ref_last": _ref_last,
        "pred_scan": pred_scan,
        "ctrl_slot_cells": tuple(_ctrl_slot_cells),
        "next_row": _row,
    }


def _xl_grafici_sel_row_pct(
    inc_cell: str,
    tick_cell: str,
    wide_data0: int,
    wide_last: int,
    col_idx: int,
) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _cl = _gcl(3 + col_idx)
    _src = _grafici_src_bang()
    return (
        f"=IF({inc_cell}<>1,NA(),IFERROR(INDEX("
        f"{_src}{_cl}${wide_data0}:{_src}{_cl}${wide_last},"
        f"MATCH({tick_cell},{_src}$A${wide_data0}:{_src}$A${wide_last},0)),NA()))"
    )


def _append_mu_rows_to_pct_chart(
    ws_p,
    row: int,
    n: int,
    *,
    gate_col: int,
    sel: dict[str, Any],
    offsets: tuple[int, ...],
    pct_fmt: str,
    wb,
    wide: dict[str, Any] | None,
    pred_scan: list[tuple[str, int]] | None = None,
) -> tuple[int, int]:
    """Aggiunge righe μ al blocco dati grafico (gate su colonna coorte)."""
    from openpyxl.utils import get_column_letter as _gcl

    _rd0 = int(sel.get("ref_data0") or 0)
    if not _rd0:
        return row, n
    _n_std = len(offsets)
    _rc = _gcl(gate_col)
    _scan = pred_scan if pred_scan is not None else (sel.get("pred_scan") or [])
    _mu_rows = dict((wide or {}).get("mu_rows") or {})
    _mu_entries: list[tuple[str, int | None]] = (
        list(_scan) if _scan else [(_mlab, None) for _mlab, _ in _REF_SPECS]
    )
    _pred_rows = dict((wide or {}).get("pred_mu_rows") or {})
    for _j, (_mlab, _pr) in enumerate(_mu_entries):
        _rr = _rd0 + _j
        _inc = f"${_rc}${_rr}"
        _cr = row
        ws_p.cell(row=_cr, column=1).value = f'=IF({_inc}<>1,"",A{_rr})'
        _rn = int(_pr or _pred_rows.get(_mlab) or 0)
        for _ji in range(_n_std):
            _ce = ws_p.cell(row=_cr, column=3 + _ji)
            if _rn and wb is not None and _predizione_guida_sheet_ref(wb):
                _mu = _xl_grafici_mu_from_predizione_row(wb, _rn, _ji, offsets)
                _inner = _mu[1:] if str(_mu).startswith("=") else str(_mu)
                _ce.value = f"=IF({_inc}<>1,NA(),{_inner})"
            else:
                _mr = int(_mu_rows.get(_mlab) or 0)
                if _mr:
                    _cl = _gcl(3 + _ji)
                    _ce.value = (
                        f"=IF({_inc}<>1,NA(),IFERROR("
                        f"{_grafici_src_bang()}{_cl}${_mr},NA()))"
                    )
                else:
                    _ce.value = f"=IF({_inc}<>1,NA(),NA())"
            _ce.number_format = pct_fmt
        row += 1
        n += 1
    return row, n


def _append_combo_ctrl_slots_to_pct_chart(
    ws_p,
    row: int,
    n: int,
    *,
    slot_cells: tuple[str, ...],
    offsets: tuple[int, ...],
    pct_fmt: str,
    combo_reg_first: int,
    combo_reg_last: int,
    combo_reg_lab_col: str,
) -> tuple[int, int]:
    """Serie scelte nei menu «Controllo 1…6» sul grafico combinato."""
    from openpyxl.utils import get_column_letter as _gcl

    _n_std = len(offsets)
    for _slot_cell in slot_cells or ():
        _cr = row
        ws_p.cell(row=_cr, column=1).value = (
            f'=IF(OR({_slot_cell}="",{_slot_cell}="{_GRAF_COMBO_NONE_XL}"),"",'
            f"{_slot_cell})"
        )
        for _ji in range(_n_std):
            _val_col = _gcl(_GRAF_COMBO_REG_COL0 + _ji)
            _ce = ws_p.cell(row=_cr, column=3 + _ji)
            _ce.value = _xl_grafici_combo_slot_cell(
                _slot_cell,
                combo_reg_lab_col,
                combo_reg_first,
                combo_reg_last,
                _val_col,
                _ji,
            )
            _ce.number_format = pct_fmt
        row += 1
        n += 1
    return row, n


def _write_grafici_pct_chart_from_selector(
    ws_p,
    start_row: int,
    *,
    block_title: str,
    include_col: int,
    wide_data0: int,
    wide_last: int,
    sel: dict[str, Any],
    company_curves: list[dict[str, Any]],
    offsets: tuple[int, ...],
    pct_fmt: str,
    wb=None,
    combo_reg_first: int = 0,
    combo_reg_last: int = 0,
    combo_reg_lab_col: str = "",
    combo_reg_val_col: str = "",
    use_combo_registry: bool = False,
    wide: dict[str, Any] | None = None,
    curve_field: str | None = None,
) -> tuple[int, int, int]:
    """Blocchi G1–G3 (letterali da ``company_curves``) o G4 (combinato) + μ opzionali."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as _gcl

    _n_std = len(offsets)
    _row = start_row
    ws_p.cell(row=_row, column=1, value=block_title).font = Font(
        bold=True, size=10, color="1F3864"
    )
    _row += 1
    _axis = _row
    _write_x_session_axis_row(ws_p, _axis, offsets)
    _row += 1
    _data0 = _row
    _n = 0
    _ic = _gcl(include_col)
    _tc = _gcl(_GRAF_SEL_COL_TICKER)
    if use_combo_registry and isinstance(wide, dict):
        for _i, _ent in enumerate(company_curves):
            _sr = int(sel["sel_data0"]) + _i
            _tick = f"${_tc}${_sr}"
            for _slab, _tcol, _wk in _G4_COMBO_TIPO_SPECS:
                _cond = _xl_grafici_g4_tipo_cond(_sr, _tcol)
                _fld = _G4_FIELD_BY_WIDE_KEY[_wk]
                _cr = _row
                ws_p.cell(row=_cr, column=1).value = (
                    f'=IF(NOT({_cond}),NA(),{_tick}&" · {_slab}")'
                )
                for _ji, _off in enumerate(offsets):
                    _ce = ws_p.cell(row=_cr, column=3 + _ji)
                    _frac = _pct_fraction_from_curve(_ent, _fld, int(_off))
                    _ce.value = _xl_grafici_g4_combo_pct_cell(_cond, _frac)
                    _ce.number_format = pct_fmt
                _row += 1
                _n += 1
        _row, _n = _append_mu_rows_to_pct_chart(
            ws_p,
            _row,
            _n,
            gate_col=_GRAF_SEL_COL_G4_MU,
            sel=sel,
            offsets=offsets,
            pct_fmt=pct_fmt,
            wb=wb,
            wide=wide,
        )
        _row, _n = _append_combo_ctrl_slots_to_pct_chart(
            ws_p,
            _row,
            _n,
            slot_cells=tuple(sel.get("ctrl_slot_cells") or ()),
            offsets=offsets,
            pct_fmt=pct_fmt,
            combo_reg_first=combo_reg_first,
            combo_reg_last=combo_reg_last,
            combo_reg_lab_col=combo_reg_lab_col,
        )
    else:
        _curve_field = curve_field or _GRAF_CHART_FIELD_BY_INCLUDE_COL.get(
            include_col, "pct_curva"
        )
        for _i, _ent in enumerate(company_curves):
            _sr = int(sel["sel_data0"]) + _i
            _inc = f"${_ic}${_sr}"
            _tick = f"${_tc}${_sr}"
            _cr = _row
            ws_p.cell(row=_cr, column=1).value = f'=IF({_inc}<>1,"",{_tick})'
            _cond_on = f"{_inc}=1"
            for _ji, _off in enumerate(offsets):
                _ce = ws_p.cell(row=_cr, column=3 + _ji)
                _frac = _pct_fraction_from_curve(_ent, _curve_field, int(_off))
                _ce.value = _xl_grafici_g4_combo_pct_cell(_cond_on, _frac)
                _ce.number_format = pct_fmt
            _row += 1
            _n += 1

    if not use_combo_registry:
        _row, _n = _append_mu_rows_to_pct_chart(
            ws_p,
            _row,
            _n,
            gate_col=_GRAF_SEL_COL_REF_INC,
            sel=sel,
            offsets=offsets,
            pct_fmt=pct_fmt,
            wb=wb,
            wide=wide,
        )

    return _axis, _data0, _n


def _write_x_session_axis_row(ws, row: int, offsets: tuple[int, ...]) -> None:
    """Riga categorie X numeriche (sessioni) per scatter con marker."""
    from openpyxl.styles import Alignment

    for _ji, _off in enumerate(offsets):
        _c = ws.cell(row=row, column=3 + _ji, value=int(_off))
        _c.number_format = "0"
        _c.alignment = Alignment(horizontal="center", vertical="center")


def _add_session_scatter_chart(
    ws,
    anchor: str,
    title: str,
    cat_row: int,
    data0: int,
    n_series: int,
    n_x: int,
    *,
    x_col0: int = 3,
    pct_fmt: str = "+0.0%;-0.0%;0.0%",
    y_title: str = "Δ% vs T−60",
    x_title: str | None = None,
    y_min: float | None = None,
    y_max: float | None = None,
    y_major: float | None = None,
) -> None:
    """Grafico a linee con marker; asse X = sessioni −60…+7 (valori numerici)."""
    from openpyxl.chart import Reference, ScatterChart
    from openpyxl.chart.marker import Marker
    from openpyxl.chart.series_factory import SeriesFactory

    if n_series < 1 or n_x < 1:
        return
    _x_title = x_title if x_title is not None else _GRAFIC_SESSION_X_TITLE
    _x_col1 = x_col0 + n_x - 1
    ch = ScatterChart()
    ch.title = title
    ch.scatterStyle = "lineMarker"
    ch.y_axis.title = y_title
    ch.x_axis.title = _x_title
    ch.height = 12
    ch.width = 22
    _style_grafici_excel_chart(ch)
    _xvalues = Reference(
        ws, min_col=x_col0, min_row=cat_row, max_col=_x_col1, max_row=cat_row
    )
    for _si in range(n_series):
        _cr = data0 + _si
        _yvalues = Reference(
            ws, min_col=x_col0, min_row=_cr, max_col=_x_col1, max_row=_cr
        )
        _ser = SeriesFactory(_yvalues, _xvalues)
        _apply_chart_series_title(_ser, ws, _cr)
        _ser.marker = Marker("circle", size=7)
        try:
            _ser.graphicalProperties.line.width = 20000
        except Exception:
            pass
        ch.series.append(_ser)
    try:
        ch.display_blanks = "gap"
    except Exception:
        pass
    try:
        ch.y_axis.numFmt = pct_fmt
        _apply_grafici_mesh_axes(
            ch,
            x_min=_GRAFIC_SESSION_X_MIN,
            x_max=_GRAFIC_SESSION_X_MAX,
            y_min=y_min,
            y_max=y_max,
            x_major=_GRAFIC_SESSION_X_MAJOR,
            y_major=y_major,
            y_num_fmt=pct_fmt,
        )
    except Exception:
        pass
    ws.add_chart(ch, anchor)


def _write_x_var_horizon_axis_row(ws, row: int, *, x_col0: int = _GRAF_VAR_COL0) -> None:
    """Intestazioni colonne: var 1day · 1 months · 3 months · 6 months."""
    from openpyxl.styles import Alignment, Font

    for _ji, _lbl in enumerate(_GRAF_VAR_SERIES_LABELS):
        _c = ws.cell(row=row, column=x_col0 + _ji, value=_lbl)
        _c.font = Font(size=9, bold=True)
        _c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_variation_bar_chart(
    ws,
    anchor: str,
    title: str,
    header_row: int,
    data0: int,
    n_companies: int,
    *,
    x_col0: int = _GRAF_VAR_COL0,
    y_min: float | None = None,
    y_max: float | None = None,
    y_major: float | None = None,
) -> None:
    """Var.%: barre raggruppate per società (4 barre = 1g, 1M, 3M, 6M)."""
    from openpyxl.chart import BarChart, Reference

    if n_companies < 1:
        return
    _x_col1 = x_col0 + _GRAF_VAR_N - 1
    _data_last = data0 + n_companies - 1
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.overlap = 0
    ch.gapWidth = 80
    try:
        ch.display_blanks = "gap"
    except Exception:
        pass
    ch.title = title
    ch.y_axis.title = "Var. %"
    ch.x_axis.title = _GRAF_VAR_X_TITLE
    ch.height = 12
    ch.width = 22
    _style_grafici_excel_chart(ch)
    _data = Reference(
        ws,
        min_col=x_col0,
        min_row=header_row,
        max_col=_x_col1,
        max_row=_data_last,
    )
    ch.add_data(_data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=data0, max_row=_data_last))
    try:
        ch.y_axis.numFmt = FMT_VARIATION
        if y_min is not None and y_max is not None:
            ch.y_axis.scaling.min = float(y_min)
            ch.y_axis.scaling.max = float(y_max)
            ch.y_axis.majorUnit = (
                float(y_major) if y_major is not None else float(_GRAFIC_VAR_Y_MAJOR)
            )
            ch.y_axis.minorUnit = None
        ch.y_axis.crosses = "autoZero"
    except Exception:
        pass
    for _ser in ch.series:
        try:
            _ser.invertIfNegative = False
        except Exception:
            pass
    ws.add_chart(ch, anchor)


def _add_variation_scatter_chart(
    ws,
    anchor: str,
    title: str,
    header_row: int,
    data0: int,
    n_series: int,
    *,
    x_col0: int = _GRAF_VAR_COL0,
) -> None:
    """Alias: grafico Var.% a barre raggruppate."""
    _add_variation_bar_chart(
        ws,
        anchor,
        title,
        header_row,
        data0,
        n_series,
        x_col0=x_col0,
    )


_GRAF_VAR_CHART_COL = 10  # J: grafico subito a destra del blocco dati (righe ~23–26)
_GRAF_VAR_LEGACY_COL0 = 20  # T–W (vecchia posizione — va svuotata)


def _simulation_ws(wb):
    if wb is None:
        return None
    for _sn in wb.sheetnames:
        if str(_sn).strip().lower() == "simulation":
            return wb[_sn]
    return None


def _grafici_sim_var_context(
    wb,
    company_curves: list[dict[str, Any]] | None = None,
) -> tuple[str, int, int, int] | None:
    """Riferimento foglio Simulation + righe ticker per INDEX/MATCH Var.%."""
    ws = _simulation_ws(wb)
    if ws is None:
        return None
    _rows: list[int] = []
    _cohort = {
        str(_e.get("ticker") or "").strip().upper()
        for _e in (company_curves or [])
        if str(_e.get("ticker") or "").strip()
    }
    _max_r = int(ws.max_row or 0)
    for _rn in range(4, _max_r + 1):
        _tk = str(ws.cell(row=_rn, column=1).value or "").strip().upper()
        if not _tk or _tk in ("—", "-", "N/D"):
            continue
        if _cohort and _tk not in _cohort:
            continue
        _rows.append(_rn)
    if not _rows and _cohort:
        for _rn in range(4, _max_r + 1):
            _tk = str(ws.cell(row=_rn, column=1).value or "").strip().upper()
            if _tk and _tk in _cohort:
                _rows.append(_rn)
    if not _rows:
        for _rn in range(4, _max_r + 1):
            _tk = str(ws.cell(row=_rn, column=1).value or "").strip().upper()
            if _tk and _tk not in ("—", "-", "N/D"):
                _rows.append(_rn)
    if not _rows:
        return None
    _c_lo = SIMULATION_36_HEADERS.index("Var. Giorn. %") + 1
    _first = 4
    _last = max(int(ws.max_row or 0), max(_rows))
    return (_excel_sheet_ref(ws), _first, _last, _c_lo)


def _clear_legacy_grafici_var_columns(ws_p) -> None:
    """Rimuove formule Var.% rimaste in T–W (grafico ora legge E–H)."""
    for _col in range(_GRAF_VAR_LEGACY_COL0, _GRAF_VAR_LEGACY_COL0 + _GRAF_VAR_N):
        for _row in range(15, 250):
            _c = ws_p.cell(row=_row, column=_col)
            if isinstance(_c.value, str) and "Simulation!" in _c.value:
                _c.value = None


def _var_pct_for_excel(raw: Any) -> float | None:
    """Valore Var.% Simulation → numero per formato '+0.00\"%\"'."""
    if raw is None:
        return None
    if isinstance(raw, str) and str(raw).strip() in ("", "—", "-", "N/D"):
        return None
    try:
        f = float(raw)
    except (TypeError, ValueError):
        return None
    if f != f:
        return None
    if abs(f) <= 1.5:
        return round(f * 100.0, 4)
    return round(f, 4)


def _seed_grafici_var_row_from_sim(
    ws_p,
    ws_sim,
    row: int,
    ticker: str,
    c_var_lo: int,
    sim_first: int,
    sim_last: int,
) -> None:
    """Valori iniziali (prima del ricalcolo Excel) così il grafico non resta vuoto."""
    _tk = str(ticker or "").strip().upper()
    if not _tk or _tk in ("—", "-", "N/D"):
        return
    _match_row: int | None = None
    for _rn in range(sim_first, sim_last + 1):
        if str(ws_sim.cell(row=_rn, column=1).value or "").strip().upper() == _tk:
            _match_row = _rn
            break
    if not _match_row:
        return
    for _ji in range(_GRAF_VAR_N):
        _vc = int(c_var_lo) + _ji
        _v = _var_pct_for_excel(ws_sim.cell(row=_match_row, column=_vc).value)
        if _v is not None:
            ws_p.cell(row=row, column=_GRAF_VAR_COL0 + _ji).value = _v


def _write_grafici_variation_block(
    ws_p,
    start_row: int,
    *,
    sim_ref: str,
    sim_first: int,
    sim_last: int,
    c_var_lo: int,
    sel_data0: int,
    sel_last: int,
    company_curves: list[dict[str, Any]] | None = None,
    ws_sim=None,
    x_col0: int = _GRAF_VAR_COL0,
    include_title: bool = True,
) -> tuple[int, int, int]:
    """
    Var.% 1g·1M·3M·6M — una riga per società in tabella coorte (righe sel_data0…sel_last).
    Filtro: col. G5 Var.% e/o menu Società A/B (C2, C3).
    """
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as _gcl

    _clear_legacy_grafici_var_columns(ws_p)
    _row = start_row
    if include_title:
        ws_p.cell(
            row=_row,
        column=1,
        value=(
                "5 · Variazioni % — 1g · 1M · 3M · 6M "
                "(tabella coorte: G5 Var.% e/o menu Società A/B righe 2–3)"
            ),
        ).font = Font(bold=True, size=10, color="1F3864")
        _row += 1
    _axis = _row
    _write_x_var_horizon_axis_row(ws_p, _axis, x_col0=x_col0)
    _row += 1
    _data0 = _row
    _n = 0
    _tc = _gcl(_GRAF_SEL_COL_TICKER)
    for _sr in range(int(sel_data0), int(sel_last) + 1):
        _gate = _xl_grafici_var_gate_expr(_sr)
        _tick = f"${_tc}${_sr}"
        _cr = _row
        ws_p.cell(row=_cr, column=1).value = f'=IF({_gate}<>1,"",{_tick})'
        _tk_plain = str(ws_p.cell(row=_sr, column=_GRAF_SEL_COL_TICKER).value or "").strip()
        for _ji in range(_GRAF_VAR_N):
            _vcol = int(c_var_lo) + _ji
            _pct: float | None = None
            if ws_sim is not None and _tk_plain:
                _pct = _read_sim_var_pct(
                    ws_sim, _tk_plain, _vcol, sim_first, sim_last
                )
            _ce = ws_p.cell(row=_cr, column=x_col0 + _ji)
            _ce.value = _xl_grafici_var_row_cell(_sr, pct_value=_pct)
            _ce.number_format = FMT_VARIATION
        _row += 1
        _n += 1
    return _axis, _data0, _n


def _grafici_src_bang() -> str:
    return f"'{GRAFICI_SHEET}'!"


def _predizione_guida_sheet_ref(wb) -> str | None:
    if wb is None:
        return None
    for _sn in wb.sheetnames:
        if str(_sn).strip() == _PREDIZIONE_GUIDA_SHEET:
            return _excel_sheet_ref(wb[_sn])
    return None


def _grafici_mu_links_broken(wb) -> bool:
    """
    True se il foglio Grafici punta ancora a celle vuote (cache SN7 / IFERROR su 2a)
    invece che a «Predizione — guida».
    """
    if wb is None or GRAFICI_SHEET not in wb.sheetnames:
        return False
    if not _predizione_guida_sheet_ref(wb):
        return False
    ws = wb[GRAFICI_SHEET]
    for _rn in range(1, min(int(ws.max_row or 0) + 1, 250)):
        _raw = str(ws.cell(row=_rn, column=1).value or "")
        if "cache SN7" in _raw:
            return True
    for _rn in range(70, min(int(ws.max_row or 0) + 1, 130)):
        _v = ws.cell(row=_rn, column=3).value
        if isinstance(_v, str) and "Grafici'!C$" in _v.replace(" ", ""):
            return True
        if isinstance(_v, str) and "Grafici!C$" in _v.replace(" ", ""):
            return True
    return False


def _predizione_mu_entries_for_grafici(
    wb,
) -> tuple[list[tuple[str, int | None]], bool, dict[str, int]]:
    """
    Serie μ per Grafici: (etichetta menu, riga guida o None), flag link guida, mappa REF_SPECS→riga.

    Preferisce lo scan di tutte le righe μ sulla guida; se assente usa ``_REF_SPECS`` +
    ``_find_predizione_mu_rows`` (evita cache SN7 vuota quando la guida esiste già).
    """
    _rows = _find_predizione_mu_rows(wb) if wb is not None else {}
    _scan = _scan_predizione_mu_series(wb) if wb is not None else []
    if _scan:
        return _scan, True, _rows
    if _rows:
        _entries: list[tuple[str, int | None]] = []
        for _mlab, _ in _REF_SPECS:
            _rn = int(_rows.get(_mlab) or 0)
            if _rn:
                _entries.append((_mlab, _rn))
        if _entries:
            return _entries, True, _rows
    return [(_mlab, None) for _mlab, _ in _REF_SPECS], False, _rows


def _find_predizione_mu_rows(wb) -> dict[str, int]:
    """Mappa etichetta REF_SPECS → riga μ sulla sheet «Predizione — guida»."""
    _out: dict[str, int] = {}
    if wb is None:
        return _out
    _ws = None
    for _sn in wb.sheetnames:
        if str(_sn).strip() == _PREDIZIONE_GUIDA_SHEET:
            _ws = wb[_sn]
            break
    if _ws is None:
        return _out
    for _rn in range(1, int(_ws.max_row or 0) + 1):
        _raw = str(_ws.cell(row=_rn, column=1).value or "").strip()
        if not _raw or not _raw.startswith("μ"):
            continue
        _low = _raw.lower()
        for _mlab, _needles in _PRED_MU_ROW_NEEDLES.items():
            if _mlab in _out:
                continue
            if any(_n in _low for _n in _needles):
                _out[_mlab] = _rn
    return _out


def _xl_grafici_mu_from_predizione_row(
    wb,
    row_num: int,
    col_idx: int,
    offsets: tuple[int, ...],
) -> str:
    """Formula μ da riga «Predizione — guida» (punti % → frazione grafico)."""
    from openpyxl.utils import get_column_letter as _gcl

    if not row_num or col_idx >= len(offsets):
        return "=NA()"
    _pref = _predizione_guida_sheet_ref(wb)
    if not _pref:
        return "=NA()"
    try:
        _ix = list(_SN7_REF_CAL_X).index(int(offsets[col_idx]))
    except ValueError:
        return "=NA()"
    _cl = _gcl(2 + _ix)
    return f"={_pref}{_cl}${int(row_num)}/100"


def _xl_grafici_mu_from_predizione(
    wb,
    mlab: str,
    col_idx: int,
    offsets: tuple[int, ...],
    *,
    pred_rows: dict[str, int] | None = None,
) -> str:
    """Legge μ dalla tabella numerica «Predizione — guida» (stessi % della guida)."""
    _rows = pred_rows if pred_rows is not None else _find_predizione_mu_rows(wb)
    _pr = int(_rows.get(mlab) or 0)
    if not _pr:
        return "=NA()"
    return _xl_grafici_mu_from_predizione_row(wb, _pr, col_idx, offsets)


def _write_grafici_mu_source_rows(
    ws_g,
    row: int,
    *,
    wb,
    offsets: tuple[int, ...],
    sn7_ref_mu: dict | None,
    num_fmt: str,
    value_divisor: float = 100.0,
    detail_layout: bool = False,
) -> int:
    """Righe μ riferimento in tabella 2a/2b/2c (formule → Predizione — guida)."""
    _pred, _linked, _pred_rows = (
        _predizione_mu_entries_for_grafici(wb) if wb is not None else ([], False, {})
    )
    if _linked and _pred:
        for _mlab, _pr in _pred:
            _rn = int(_pr or _pred_rows.get(_mlab) or 0)
            if not _rn:
                continue
            if detail_layout:
                for _ji, _off in enumerate(offsets):
                    _lab = f"{int(_off):+d}" if int(_off) > 0 else str(int(_off))
                    _vals = (
                        _mlab,
                        "—",
                        _lab,
                        int(_off),
                        None,
                        "—",
                        None,
                        None,
                        "—",
                        "standard",
                        None,
                        None,
                        "μ storico riferimento",
                        None,
                        None,
                        None,
                        None,
                    )
                    for _ci, _vv in enumerate(_vals, start=1):
                        _c = ws_g.cell(row=row, column=_ci)
                        if _ci == 14:
                            _c.value = _xl_grafici_mu_from_predizione_row(
                                wb, _rn, _ji, offsets
                            )
                            _c.number_format = num_fmt
                        else:
                            _c.value = _vv if _vv is not None else "—"
                    row += 1
            else:
                ws_g.cell(row=row, column=1, value=_mlab)
                ws_g.cell(row=row, column=2, value="—")
                for _ji in range(len(offsets)):
                    _ce = ws_g.cell(row=row, column=3 + _ji)
                    _ce.value = _xl_grafici_mu_from_predizione_row(
                        wb, _rn, _ji, offsets
                    )
                    _ce.number_format = num_fmt
                row += 1
        return row
    if not isinstance(sn7_ref_mu, dict):
        return row
    for _mlab, _key in _REF_SPECS:
        _rlist = _pad10(sn7_ref_mu.get(_key))
        _mu_row = _sn7_ref_mu_at_predizione_table_offsets(_rlist, offsets)
        if detail_layout:
            for _ji, _off in enumerate(offsets):
                _mu = _mu_row[_ji] if _ji < len(_mu_row) else None
                _lab = f"{int(_off):+d}" if int(_off) > 0 else str(int(_off))
                _vals = (
                    _mlab,
                    "—",
                    _lab,
                    int(_off),
                    None,
                    "—",
                    None,
                    None,
                    "—",
                    "standard",
                    None,
                    None,
                    "μ storico riferimento",
                    _mu,
                    None,
                    None,
                    None,
                )
                for _ci, _vv in enumerate(_vals, start=1):
                    _c = ws_g.cell(row=row, column=_ci)
                    if _ci == 14 and isinstance(_vv, (int, float)) and _vv == _vv:
                        _c.value = float(_vv) / value_divisor
                        _c.number_format = num_fmt
                    else:
                        _c.value = _vv if _vv is not None else "—"
                row += 1
        else:
            ws_g.cell(row=row, column=1, value=_mlab)
            ws_g.cell(row=row, column=2, value="—")
            for _ji, _mu in enumerate(_mu_row):
                _ce = ws_g.cell(row=row, column=3 + _ji)
                if _mu is not None and _mu == _mu:
                    _ce.value = float(_mu) / value_divisor
                    _ce.number_format = num_fmt
                else:
                    _ce.value = None
            row += 1
    return row


def _grafici_mu_values_for_offsets(
    sn7_ref_mu: dict[str, list] | None,
    mlab: str,
    offsets: tuple[int, ...],
) -> list[float | None]:
    """Fallback: bundle sn7_ref_mu (punti % assoluti, come Predizione — guida)."""
    _key = dict(_REF_SPECS).get(mlab)
    if not _key or not isinstance(sn7_ref_mu, dict):
        return [None] * len(offsets)
    _rlist = _pad10(sn7_ref_mu.get(_key))
    return _sn7_ref_mu_at_predizione_table_offsets(_rlist, offsets)


def _xl_grafici_vis_col(ticker_cell: str) -> str:
    """Vis. colonna matrice: coorte intera oppure ticker = A/B."""
    _ca = f"$C${_GRAF_PANEL_CO_A}"
    _cb = f"$C${_GRAF_PANEL_CO_B}"
    _all = f"$C${_GRAF_PANEL_ALL}"
    return (
        f"=IF({_all}=1,1,IF(OR({ticker_cell}={_ca},AND({_cb}<>\"\","
        f"{_cb}<>\"{_CO_B_NONE_XL}\",{ticker_cell}={_cb})),1,0))"
    )


def _xl_grafici_cross_pct(
    ticker_cell: str,
    vis_cell: str,
    wide_data0: int,
    wide_last: int,
    col_idx: int,
) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _src = _grafici_src_bang()
    _cl = _gcl(3 + col_idx)
    _inner = (
        f"IFERROR(INDEX({_src}{_cl}${wide_data0}:{_src}{_cl}${wide_last},"
        f"MATCH({ticker_cell},{_src}$A${wide_data0}:{_src}$A${wide_last},0)),NA())"
    )
    return f"=IF({vis_cell}=1,{_inner},NA())"


def _xl_grafici_cross_mu(
    vis_mu: str,
    mu_row: int,
    col_idx: int,
) -> str:
    from openpyxl.utils import get_column_letter as _gcl

    _src = _grafici_src_bang()
    _cl = _gcl(3 + col_idx)
    return f"=IF({vis_mu}=1,{_src}{_cl}${mu_row},NA())"


_G4_FIELD_BY_WIDE_KEY = {
    "curva": "pct_curva",
    "storico": "pct_reale",
    "modello": "pct_modello",
}

_GRAF_CHART_FIELD_BY_INCLUDE_COL: dict[int, str] = {
    _GRAF_SEL_COL_G1: "pct_curva",
    _GRAF_SEL_COL_G2: "pct_reale",
    _GRAF_SEL_COL_G3: "pct_modello",
}

# Prime N società in tabella coorte: G1–G3 attivi di default (resto 0 se si espande).
_GRAF_DEFAULT_G123_ON_N = 12


def _pct_fraction_from_curve(ent: dict[str, Any], field: str, offset: int) -> float | None:
    """Frazione Excel (0.05 = 5%) da ``points`` società."""
    _fld = field
    if _fld == "pct_curva":
        _fld = "pct_seq"
    for _p in ent.get("points") or []:
        if _p.get("nodo") != "standard":
            continue
        try:
            if int(_p["offset"]) != int(offset):
                continue
        except (TypeError, ValueError, KeyError):
            continue
        _v = _p.get(_fld)
        if _fld == "pct_seq" and (_v is None or _v != _v):
            _v = _p.get("pct_curva")
        if _v is not None and _v == _v:
            return float(_v) / 100.0
    return None


def _xl_grafici_g4_tipo_cond(cohort_row: int, tipo_col: int) -> str:
    """
    Serie visibile se G4 comb.=1 e (spunta tipo=1 oppure tutte e tre le spunte tipo a 0 → mostra tutte).
    """
    from openpyxl.utils import get_column_letter as _gcl

    _g4 = f"${_gcl(_GRAF_SEL_COL_G4)}${cohort_row}"
    _tt = f"${_gcl(tipo_col)}${cohort_row}"
    _cu = f"${_gcl(_GRAF_SEL_COL_G4_CURVA)}${cohort_row}"
    _st = f"${_gcl(_GRAF_SEL_COL_G4_STOR)}${cohort_row}"
    _mo = f"${_gcl(_GRAF_SEL_COL_G4_MOD)}${cohort_row}"
    _all_off = f"AND({_cu}=0,{_st}=0,{_mo}=0)"
    return f"AND({_g4}=1,OR({_tt}=1,{_all_off}))"


def _xl_grafici_g4_combo_pct_cell(cond: str, frac: float | None) -> str:
    """Δ% combinato con literal (visibile subito in Excel senza INDEX)."""
    if frac is None or frac != frac:
        return f"=IF(NOT({cond}),NA(),NA())"
    _lit = str(round(float(frac), 8)).replace(",", ".")
    return f"=IF(NOT({cond}),NA(),{_lit})"


def _xl_grafici_cross_pct_raw(
    ticker_cell: str,
    wide_data0: int,
    wide_last: int,
    col_idx: int,
) -> str:
    """INDEX/MATCH sorgente senza filtro vis (registro serie combinato)."""
    from openpyxl.utils import get_column_letter as _gcl

    _src = _grafici_src_bang()
    _cl = _gcl(3 + col_idx)
    return (
        f"=IFERROR(INDEX({_src}{_cl}${wide_data0}:{_src}{_cl}${wide_last},"
        f"MATCH({ticker_cell},{_src}$A${wide_data0}:{_src}$A${wide_last},0)),NA())"
    )


def _xl_grafici_extra_slot_cell(
    slot_cell: str,
    reg_lab_col: str,
    reg_start: int,
    reg_end: int,
    val_col: str,
    col_idx: int,
) -> str:
    """Come combinato: una serie scelta nel menu (μ o società)."""
    return _xl_grafici_combo_slot_cell(
        slot_cell, reg_lab_col, reg_start, reg_end, val_col, col_idx
    )


def _xl_section5_slot_mu(
    slot_cell: str,
    mu_rows: dict[str, int],
    col_idx: int,
) -> str:
    """Menu sez. 5a → riga μ tabella 2a (solo etichette ``_REF_SPECS``)."""
    from openpyxl.utils import get_column_letter as _gcl

    _cl = _gcl(3 + col_idx)
    _src = _grafici_src_bang()
    _inner = "NA()"
    for _mlab, _ in reversed(_REF_SPECS):
        _mr = int(mu_rows.get(_mlab) or 0)
        if _mr:
            _inner = f'IF({slot_cell}="{_mlab}",{_src}{_cl}${_mr},{_inner})'
    return (
        f"=IF(OR({slot_cell}=\"\",{slot_cell}=\"{_GRAF_COMBO_NONE_XL}\"),"
        f"NA(),{_inner})"
    )


def _xl_grafici_combo_slot_cell(
    slot_cell: str,
    reg_lab_col: str,
    reg_start: int,
    reg_end: int,
    val_col: str,
    col_idx: int,
) -> str:
    """Una cella del grafico combinato: serie scelta nel menu della riga."""
    _rng_lab = f"${reg_lab_col}${reg_start}:${reg_lab_col}${reg_end}"
    _rng_val = f"${val_col}${reg_start}:${val_col}${reg_end}"
    _match = f"MATCH({slot_cell},{_rng_lab},0)"
    return (
        f'=IF(OR({slot_cell}="",{slot_cell}="{_GRAF_COMBO_NONE_XL}"),NA(),'
        f"IFERROR(INDEX({_rng_val},{_match}),NA()))"
    )


def _write_grafici_combo_registry(
    ws_p,
    reg_start: int,
    *,
    company_curves: list[dict[str, Any]],
    wide: dict[str, Any],
    offsets: tuple[int, ...],
    pct_fmt: str,
    mu_rows: dict[str, int],
    wb=None,
    pred_mu_rows: dict[str, int] | None = None,
    sn7_ref_mu: dict[str, list] | None = None,
) -> tuple[int, int, str]:
    """
    Registro nascosto: una riga per ogni opzione «TICKER · curva|storico|modello» (+ μ).
    Ritorna (start_row, end_row, formula1 elenco menu).
    """
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as _gcl

    _lab_c = _gcl(_GRAF_COMBO_REG_COL_LAB)
    _val_c0 = _GRAF_COMBO_REG_COL0
    _n_std = len(offsets)
    _row = reg_start
    _hdr = _grafici_set_cell(
        ws_p,
        _row,
        _GRAF_COMBO_REG_COL_LAB,
        value="Etichetta serie",
        font=Font(bold=True, size=8),
    )
    _row += 1
    _first = _row
    _specs: list[tuple[str, str, int, int]] = [
        ("curva ricalibr.", "curva", int(wide["curva"]["data0"]), int(wide["curva"]["last"])),
        ("storico", "storico", int(wide["storico"]["data0"]), int(wide["storico"]["last"])),
        ("modello", "modello", int(wide["modello"]["data0"]), int(wide["modello"]["last"])),
    ]
    for _ent in company_curves:
        _tk = str(_ent["ticker"])
        for _slab, _wk, _d0, _d1 in _specs:
            _lbl = f"{_tk} · {_slab}"
            _grafici_set_cell(ws_p, _row, _GRAF_COMBO_REG_COL_LAB, value=_lbl)
            _tk_ref = f'"{_tk}"'
            for _ji in range(_n_std):
                _vc = _gcl(_val_c0 + _ji)
                _grafici_set_cell(
                    ws_p,
                    _row,
                    _val_c0 + _ji,
                    value=_xl_grafici_cross_pct_raw(_tk_ref, _d0, _d1, _ji),
                )
                ws_p.cell(row=_row, column=_val_c0 + _ji).number_format = pct_fmt
            _row += 1
    _pred_series, _pred_linked, _pred_rows_map = (
        _predizione_mu_entries_for_grafici(wb) if wb is not None else ([], False, {})
    )
    if _pred_linked and _pred_series:
        for _mlab, _pr in _pred_series:
            _rn = int(_pr or (pred_mu_rows or {}).get(_mlab) or 0)
            if not _rn:
                continue
            _lbl = f"{_mlab} · curva"
            _grafici_set_cell(ws_p, _row, _GRAF_COMBO_REG_COL_LAB, value=_lbl)
            for _ji in range(_n_std):
                _ce = _grafici_set_cell(ws_p, _row, _val_c0 + _ji)
                _ce.value = _xl_grafici_mu_from_predizione_row(wb, _rn, _ji, offsets)
                _ce.number_format = pct_fmt
            _row += 1
    else:
        for _mlab, _ in _REF_SPECS:
            _lbl = f"{_mlab} · curva"
            _grafici_set_cell(ws_p, _row, _GRAF_COMBO_REG_COL_LAB, value=_lbl)
            _mr = int(mu_rows.get(_mlab) or 0)
            for _ji in range(_n_std):
                _ce = _grafici_set_cell(ws_p, _row, _val_c0 + _ji)
                if _predizione_guida_sheet_ref(wb) and (pred_mu_rows or {}).get(_mlab):
                    _ce.value = _xl_grafici_mu_from_predizione(
                        wb, _mlab, _ji, offsets, pred_rows=pred_mu_rows
                    )
                elif _mr:
                    _cl_mu = _gcl(3 + _ji)
                    _ce.value = f"={_grafici_src_bang()}{_cl_mu}${_mr}"
                else:
                    _vals = _grafici_mu_values_for_offsets(sn7_ref_mu, _mlab, offsets)
                    _v = _vals[_ji] if _ji < len(_vals) else None
                    _ce.value = (
                        (float(_v) / 100.0) if _v is not None and _v == _v else None
                    )
                _ce.number_format = pct_fmt
            _row += 1
    _last = _row - 1
    for _ri, _lbl in enumerate([_GRAF_COMBO_NONE_XL], start=1):
        _grafici_set_cell(ws_p, _ri, _GRAF_COMBO_LIST_COL, value=_lbl)
    _li = 2
    for _rr in range(_first, _last + 1):
        _lb = ws_p.cell(row=_rr, column=_GRAF_COMBO_REG_COL_LAB).value
        _grafici_set_cell(ws_p, _li, _GRAF_COMBO_LIST_COL, value=_lb)
        _li += 1
    _cl_list = _gcl(_GRAF_COMBO_LIST_COL)
    _list_formula = f"${_cl_list}$1:${_cl_list}${_li - 1}"
    return _first, _last, _list_formula


def _write_grafici_slot_chart_block(
    ws_p,
    start_row: int,
    *,
    block_title: str,
    offsets: tuple[int, ...],
    slot_cells: tuple[str, ...],
    pct_fmt: str,
    combo_reg_first: int,
    combo_reg_last: int,
    combo_reg_lab_col: str,
    combo_reg_val_col: str,
) -> tuple[int, int, int]:
    """Dati grafico %: solo serie scelte nei menu (fuori dall'area del grafico)."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as _gcl

    _n_std = len(offsets)
    _row = start_row
    ws_p.cell(row=_row, column=1, value=block_title).font = Font(bold=True, size=10, color="1F3864")
    _row += 1
    _axis = _row
    _write_x_session_axis_row(ws_p, _axis, offsets)
    _row += 1
    _data0 = _row
    _n = 0
    for _slot_cell in slot_cells:
        _cr = _row
        ws_p.cell(row=_cr, column=1).value = (
            f'=IF(OR({_slot_cell}="",{_slot_cell}="{_GRAF_COMBO_NONE_XL}"),"",'
            f"{_slot_cell})"
        )
        for _ji in range(_n_std):
            _val_col_letter = _gcl(_GRAF_COMBO_REG_COL0 + _ji)
            _ce = ws_p.cell(row=_cr, column=3 + _ji)
            _ce.value = _xl_grafici_combo_slot_cell(
                _slot_cell,
                combo_reg_lab_col,
                combo_reg_first,
                combo_reg_last,
                _val_col_letter,
                _ji,
            )
            _ce.number_format = pct_fmt
        _row += 1
        _n += 1
    return _axis, _data0, _n


_SECTION_STOP_PREFIXES = (
    "1 —",
    "2a —",
    "2b —",
    "2c —",
    "2d —",
    "3 —",
    "4 —",
    "5a —",
    "5b —",
    "A —",
)


def _cell_a(ws_g, row: int) -> str:
    v = ws_g.cell(row=row, column=1).value
    return str(v or "").strip()


def _find_wide_table_range(ws_g, section_no: str, *, title_hint: str = "") -> tuple[int, int] | None:
    """Ritorna (first_data_row, last_data_row) per sez. 2a/2b/2c sul foglio grafici."""
    _prefix = f"{section_no} —"
    _max_r = int(ws_g.max_row or 0)
    for _r in range(1, _max_r + 1):
        _a = _cell_a(ws_g, _r)
        if not _a.startswith(_prefix):
            continue
        if title_hint and title_hint not in _a:
            continue
        _data0 = _r + 2
        if _data0 > _max_r:
            return None
        _last = _data0
        for _rr in range(_data0, _max_r + 1):
            _av = _cell_a(ws_g, _rr)
            if not _av:
                break
            if _rr > _data0 and any(_av.startswith(p) for p in _SECTION_STOP_PREFIXES):
                if not _av.startswith(_prefix):
                    break
            _last = _rr
        if _last >= _data0:
            return _data0, _last
    return None


def _grafici_std_points_by_offset(
    ent: dict[str, Any],
) -> dict[int, dict[str, Any]]:
    return {
        int(p["offset"]): p
        for p in (ent.get("points") or [])
        if p.get("nodo") == "standard"
    }


def _grafici_log_curve_build_summary(
    company_curves: list[dict[str, Any]],
    offsets: tuple[int, ...],
) -> None:
    """Log ticker saltati o con serie vuote (storico / curva / modello)."""
    if not company_curves:
        _grafici_step("[Grafici] Nessuna società in coorte — verificare foglio Simulation.")
        return
    _probe = tuple(int(o) for o in offsets[: min(4, len(offsets))])
    _n_ok_c = _n_ok_s = _n_ok_m = 0
    for _ent in company_curves:
        _tk = str(_ent.get("ticker") or "?")
        _std = _grafici_std_points_by_offset(_ent)
        _has_c = any(
            (_std.get(_o) or {}).get("pct_curva") is not None
            and (_std.get(_o) or {}).get("pct_curva") == (_std.get(_o) or {}).get("pct_curva")
            for _o in _probe
        )
        _has_s = any(
            (_std.get(_o) or {}).get("pct_reale") is not None
            and (_std.get(_o) or {}).get("pct_reale") == (_std.get(_o) or {}).get("pct_reale")
            for _o in _probe
        )
        _has_m = any(
            (_std.get(_o) or {}).get("pct_modello") is not None
            and (_std.get(_o) or {}).get("pct_modello") == (_std.get(_o) or {}).get("pct_modello")
            for _o in _probe
        )
        if _has_c:
            _n_ok_c += 1
        if _has_s:
            _n_ok_s += 1
        if _has_m:
            _n_ok_m += 1
        if _has_c or _has_s or _has_m:
            continue
        _reasons: list[str] = []
        if not _has_s:
            _reasons.append(
                "storico vuoto (eseguire --refresh senza GRAFICI_FAST=1; HistLib/yfinance)"
            )
        if not _has_c:
            _reasons.append("curva ricalibrata vuota (seq_curve / enrich)")
        if not _has_m:
            _reasons.append(
                "modello vuoto (Pred±N su Simulation o model_curve_valid=false)"
            )
        _grafici_step(f"[Grafici] Saltato {_tk}: {'; '.join(_reasons)}.")
    _grafici_step(
        f"[Grafici] Serie con dati: curva {_n_ok_c}/{len(company_curves)}, "
        f"storico {_n_ok_s}/{len(company_curves)}, "
        f"modello {_n_ok_m}/{len(company_curves)}."
    )


def build_company_curves_from_simulation(
    wb,
    ws_sim,
    main_row_positions: list[int],
    *,
    universe: dict | None = None,
    sim_pred_data: dict | None = None,
) -> list[dict[str, Any]]:
    """Costruisce curve società (punti per nodo) dal foglio Simulation."""
    if not main_row_positions:
        return []
    _today = date.today()
    _offsets = SIM_GRAFICI_CURVE_OFFSETS
    _pred_log: set[str] = set()
    _harvest_by_pk: dict[str, dict[str, Any]] = {}
    try:
        for _hr in _sim_harvest_rows_from_workbook(wb):
            _hk = _past_pred_key_from_harvest_row(_hr)
            if _hk:
                _harvest_by_pk[_hk] = _hr
    except Exception:
        _harvest_by_pk = {}

    _graf_accum: dict[str, dict[str, Any]] = {}
    for _rn0 in main_row_positions:
        _tkv0 = ws_sim.cell(row=_rn0, column=1).value
        _cd0 = _sn7_parse_sheet_date_ddmmyyyy(ws_sim.cell(row=_rn0, column=3).value)
        if not _tkv0 or _cd0 is None:
            continue
        _tk0 = str(_tkv0).strip().upper()
        _pk0 = f"{_tk0}|{_cd0.isoformat()}"
        if _pk0 in _graf_accum:
            continue
        if isinstance(universe, dict) and _pk0 in universe:
            _graf_accum[_pk0] = dict(universe[_pk0])
        elif _pk0 in _harvest_by_pk:
            _graf_accum[_pk0] = dict(
                _stub_past_pred_row_from_sim_harvest(_harvest_by_pk[_pk0])
            )
        else:
            _graf_accum[_pk0] = {
                "ticker": _tk0,
                "completion_date": _cd0.isoformat()[:10],
                "company_name_full": (
                    str(ws_sim.cell(row=_rn0, column=2).value or "").strip() or "—"
                ),
                "nct_id": str(ws_sim.cell(row=_rn0, column=7).value or "").strip().upper(),
                "sponsor_match": str(ws_sim.cell(row=_rn0, column=4).value or "").strip()
                or "N/D",
                "nct_relation_type": str(ws_sim.cell(row=_rn0, column=9).value or "").strip()
                or "—",
            }
    try:
        _accuracy_sim_merge_live_pred(_graf_accum, _harvest_by_pk, sim_pred_data)
    except Exception as exc:
        log_prediction_error("grafici_merge_live_pred", exc)
    _nb_hf = 0
    if not _grafici_fast_enabled():
        try:
            _nb_hf = int(
                _accuracy_sim_histlib_backfill_session_closes(_graf_accum) or 0
            )
        except Exception as exc:
            log_prediction_error("grafici_histlib_backfill", exc)
        if _nb_hf:
            _grafici_step(f"[Grafici] HistLib backfill: {_nb_hf} celle close.")
    else:
        _grafici_step("[Grafici] HistLib backfill: skip (GRAFICI_FAST=1).")

    _company_curves: list[dict[str, Any]] = []
    _ser_cache: dict = {}
    _tk_refresh: list[str] = []
    for _rn in main_row_positions:
        _tkv = ws_sim.cell(row=_rn, column=1).value
        if _tkv:
            _tk_refresh.append(str(_tkv).strip().upper())
    if _tk_refresh and not _grafici_fast_enabled():
        _grafici_step(
            f"[Grafici] Yahoo/HistLib refresh per {len(set(_tk_refresh))} ticker "
            "(rete, può richiedere diversi minuti)…"
        )
        _grafici_refresh_histlib_from_yfinance(_tk_refresh, force=False, min_lag_days=3)
    elif _tk_refresh and _grafici_fast_enabled():
        _grafici_step("[Grafici] Yahoo/HistLib refresh: skip (GRAFICI_FAST=1).")

    if not _grafici_fast_enabled():
        try:
            _n_enr, _n_seq = _accuracy_sim_enrich_curve_records(
                _graf_accum,
                sim_pred_data=sim_pred_data,
                refresh_histlib=False,
                log_label="Grafici",
            )
            if _n_enr or _n_seq:
                _grafici_step(
                    f"[Grafici] Enrich seq_curve/calendario: {_n_enr} celle close, "
                    f"{_n_seq} curve seq."
                )
        except Exception as exc:
            log_prediction_error("grafici_enrich_curve_records", exc)
    else:
        _grafici_step(
            "[Grafici] Enrich seq_curve/calendario: skip (GRAFICI_FAST=1)."
        )

    _n_pos = len(main_row_positions)
    _grafici_step(f"[Grafici] Costruzione curve: 0/{_n_pos} righe Simulation…")
    for _idx, _rn in enumerate(main_row_positions, start=1):
        _tkv = ws_sim.cell(row=_rn, column=1).value
        _cd = _sn7_parse_sheet_date_ddmmyyyy(ws_sim.cell(row=_rn, column=3).value)
        if not _tkv or _cd is None:
            continue
        _tk = str(_tkv).strip().upper()
        _pk = f"{_tk}|{_cd.isoformat()}"
        _pr = dict(_graf_accum.get(_pk) or {})
        _sn7_cd7_overlay_seq_curve_from_live(
            _pr, row_key=_pk, ticker_upper=_tk, sim_pred_data=sim_pred_data
        )
        _pw = dict(_pr)
        _overlay_model_pcts_from_simulation_row(ws_sim, _rn, _pw)
        _grafici_ensure_seq_curve_on_pairwise(
            _pw, _cd, row_key=_pk, today=_today, ser_cache=_ser_cache
        )
        _accuracy_sim_synthesize_interp_nodes_from_post_d_only(_pw)
        _accuracy_sim_impute_missing_pre_cd_model_pcts(_pw)
        _graf_accum[_pk] = {**(_graf_accum.get(_pk) or {}), **_pw}
        _pts = _ristretta_curve_table_points(
            _pw,
            _cd,
            _today,
            pk=_pk,
            pred_log=_pred_log,
            wb=wb,
            offsets=_offsets,
            with_price_usd=True,
            ser_cache=_ser_cache,
            storico_basis="calendar",
        )
        _company_curves.append(
            {"ticker": _tk, "cd": _cd, "pk": _pk, "sim_row": _rn, "points": _pts}
        )
        if _idx == 1 or _idx == _n_pos or _idx % 5 == 0:
            _grafici_step(
                f"[Grafici] Costruzione curve: {_idx}/{_n_pos} "
                f"({len(_company_curves)} società)…"
            )
    _grafici_log_curve_build_summary(_company_curves, _offsets)
    return _company_curves


def _write_grafici_pct_wide_bundle(
    ws,
    start_row: int,
    company_curves: list[dict[str, Any]],
    *,
    wb,
    offsets: tuple[int, ...],
    sn7_ref_mu: dict | None,
    universe: dict | None,
    pct_fmt: str = "+0.0%;-0.0%;0.0%",
) -> dict[str, Any]:
    """Scrive tabelle 2a–2c sul foglio Grafici; ritorna indici righe per formule."""
    _sn7_mu = _resolve_sn7_ref_mu(sn7_ref_mu, wb)
    _n_co = len(company_curves)
    _univ = universe if isinstance(universe, dict) else None
    _row = start_row
    _, _wide_data0_c, _row = _write_curve_wide_table(
        ws,
        _row,
        "2a",
        company_curves,
        _sn7_mu,
        offsets,
        num_fmt=pct_fmt,
        value_field="pct_curva",
        include_mu=True,
        table_label="curva ricalibr.",
        universe=_univ,
        wb=wb,
    )
    _pred_entries, _pred_linked, _pred_mu_rows = _predizione_mu_entries_for_grafici(wb)
    _mu_rows = {
        _mlab: _wide_data0_c + _n_co + _ii
        for _ii, (_mlab, _) in enumerate(
            _pred_entries if _pred_entries else [(_m, None) for _m, _ in _REF_SPECS]
        )
    }
    _row += 2
    _, _wide_data0_s, _row = _write_curve_wide_table(
        ws,
        _row,
        "2b",
        company_curves,
        _sn7_mu,
        offsets,
        num_fmt=pct_fmt,
        value_field="pct_reale",
        include_mu=False,
        table_label="storico (dato vero)",
        universe=_univ,
    )
    _row += 2
    _, _wide_data0_m, _row = _write_curve_wide_table(
        ws,
        _row,
        "2c",
        company_curves,
        _sn7_mu,
        offsets,
        num_fmt=pct_fmt,
        value_field="pct_modello",
        include_mu=False,
        table_label="modello T−60",
        universe=_univ,
    )
    return {
        "curva": {"data0": _wide_data0_c, "last": _wide_data0_c + max(0, _n_co - 1)},
        "storico": {"data0": _wide_data0_s, "last": _wide_data0_s + max(0, _n_co - 1)},
        "modello": {"data0": _wide_data0_m, "last": _wide_data0_m + max(0, _n_co - 1)},
        "mu_rows": _mu_rows,
        "pred_mu_rows": _pred_mu_rows,
        "pred_linked": _pred_linked,
        "pred_entries": _pred_entries,
    }


def parse_grafici_source_from_workbook(wb) -> tuple[list[dict[str, Any]], dict[str, Any], tuple[int, ...]]:
    """
    Legge tabelle 2a–2c dal foglio «Grafici» (workbook già aperto).
    Ritorna (company_curves, wide_bundle, offsets).
    """
    if GRAFICI_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Foglio «{GRAFICI_SHEET}» assente — eseguire l'orchestrator o "
            "launch_grafici_sheet.py."
        )
    ws_g = wb[GRAFICI_SHEET]
    _rng_c = _find_wide_table_range(ws_g, "2a", title_hint="curva")
    _rng_s = _find_wide_table_range(ws_g, "2b", title_hint="storico")
    _rng_m = _find_wide_table_range(ws_g, "2c", title_hint="modello")
    if not _rng_c or not _rng_s or not _rng_m:
        raise ValueError(
            "Tabelle 2a/2b/2c non trovate su «Grafici». "
            "Rigenerare con launch_grafici_sheet.py o orchestrator."
        )
    _d0c, _d1c = _rng_c
    _mu_labels = {m[0] for m in _REF_SPECS}
    _company_curves: list[dict[str, Any]] = []
    _mu_rows: dict[str, int] = {}
    for _rr in range(_d0c, _d1c + 1):
        _tk = _cell_a(ws_g, _rr)
        if not _tk or _tk in ("Ticker / serie", "—"):
            continue
        if _tk in _mu_labels:
            _mu_rows[_tk] = _rr
        else:
            _company_curves.append({"ticker": _tk.split()[0].upper()})
    if not _company_curves:
        raise ValueError("Nessuna società nella tabella 2a.")
    _wide = {
        "curva": {"data0": _d0c, "last": _d1c},
        "storico": {"data0": _rng_s[0], "last": _rng_s[1]},
        "modello": {"data0": _rng_m[0], "last": _rng_m[1]},
        "mu_rows": _mu_rows,
    }
    return _company_curves, _wide, SIM_GRAFICI_CURVE_OFFSETS


def _wide_company_pct_rows_have_signal(
    wb,
    wide_key: str,
    *,
    check_cols: tuple[int, ...] = (4, 5, 8, 9),
) -> bool:
    """True se almeno una riga società nella tabella wide ha % numerici non nulli."""
    try:
        _co, _wide, _ = parse_grafici_source_from_workbook(wb)
    except ValueError:
        return False
    if not _co:
        return False
    ws_g = wb[GRAFICI_SHEET]
    _rng = _wide.get(wide_key)
    if not _rng:
        return False
    _d0, _d1 = int(_rng["data0"]), int(_rng["last"])
    _mu_labels = {m[0] for m in _REF_SPECS}
    for _rr in range(_d0, _d1 + 1):
        _tk = _cell_a(ws_g, _rr)
        if not _tk or _tk in _mu_labels:
            continue
        for _ci in check_cols:
            _v = ws_g.cell(row=_rr, column=_ci).value
            if isinstance(_v, (int, float)) and _v == _v and abs(float(_v)) > 1e-9:
                return True
    return False


def _wide_modello_table_looks_empty(wb) -> bool:
    """True se 2c esiste ma % modello società tutti vuoti/zero."""
    return not _wide_company_pct_rows_have_signal(wb, "modello")


def _grafici_source_needs_rebuild(wb) -> bool:
    """True se mancano tabelle 2a–2c o % società vuoti (curva/storico/modello)."""
    try:
        parse_grafici_source_from_workbook(wb)
    except ValueError:
        return True
    if not _wide_company_pct_rows_have_signal(wb, "curva"):
        return True
    if not _wide_company_pct_rows_have_signal(wb, "storico"):
        return True
    if not _wide_company_pct_rows_have_signal(wb, "modello"):
        return True
    return False


def _load_grafici_inputs_from_simulation(
    wb,
    *,
    sim_pred_data: dict | None = None,
    universe: dict | None = None,
) -> tuple[list[dict[str, Any]], tuple[int, ...]]:
    if "Simulation" not in wb.sheetnames:
        raise ValueError(
            "Foglio «Simulation» assente nel workbook — eseguire l'orchestrator completo."
        )
    from data_orchestrator import _simulation_sheet_pred_graph_row_positions

    ws_sim = wb["Simulation"]
    _pos = _simulation_sheet_pred_graph_row_positions(ws_sim)
    if not _pos:
        raise ValueError("Nessuna riga catalyst sul foglio «Simulation».")
    _co = build_company_curves_from_simulation(
        wb,
                ws_sim,
        _pos,
                universe=universe,
        sim_pred_data=sim_pred_data,
    )
    if not _co:
        raise ValueError("Nessuna curva società da Simulation.")
    return _co, SIM_GRAFICI_CURVE_OFFSETS


def rebuild_grafici_sheet_only(
    xlsx_path: str | Path | None = None,
    *,
    save: bool = True,
    ensure_source: bool = True,
    force_refresh_source: bool = False,
) -> Path:
    """Rigenera il foglio «Grafici» (curve da Simulation, tabelle 2a–2c sullo stesso foglio)."""
    from pathlib import Path as _Path

    from orchestrator_io_paths import FINAL_XLSX

    _path = _Path(xlsx_path or FINAL_XLSX)
    if not _path.is_file():
        raise FileNotFoundError(f"Workbook non trovato: {_path}")
    import openpyxl

    _mb = _path.stat().st_size / (1024 * 1024)
    _grafici_step(
        f"[Grafici] 1/5 Apertura workbook ({_mb:.1f} MB, openpyxl)…"
    )
    wb = openpyxl.load_workbook(_path, keep_links=False)
    _grafici_step(
        f"[Grafici] 2/5 Workbook aperto — {len(wb.sheetnames)} fogli."
    )
    if _grafici_mu_links_broken(wb):
        print(
            "[Grafici] Rilevati controlli μ non collegati a «Predizione — guida» "
            "(curve a zero) — riscrittura foglio Grafici.",
            flush=True,
        )
    _saved = _path
    try:
        if ensure_source and not force_refresh_source and not _grafici_source_needs_rebuild(wb):
            _grafici_step("[Grafici] Tabelle 2a–2c presenti; solo layout/grafici.")
        elif force_refresh_source:
            _grafici_step("[Grafici] 3/5 --refresh: ricalcolo curve da Simulation…")
        else:
            _grafici_step("[Grafici] 3/5 Ricalcolo curve da Simulation…")
        _co, _offsets = _load_grafici_inputs_from_simulation(wb)
        _grafici_step(
            f"[Grafici] 4/5 Scrittura foglio «Grafici» ({len(_co)} società)…"
        )
        write_grafici_sheet(
            wb,
            company_curves=_co,
            wide=None,
            offsets=_offsets,
        )
        if save:
            _grafici_step("[Grafici] 5/5 Salvataggio workbook…")
            try:
                try:
                    wb.calculation.fullCalcOnLoad = True
                    wb.calculation.calcMode = "auto"
                except Exception:
                    pass
                wb.save(_path)
            except PermissionError:
                from datetime import datetime

                _staged = _path.parent / (
                    f"{_path.stem}__grafici_staged_{datetime.now():%Y%m%d_%H%M%S}{_path.suffix}"
                )
                wb.save(_staged)
                _saved = _staged
                print(
                    f"[Grafici] File principale bloccato (Excel aperto). "
                    f"Salvato su:\n  {_staged.resolve()}",
                    flush=True,
                )
                print(
                    "[Grafici] Chiudi Excel sul file principale, poi sostituisci "
                    f"data\\{_path.name} con il file staged (o rilancia questo script).",
                    flush=True,
                )
            else:
                print(
                    f"[Grafici] Salvato: {_path.resolve()} — controlli μ collegati a "
                    f"«{_PREDIZIONE_GUIDA_SHEET}».",
                    flush=True,
                )
    finally:
        wb.close()
    return _saved


def write_grafici_sheet(
    wb,
    *,
    company_curves: list[dict[str, Any]],
    wide: dict[str, Any] | None = None,
    offsets: tuple[int, ...],
    pct_fmt: str = "+0.0%;-0.0%;0.0%",
    sn7_ref_mu: dict | None = None,
    universe: dict | None = None,
) -> None:
    """
    Foglio «Grafici»: coorte, tabelle sorgente 2a–2c (area bassa), grafici % e Var.%.
    Var.% dal foglio «Simulation». Rimuove il foglio legacy «Simulation — grafici» se presente.
    """
    if not company_curves:
        return
    _grafici_step(
        f"[Grafici] Layout e grafici Excel per {len(company_curves)} società…"
    )
    try:
        from openpyxl.chart import LineChart, Reference
        from openpyxl.chart.series_factory import SeriesFactory
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter as _gcl
    except Exception as exc:
        print(f"[Grafici] Modulo chart non disponibile: {exc}", flush=True)
        return

    if SIMULATION_GRAFICI_SHEET in wb.sheetnames:
        del wb[SIMULATION_GRAFICI_SHEET]

    if GRAFICI_SHEET in wb.sheetnames:
        del wb[GRAFICI_SHEET]
    try:
        _ins = wb.sheetnames.index("Simulation") + 1
    except ValueError:
        _ins = len(wb.sheetnames)
    ws_p = wb.create_sheet(GRAFICI_SHEET, _ins)
    ws_p.sheet_properties.tabColor = "4472C4"

    _n_co = len(company_curves)
    _n_std = len(offsets)
    _sn7_mu = _resolve_sn7_ref_mu(sn7_ref_mu, wb)
    _pred_entries, _pred_linked, _pred_mu_rows = _predizione_mu_entries_for_grafici(wb)
    _wide_start_row = _grafici_wide_table_start_row(
        company_curves,
        pred_linked=_pred_linked,
        pred_series_len=len(_pred_entries) if _pred_entries else len(_REF_SPECS),
    )
    if wide is None:
        wide = _write_grafici_pct_wide_bundle(
            ws_p,
            _wide_start_row,
            company_curves,
            wb=wb,
            offsets=offsets,
            sn7_ref_mu=_sn7_mu,
                universe=universe,
            pct_fmt=pct_fmt,
        )
    _w_curva = wide["curva"]
    _w_stor = wide["storico"]
    _w_mod = wide["modello"]
    _mu_rows: dict[str, int] = dict(wide.get("mu_rows") or {})
    _d0c, _d1c = int(_w_curva["data0"]), int(_w_curva["last"])
    _d0s, _d1s = int(_w_stor["data0"]), int(_w_stor["last"])
    _d0m, _d1m = int(_w_mod["data0"]), int(_w_mod["last"])

    # Elenco ticker (col. Z)
    for _i, _ent in enumerate(company_curves):
        ws_p.cell(row=1 + _i, column=_GRAF_LIST_COL_G, value=str(_ent["ticker"]))
    _list_end = _n_co + 1
    ws_p.cell(row=_list_end, column=_GRAF_LIST_COL_G, value=_CO_B_NONE_XL)
    _list_formula = (
        f"${_gcl(_GRAF_LIST_COL_G)}$1:${_gcl(_GRAF_LIST_COL_G)}${_list_end}"
    )

    ws_p.merge_cells("A1:H1")
    _t1 = ws_p.cell(
        row=1,
        column=1,
        value=(
            "Grafici · Δ% vs T−60. Tabella coorte (righe 6–): spunta G1–G5 per società; "
            "curve μ sotto la tabella. Mesh 5×4."
        ),
    )
    _t1.font = Font(bold=True, size=10, color="FFFFFF")
    _t1.fill = PatternFill("solid", fgColor="2F5496")
    _t1.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_p.row_dimensions[1].height = 36

    _sn7_mu = _resolve_sn7_ref_mu(sn7_ref_mu, wb)
    _pred_scan = _pred_entries
    if _pred_linked and not _pred_mu_rows:
        print(
            f"[Grafici] ATTENZIONE: «{_PREDIZIONE_GUIDA_SHEET}» presente ma righe μ non "
            "trovate — eseguire refresh_predizione_guida.py.",
            flush=True,
        )
    if _pred_linked:
        print(
            f"[Grafici] μ controllo → «{_PREDIZIONE_GUIDA_SHEET}» "
            f"({len(_pred_entries)} serie, formule live).",
            flush=True,
        )
    elif _sn7_mu:
        print(
            "[Grafici] μ controllo: da cache sn7_ref_mu — "
            f"rigenerare «{_PREDIZIONE_GUIDA_SHEET}» poi Grafici.",
            flush=True,
        )

    _reg_start_row = _GRAF_COMBO_REG_START_ROW
    _reg_first, _reg_last, _combo_list_formula = _write_grafici_combo_registry(
        ws_p,
        _reg_start_row,
        company_curves=company_curves,
        wide=wide,
        offsets=offsets,
        pct_fmt=pct_fmt,
        mu_rows=_mu_rows,
        wb=wb,
        pred_mu_rows=_pred_mu_rows,
        sn7_ref_mu=_sn7_mu,
    )
    _lab_col = _gcl(_GRAF_COMBO_REG_COL_LAB)
    _val_col = _gcl(_GRAF_COMBO_REG_COL0)
    # ── Pannello menu (fuori dai grafici) ───────────────────────────────────
    _def0 = company_curves[0]["ticker"] if company_curves else "—"
    ws_p.cell(row=2, column=1, value="Società A (Var.% + matrice)").font = Font(size=9)
    _ca = ws_p.cell(row=2, column=3, value=_def0)
    _ca.fill = PatternFill("solid", fgColor="DEEAF6")
    _add_list_validation(ws_p, "C2", _list_formula)
    ws_p.cell(row=3, column=1, value="Società B (Var.% + matrice)").font = Font(size=9)
    _cb = ws_p.cell(row=3, column=3, value=_CO_B_NONE_XL)
    _cb.fill = PatternFill("solid", fgColor="DEEAF6")
    _add_list_validation(ws_p, "C3", _list_formula)
    ws_p.cell(row=4, column=1, value="Mostra tutta la coorte in matrice (1/0)").font = Font(
        size=9
    )
    ws_p.cell(row=4, column=3, value=1).fill = PatternFill("solid", fgColor="DEEAF6")
    _add_yesno_validation(ws_p, "C4")

    _sel = _write_grafici_cohort_selector(
        ws_p,
        company_curves,
        _pred_scan,
        pred_linked=_pred_linked,
        combo_list_formula=_combo_list_formula,
    )

    _matrix_hdr = int(_sel["next_row"])
    _matrix_vis = _matrix_hdr + 1
    _matrix_lab0 = _matrix_vis + 1
    _fc = 4
    ws_p.cell(row=_matrix_hdr, column=1, value="Matrice visibilità").font = Font(
        bold=True, size=9, color="1F3864"
    )
    ws_p.cell(row=_matrix_hdr + 1, column=1, value="Serie →").font = Font(bold=True, size=9)
    ws_p.cell(row=_matrix_hdr + 1, column=2, value="Vis.").font = Font(bold=True, size=9)
    ws_p.cell(row=_matrix_vis, column=2, value="(1/0)").font = Font(size=8, italic=True)
    for _ci, _ent in enumerate(company_curves):
        _col = _fc + _ci
        _cl = _gcl(_col)
        ws_p.cell(row=_matrix_hdr + 1, column=_col, value=_ent["ticker"]).font = Font(
            bold=True, size=9
        )
        ws_p.cell(
            row=_matrix_vis,
            column=_col,
            value=_xl_grafici_vis_col(f"{_cl}${_matrix_hdr + 1}"),
        )
    _lab_rows = (
        ("% curva vs T−60", "2a · G1"),
        ("% storico", "2b · G2"),
        ("% modello", "2c · G3"),
        ("Combinato", "G4 · curva / stor. / mod."),
        ("Var. % 1g·1M·3M·6M", "G5 e/o menu A/B"),
    )
    for _ii, (_lbl, _sec) in enumerate(_lab_rows):
        ws_p.cell(row=_matrix_lab0 + _ii, column=1, value=_lbl).font = Font(size=9)
        ws_p.cell(
            row=_matrix_lab0 + _ii,
            column=2,
            value=f"→ {_sec}",
        ).font = Font(size=8, italic=True, color="666666")

    _row = _matrix_lab0 + len(_lab_rows) + 2
    _ws_sim = _simulation_ws(wb)
    _sim_var = _grafici_sim_var_context(wb, company_curves)
    _axis_var = 0
    _data_var = 0
    _n_var = 0
    if _sim_var:
        _sim_ref, _sim_first, _sim_last, _c_var_lo = _sim_var
        _axis_var, _data_var, _n_var = _write_grafici_variation_block(
            ws_p,
            _row,
            sim_ref=_sim_ref,
            sim_first=_sim_first,
            sim_last=_sim_last,
            c_var_lo=_c_var_lo,
            sel_data0=_sel["sel_data0"],
            sel_last=_sel["sel_last"],
            ws_sim=_ws_sim,
            company_curves=company_curves,
        )
        _row = _data_var + _n_var + 2
    else:
        print(
            "[Grafici] Var.%: foglio Simulation non trovato — grafico 5 omesso.",
            flush=True,
        )

    # ── Blocchi dati + grafici ─────────────────────────────────────────────

    def _pct_session_chart(
        title: str,
        anchor: str,
        data0: int,
        n_series: int,
        cat_row: int,
        *,
        pct_fields: tuple[str, ...],
        value_fmt: str | None = None,
    ) -> None:
        # Asse Y auto: Excel scala sui punti non-NA (G1–G4 / tipo attivi).
        _add_session_scatter_chart(
            ws_p,
            anchor,
            title,
            cat_row,
            data0,
            n_series,
            _n_std,
            pct_fmt=value_fmt if value_fmt is not None else pct_fmt,
        )

    _sel_kw = dict(
        offsets=offsets,
        pct_fmt=pct_fmt,
        sel=_sel,
        company_curves=company_curves,
        wb=wb,
        combo_reg_first=_reg_first,
        combo_reg_last=_reg_last,
        combo_reg_lab_col=_lab_col,
        combo_reg_val_col=_val_col,
        wide=wide,
    )
    _axis_c, _data_c, _n_c = _write_grafici_pct_chart_from_selector(
        ws_p,
            _row,
        block_title="1 · % curva (col. G1 tabella coorte)",
        include_col=_GRAF_SEL_COL_G1,
        wide_data0=_d0c,
        wide_last=_d1c,
        curve_field="pct_curva",
        **_sel_kw,
    )
    _row = _data_c + _n_c + 2
    _axis_s, _data_s, _n_s = _write_grafici_pct_chart_from_selector(
        ws_p,
            _row,
        block_title="2 · % storico (col. G2 tabella coorte)",
        include_col=_GRAF_SEL_COL_G2,
        wide_data0=_d0s,
        wide_last=_d1s,
        curve_field="pct_reale",
        **_sel_kw,
    )
    _row = _data_s + _n_s + 2
    _axis_m, _data_m, _n_m = _write_grafici_pct_chart_from_selector(
        ws_p,
        _row,
        block_title="3 · % modello (col. G3 tabella coorte)",
        include_col=_GRAF_SEL_COL_G3,
        wide_data0=_d0m,
        wide_last=_d1m,
        curve_field="pct_modello",
        **_sel_kw,
    )
    _row = _data_m + _n_m + 2
    _axis_cb, _data_cb, _n_cb = _write_grafici_pct_chart_from_selector(
        ws_p,
        _row,
        block_title="4 · Combinato (G4 + curva / stor. / mod.)",
        include_col=_GRAF_SEL_COL_G4,
        wide_data0=_d0c,
        wide_last=_d1c,
        use_combo_registry=True,
        **_sel_kw,
    )
    _row = _data_cb + _n_cb + 2

    _chart_anchor = _row + 1
        _gap = _CHART_STACK_GAP
    _pct_session_chart(
        "1 · % curva ricalibrata vs T−60",
        f"A{_chart_anchor}",
        _data_c,
        _n_c,
        _axis_c,
        pct_fields=("pct_curva",),
    )
    _pct_session_chart(
        "2 · % storico vs T−60",
            f"A{_chart_anchor + _gap}",
        _data_s,
        _n_s,
        _axis_s,
        pct_fields=("pct_reale",),
    )
    _pct_session_chart(
        "3 · % modello vs T−60",
            f"A{_chart_anchor + 2 * _gap}",
        _data_m,
        _n_m,
        _axis_m,
        pct_fields=("pct_modello",),
    )
    _pct_session_chart(
        "4 · Combinato — G4 (curve selezionate)",
            f"A{_chart_anchor + 3 * _gap}",
        _data_cb,
        _n_cb,
        _axis_cb,
        pct_fields=("pct_curva", "pct_reale", "pct_modello"),
    )
    if _n_var > 0 and _axis_var and _data_var:
        _add_variation_bar_chart(
            ws_p,
            f"{_gcl(_GRAF_VAR_CHART_COL)}{_axis_var}",
            "5 · Variazioni % — 1g · 1M · 3M · 6M",
            _axis_var,
            _data_var,
            _n_var,
        )

    try:
        ws_p.column_dimensions["A"].width = 28
        ws_p.column_dimensions["B"].width = 10
        ws_p.column_dimensions["C"].width = 14
        for _ci in range(4, 4 + _n_co):
            ws_p.column_dimensions[_gcl(_ci)].width = 11
        for _ci in range(3, 3 + _n_std):
            ws_p.column_dimensions[_gcl(_ci)].width = 11
        ws_p.column_dimensions[_gcl(_GRAF_LIST_COL_G)].hidden = True
        ws_p.column_dimensions[_gcl(_GRAF_COMBO_LIST_COL)].hidden = True
        for _hci in range(_GRAF_COMBO_REG_COL_LAB, _GRAF_COMBO_REG_COL0 + _n_std):
            ws_p.column_dimensions[_gcl(_hci)].hidden = True
        ws_p.column_dimensions[_gcl(_GRAF_SEL_COL_G4_MU)].width = 9
        ws_p.column_dimensions[_gcl(_GRAF_SEL_COL_COMBO_LBL)].width = 28
        for _hci in range(_GRAF_VAR_COL0, _GRAF_VAR_COL0 + _GRAF_VAR_N):
            ws_p.column_dimensions[_gcl(_hci)].width = 11
        ws_p.column_dimensions[_gcl(_GRAF_VAR_CHART_COL)].width = 2
        ws_p.freeze_panes = "A2"
    except Exception:
        pass

    _n_mu = len(_pred_entries) if _pred_entries else len(_REF_SPECS)
    print(
        f"[Grafici] 4 grafici · tabella coorte {len(company_curves)} soc."
        f" · serie dati G1–G4: {_n_c}/{_n_s}/{_n_m}/{_n_cb}"
        f" · μ riferimento: {_n_mu}"
        f"{f' + var ({_n_var})' if _n_var else ''}, anchor riga {_chart_anchor}.",
        flush=True,
    )
    if _n_c < 1 and _n_s < 1 and _n_m < 1:
        print(
            "[Grafici] ATTENZIONE: blocchi % curva/storico/modello senza serie — "
            "verificare spunte G1–G3 in tabella coorte e tabelle 2a–2c (righe ~520+).",
            flush=True,
        )
    elif _n_m < 1:
        print(
            "[Grafici] ATTENZIONE: grafico % modello senza serie — "
            "ricalcolare con --refresh o verificare model_dm* su Simulation.",
            flush=True,
        )
    if not _sn7_mu or not _sn7_ref_mu_series_has_signal(_sn7_mu.get("primary")):
        print(
            "[Grafici] ATTENZIONE: μ di riferimento assenti — "
            "eseguire refresh_predizione_guida.py poi Avvia_Grafici_Sheet.",
            flush=True,
        )


def refresh_grafici_from_simulation(
    wb,
    ws_sim,
    main_row_positions: list[int],
    *,
    sn7_ref_mu: dict | None = None,
    universe: dict | None = None,
    sim_pred_data: dict | None = None,
    financial_df=None,
) -> None:
    """Aggiorna il foglio «Grafici» da Simulation (senza foglio Simulation — grafici)."""
    _ = financial_df  # riservato al coordinator refresh (SEC K-8 / enrich futuro)
    if not main_row_positions:
        return
    _co = build_company_curves_from_simulation(
        wb,
        ws_sim,
        main_row_positions,
        universe=universe,
        sim_pred_data=sim_pred_data,
    )
    if not _co:
        return
    write_grafici_sheet(
        wb,
        company_curves=_co,
        wide=None,
        offsets=SIM_GRAFICI_CURVE_OFFSETS,
        sn7_ref_mu=sn7_ref_mu,
        universe=universe,
    )
    print(
        f"[Grafici] Foglio aggiornato: {len(_co)} società.",
        flush=True,
    )


# Alias legacy (orchestrator / script vecchi)
write_simulation_grafici_sheet = refresh_grafici_from_simulation
