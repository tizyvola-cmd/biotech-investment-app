"""
SuperNova Desktop — interfaccia Dear PyGui (sostituisce la versione Tkinter).

Avvio::

    .venv\\Scripts\\python.exe supernova_dpg.py
    scripts\\Avvia_SuperNova.bat
"""
from __future__ import annotations

import json
import os
import queue
import shutil
import subprocess
import threading
import time
from collections import deque
from datetime import datetime
from pathlib import Path

import dearpygui.dearpygui as dpg

import orchestrator_io_paths as _paths
from dpg_charts_panel import ChartsPanel, apply_supernova_theme
from orchestrator_io_paths import (
    DATA_DIR,
    FALLBACK_XLSX,
    FINAL_XLSX,
    LAST_ORCH_LOG,
    ORCHESTRATOR_SCRIPT,
    PYTHON_VENV_EXE,
)

ROOT = Path(_paths.project_root())
PYTHON = Path(PYTHON_VENV_EXE)
ORCHESTRATOR = Path(ORCHESTRATOR_SCRIPT)
DEFAULT_JSON = Path(_paths.DATA_DIR) / "model_calibration_state.json"

_BG = (28, 28, 28, 255)
_ACCENT = (0, 200, 150, 255)
_MUTED = (156, 163, 175, 255)


def _now() -> str:
    return datetime.now().strftime("%H:%M:%S")


def _cursor_exes() -> list[str]:
    seen: set[str] = set()
    out: list[str] = []

    def add(p: str | None) -> None:
        if not p:
            return
        s = str(p).strip().strip('"')
        if s and os.path.isfile(s) and s not in seen:
            seen.add(s)
            out.append(s)

    add(os.environ.get("CURSOR_EXE"))
    w = shutil.which("cursor")
    if w:
        add(w)
    la = os.environ.get("LOCALAPPDATA", "")
    if la:
        add(str(Path(la) / "Programs" / "cursor" / "Cursor.exe"))
    return out


class SuperNovaDPG:
    def __init__(self) -> None:
        self._ui_q: deque = deque()
        self._log_q: queue.Queue[str] = queue.Queue()
        self._run_lock = threading.Lock()
        self._proc_lock = threading.Lock()
        self._proc: subprocess.Popen | None = None
        self._run_thread: threading.Thread | None = None
        self._stop_requested = threading.Event()
        self._run_started: float | None = None
        self._run_profile = "quick"
        self._skip_sec_k8 = False
        self._orch_perf = False
        self._auto_enabled = False
        self._auto_minutes = 60
        self._watch_disk = True
        self._last_json_mt: float | None = None
        self._last_xlsx_mt: float | None = None
        self._next_auto_at: float | None = None
        self._charts = ChartsPanel(status_tag="charts_status")
        self._wb_sheets: list[str] = []
        self._wb_rows: list[list[str]] = []

    def _post(self, fn) -> None:
        self._ui_q.append(fn)

    def _set_status(self, msg: str) -> None:
        if dpg.does_item_exist("status_bar"):
            dpg.set_value("status_bar", f"[{_now()}] {msg}")

    def _pick_xlsx(self) -> Path | None:
        cands: list[Path] = []
        if FINAL_XLSX.exists():
            cands.append(FINAL_XLSX)
        if FALLBACK_XLSX.exists():
            cands.append(FALLBACK_XLSX)
        cands.extend(
            sorted(
                DATA_DIR.glob("biotech_orchestrated_output_*.xlsx"),
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )[:5]
        )
        return max(cands, key=lambda p: p.stat().st_mtime) if cands else None

    def _ui_tick(self) -> None:
        while self._ui_q:
            try:
                self._ui_q.popleft()()
            except Exception:
                pass
        try:
            while True:
                chunk = self._log_q.get_nowait()
                self._append_log(chunk)
        except queue.Empty:
            pass
        if self._run_started is not None:
            e = int(time.monotonic() - self._run_started)
            dpg.set_value("run_timer", f"Run: {e // 60}m {e % 60}s")
        if self._auto_enabled and self._next_auto_at and time.monotonic() >= self._next_auto_at:
            self._next_auto_at = None
            if not (self._run_thread and self._run_thread.is_alive()):
                self.run_update(auto=True)
            self._arm_auto_timer()
        if self._watch_disk:
            self._maybe_reload_disk()
        dpg.set_frame_callback(dpg.get_total_frames() + 2, self._ui_tick)

    def _append_log(self, text: str) -> None:
        if not dpg.does_item_exist("orch_log"):
            return
        cur = dpg.get_value("orch_log") or ""
        cur += text
        if len(cur) > 100_000:
            cur = cur[-90_000:]
        dpg.set_value("orch_log", cur)

    def _build_ui(self) -> None:
        with dpg.window(tag="primary", label="SuperNova"):
            dpg.add_text("SuperNova · Data Lab", color=_ACCENT)
            dpg.add_text("", tag="status_bar", color=_MUTED)

            with dpg.tab_bar(tag="main_tabs"):
                with dpg.tab(label="Pipeline"):
                    self._build_pipeline_tab()
                with dpg.tab(label="Charts"):
                    with dpg.child_window(border=True, height=-220, tag="charts_host"):
                        self._charts.build("charts_host")
                    dpg.add_text("", tag="charts_status", color=_MUTED)
                with dpg.tab(label="Calibrazione"):
                    self._build_calib_tab()
                with dpg.tab(label="Workbook"):
                    self._build_workbook_tab()

        dpg.set_primary_window("primary", True)

    def _build_pipeline_tab(self) -> None:
        dpg.add_text("Profilo aggiornamento", color=_ACCENT)
        dpg.add_radio_button(
            items=["quick", "full", "skip_fetch"],
            default_value="quick",
            tag="run_profile",
            horizontal=True,
            callback=lambda s, v: setattr(self, "_run_profile", v),
        )
        dpg.add_text(
            "quick = --quick  |  full = pipeline completa  |  skip_fetch = ORCH_SKIP_FETCH=1",
            color=_MUTED,
        )
        with dpg.group(horizontal=True):
            dpg.add_checkbox(
                label="Salta SEC K-8",
                default_value=False,
                callback=lambda s, v: setattr(self, "_skip_sec_k8", v),
            )
            dpg.add_checkbox(
                label="ORCH_PERF",
                default_value=False,
                callback=lambda s, v: setattr(self, "_orch_perf", v),
            )
        with dpg.group(horizontal=True):
            dpg.add_button(label="Avvia aggiornamento", callback=lambda: self.run_update())
            dpg.add_button(label="Interrompi", tag="btn_stop", callback=self.stop_update)
            dpg.disable_item("btn_stop")
            dpg.add_button(label="Apri Excel", callback=self.open_excel)
            dpg.add_button(label="Cartella data", callback=self.open_data_folder)
            dpg.add_button(label="Log completo", callback=self.open_log_file)
            dpg.add_button(label="Cursor progetto", callback=self.open_cursor)
        dpg.add_text("", tag="run_timer", color=_MUTED)
        with dpg.group(horizontal=True):
            dpg.add_checkbox(
                label="Timer auto",
                default_value=False,
                callback=self._on_auto_toggle,
            )
            dpg.add_input_int(
                label="min",
                default_value=60,
                width=80,
                tag="auto_minutes",
                callback=lambda s, v: setattr(self, "_auto_minutes", max(1, int(v))),
            )
            dpg.add_checkbox(
                label="Watch JSON/Excel (30s)",
                default_value=True,
                callback=lambda s, v: setattr(self, "_watch_disk", v),
            )
        dpg.add_separator()
        dpg.add_text("Output orchestrator", color=_ACCENT)
        dpg.add_input_text(
            tag="orch_log",
            multiline=True,
            readonly=True,
            width=-1,
            height=280,
        )

    def _build_calib_tab(self) -> None:
        with dpg.group(horizontal=True):
            dpg.add_button(label="Ricarica JSON", callback=self.reload_calib)
            dpg.add_button(label="Apri cartella data", callback=self.open_data_folder)
        dpg.add_input_text(
            tag="calib_text",
            multiline=True,
            readonly=True,
            width=-1,
            height=-1,
        )

    def _build_workbook_tab(self) -> None:
        with dpg.group(horizontal=True):
            dpg.add_button(label="Ricarica workbook", callback=self.reload_workbook)
            dpg.add_button(label="Apri in Excel", callback=self.open_excel)
        dpg.add_combo(
            label="Foglio",
            items=["—"],
            tag="wb_sheet",
            width=280,
            callback=self._on_sheet_pick,
        )
        with dpg.child_window(height=-1, border=True):
            with dpg.table(
                tag="wb_table",
                header_row=True,
                resizable=True,
                borders_innerH=True,
                borders_outerH=True,
            ):
                pass

    def _on_auto_toggle(self, sender, app_data) -> None:
        self._auto_enabled = bool(app_data)
        if self._auto_enabled:
            self._arm_auto_timer()
            self._set_status(f"Timer ogni {self._auto_minutes} min.")
        else:
            self._next_auto_at = None
            self._set_status("Timer disattivato.")

    def _arm_auto_timer(self) -> None:
        if self._auto_enabled:
            self._next_auto_at = time.monotonic() + self._auto_minutes * 60

    def reload_calib(self) -> None:
        if not DEFAULT_JSON.is_file():
            dpg.set_value("calib_text", f"File assente:\n{DEFAULT_JSON}")
            return
        try:
            doc = json.loads(DEFAULT_JSON.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            dpg.set_value("calib_text", f"Errore: {exc}")
            return
        cur = doc.get("current") if isinstance(doc.get("current"), dict) else doc
        cal = cur.get("cal_factor") if isinstance(cur, dict) else {}
        lines = [
            f"File: {DEFAULT_JSON.name}",
            f"Modificato: {datetime.fromtimestamp(DEFAULT_JSON.stat().st_mtime)}",
            "",
            f"cal_factor v4_options: {cal.get('v4_options', '—') if isinstance(cal, dict) else '—'}",
            f"population: {cur.get('population_filter', '—') if isinstance(cur, dict) else '—'}",
            "",
            "Per grafici e curve usa la scheda Charts.",
        ]
        dpg.set_value("calib_text", "\n".join(lines))
        self._last_json_mt = DEFAULT_JSON.stat().st_mtime
        self._set_status("Calibrazione ricaricata.")

    def reload_workbook(self) -> None:
        self._set_status("Caricamento workbook…")
        threading.Thread(target=self._wb_worker, daemon=True).start()

    def _wb_worker(self) -> None:
        p = self._pick_xlsx()
        if p is None:
            self._post(lambda: self._set_status("Nessun workbook in data/."))
            return
        try:
            import openpyxl

            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            sheets = list(wb.sheetnames)
            sn = sheets[0] if sheets else None
            rows: list[list[str]] = []
            hdr: list[str] = []
            if sn:
                ws = wb[sn]
                for i, row in enumerate(ws.iter_rows(max_row=80, values_only=True)):
                    cells = ["" if c is None else str(c)[:80] for c in (row or [])[:24]]
                    if i == 0:
                        hdr = cells
                    else:
                        rows.append(cells)
            wb.close()
        except Exception as exc:
            self._post(lambda e=exc: self._set_status(f"Workbook: {e}"))
            return

        def apply() -> None:
            self._wb_sheets = sheets
            dpg.configure_item("wb_sheet", items=sheets or ["—"])
            if sheets:
                dpg.set_value("wb_sheet", sheets[0])
            self._fill_wb_table(hdr, rows)
            self._last_xlsx_mt = p.stat().st_mtime
            self._set_status(f"Workbook: {p.name} · {len(sheets)} fogli")

        self._post(apply)

    def _fill_wb_table(self, headers: list[str], rows: list[list[str]]) -> None:
        if dpg.does_item_exist("wb_table"):
            dpg.delete_item("wb_table", children_only=True)
        cols = headers or (rows[0] if rows else ["A"])
        for h in cols[:24]:
            dpg.add_table_column(label=str(h)[:20], parent="wb_table")
        for row in rows[:60]:
            with dpg.table_row(parent="wb_table"):
                for c in row[: len(cols)]:
                    dpg.add_text(str(c))

    def _on_sheet_pick(self, sender, app_data) -> None:
        sn = str(app_data or "")
        if not sn or sn == "—":
            return
        p = self._pick_xlsx()
        if not p:
            return
        threading.Thread(
            target=self._wb_sheet_worker, args=(p, sn), daemon=True
        ).start()

    def _wb_sheet_worker(self, path: Path, sheet: str) -> None:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet]
            hdr: list[str] = []
            rows: list[list[str]] = []
            for i, row in enumerate(ws.iter_rows(max_row=80, values_only=True)):
                cells = ["" if c is None else str(c)[:80] for c in (row or [])[:24]]
                if i == 0:
                    hdr = cells
                else:
                    rows.append(cells)
            wb.close()
        except Exception as exc:
            self._post(lambda e=exc: self._set_status(str(e)))
            return
        self._post(lambda: self._fill_wb_table(hdr, rows))

    def _maybe_reload_disk(self) -> None:
        now = time.monotonic()
        if not hasattr(self, "_last_watch_t"):
            self._last_watch_t = 0.0
        if now - self._last_watch_t < 28:
            return
        self._last_watch_t = now
        try:
            if DEFAULT_JSON.is_file():
                mt = DEFAULT_JSON.stat().st_mtime
                if self._last_json_mt and mt > self._last_json_mt + 0.5:
                    self.reload_calib()
                    self._charts.load_mode("simulation")
                self._last_json_mt = mt
        except OSError:
            pass
        p = self._pick_xlsx()
        if p:
            try:
                mt = p.stat().st_mtime
                if self._last_xlsx_mt and mt > self._last_xlsx_mt + 0.5:
                    self.reload_workbook()
                    self._charts.load_mode("simulation")
                self._last_xlsx_mt = mt
            except OSError:
                pass

    def run_update(self, auto: bool = False) -> None:
        if not PYTHON.is_file():
            self._set_status(f"Python non trovato: {PYTHON}")
            return
        if not ORCHESTRATOR.is_file():
            self._set_status(f"Orchestrator non trovato: {ORCHESTRATOR}")
            return
        with self._run_lock:
            if self._run_thread and self._run_thread.is_alive():
                self._set_status("Run già in corso.")
                return
            self._stop_requested.clear()
            self._run_started = time.monotonic()
            dpg.enable_item("btn_stop")
            self._set_status(
                f"Avvio {'auto' if auto else 'manuale'} ({self._run_profile})…"
            )
            self._run_thread = threading.Thread(target=self._run_worker, daemon=True)
            self._run_thread.start()

    def stop_update(self) -> None:
        self._stop_requested.set()
        with self._proc_lock:
            p = self._proc
        if p and p.poll() is None:
            try:
                p.terminate()
            except Exception:
                pass
        self._set_status("Interruzione inviata.")

    def _run_worker(self) -> None:
        cmd = [str(PYTHON), "-u", str(ORCHESTRATOR)]
        env = os.environ.copy()
        env["PYTHONUNBUFFERED"] = "1"
        env.setdefault("PYTHONIOENCODING", "utf-8")
        if self._run_profile == "quick":
            cmd.append("--quick")
        elif self._run_profile == "skip_fetch":
            env["ORCH_SKIP_FETCH"] = "1"
        if self._skip_sec_k8:
            env["ORCH_SKIP_SEC_K8"] = "1"
        if self._orch_perf:
            env["ORCH_PERF"] = "1"
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        rc = -1
        try:
            with open(LAST_ORCH_LOG, "w", encoding="utf-8", errors="replace") as logf:
                logf.write(f"# cmd={' '.join(cmd)}\n\n")
                proc = subprocess.Popen(
                    cmd,
                    cwd=str(ROOT),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    env=env,
                )
                with self._proc_lock:
                    self._proc = proc
                self._log_q.put("\n>>> Avviato\n")
                assert proc.stdout
                for line in proc.stdout:
                    logf.write(line)
                    logf.flush()
                    self._log_q.put(line)
                rc = proc.wait()
        except Exception as exc:
            self._log_q.put(f"\nERRORE: {exc}\n")
            rc = -1
        finally:
            with self._proc_lock:
                self._proc = None
            self._run_started = None
            self._post(lambda: dpg.disable_item("btn_stop"))
            self._post(lambda: dpg.set_value("run_timer", ""))

        def done() -> None:
            if rc == 0:
                self._set_status("Completato.")
                self.reload_calib()
                self.reload_workbook()
                self._charts.load_mode("simulation")
            elif rc in (130, -15):
                self._set_status("Interrotto.")
            else:
                self._set_status(f"Uscita codice {rc}")

        self._post(done)

    def open_excel(self) -> None:
        p = self._pick_xlsx()
        if not p:
            self._set_status("Nessun Excel in data/.")
            return
        try:
            os.startfile(str(p))  # type: ignore[attr-defined]
        except Exception as exc:
            self._set_status(str(exc))

    def open_data_folder(self) -> None:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(str(DATA_DIR))  # type: ignore[attr-defined]
        except Exception as exc:
            self._set_status(str(exc))

    def open_log_file(self) -> None:
        if LAST_ORCH_LOG.is_file():
            os.startfile(str(LAST_ORCH_LOG))  # type: ignore[attr-defined]

    def open_cursor(self) -> None:
        for exe in _cursor_exes():
            try:
                subprocess.Popen([exe, str(ROOT)], cwd=str(ROOT))
                self._set_status("Cursor avviato.")
                return
            except Exception:
                continue
        self._set_status("Cursor non trovato (CURSOR_EXE / PATH).")

    def run(self) -> None:
        os.chdir(ROOT)
        dpg.create_context()
        apply_supernova_theme()
        self._build_ui()
        dpg.create_viewport(
            title="SuperNova · Dear PyGui",
            width=1320,
            height=880,
            clear_color=(0.11, 0.11, 0.11, 1.0),
        )
        dpg.setup_dearpygui()
        dpg.show_viewport()
        dpg.set_frame_callback(8, self._ui_tick)
        self.reload_calib()
        self._charts.load_mode("simulation")
        dpg.start_dearpygui()
        dpg.destroy_context()


def main() -> None:
    SuperNovaDPG().run()


if __name__ == "__main__":
    main()
