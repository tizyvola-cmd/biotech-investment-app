"""Pannello lettura fogli Excel — layout «data lab» (sidebar fogli + griglia) o classico."""
from __future__ import annotations

import os
import threading
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk
from typing import Literal

import tkinter as tk

import openpyxl

from supernova_theme import DATALAB_SIDEBAR, NAVY, ORANGE, TEXT, TEXT_MUTED

# Limite pratico Treeview Tk su Windows
_MAX_COLS = 40
_MAX_ROWS = 8000


def _cell_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
    s = str(v)
    if len(s) > 500:
        return s[:497] + "…"
    return s


class WorkbookPanel(ttk.Frame):
    """
    Elenco fogli + griglia (Treeview) da ``biotech_orchestrated_output.xlsx``.

    ``layout="datalab"``: ``PanedWindow`` orizzontale con sidebar (lista fogli) e area griglia.
    ``layout="classic"``: toolbar con Combobox fogli (comportamento precedente).
    """

    def __init__(
        self,
        master: tk.Misc,
        *,
        get_xlsx_path: Callable[[], Path | None],
        on_open_excel: Callable[[], None] | None = None,
        layout: Literal["datalab", "classic"] = "datalab",
    ) -> None:
        super().__init__(master)
        self._get_xlsx_path = get_xlsx_path
        self._on_open_excel = on_open_excel
        self._layout = layout
        self._loading = False
        self._row_buffer: list[tuple[str, ...]] = []
        self._headers: tuple[str, ...] = ()
        self._header_display: list[str] = []
        self._filter_norm = ""
        self._sheet_var = tk.StringVar()

        self._sheet_combo: ttk.Combobox | None = None
        self._sheet_list: tk.Listbox | None = None

        if layout == "datalab":
            self._build_datalab()
        else:
            self._build_classic()

        self.after(200, self.reload_paths_only)

    def _build_datalab(self) -> None:
        # Stesso vincolo di ``supernova_desk``: niente highlightthickness su PanedWindow (Tcl 9+).
        pw = tk.PanedWindow(
            self,
            orient=tk.HORIZONTAL,
            sashwidth=5,
            bg=NAVY,
            bd=0,
        )
        pw.pack(fill="both", expand=True)

        left = tk.Frame(pw, bg=DATALAB_SIDEBAR, width=220)
        tk.Label(
            left,
            text="FOGLI",
            font=("Segoe UI", 9, "bold"),
            fg=TEXT_MUTED,
            bg=DATALAB_SIDEBAR,
        ).pack(anchor="w", padx=10, pady=(10, 4))

        self._sheet_list = tk.Listbox(
            left,
            bg=DATALAB_SIDEBAR,
            fg=TEXT,
            selectbackground=ORANGE,
            selectforeground="#ffffff",
            activestyle="none",
            highlightthickness=0,
            bd=0,
            font=("Segoe UI", 10),
            exportselection=0,
        )
        self._sheet_list.pack(fill="both", expand=True, padx=6, pady=(0, 6))
        self._sheet_list.bind("<<ListboxSelect>>", self._on_sheet_list_select)

        bf = tk.Frame(left, bg=DATALAB_SIDEBAR)
        bf.pack(fill="x", padx=8, pady=(0, 10))
        ttk.Button(bf, text="Ricarica", command=self.reload_paths_only).pack(fill="x", pady=2)
        ttk.Button(bf, text="Apri in Excel", command=self._open_excel).pack(fill="x", pady=2)

        right = ttk.Frame(pw)
        pw.add(left, minsize=168, stretch="never")
        pw.add(right, stretch="always", minsize=320)

        self._build_grid_area(right)

    def _build_classic(self) -> None:
        self._build_grid_area(self)

    def _build_grid_area(self, parent: ttk.Frame) -> None:
        bar = ttk.Frame(parent)
        bar.pack(fill="x", padx=8, pady=(8, 4))

        ttk.Button(bar, text="Ricarica fogli", command=self.reload_paths_only).pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(bar, text="Apri griglia in Excel", command=self._open_excel).pack(
            side="left", padx=(0, 8)
        )

        if self._layout == "classic":
            ttk.Label(bar, text="Foglio:").pack(side="left", padx=(12, 4))
            self._sheet_combo = ttk.Combobox(
                bar,
                textvariable=self._sheet_var,
                width=42,
                state="readonly",
                values=(),
            )
            self._sheet_combo.pack(side="left", padx=(0, 12))
            self._sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self._load_current_sheet())

        ttk.Label(bar, text="Filtro righe:").pack(side="left", padx=(12 if self._layout == "classic" else 0, 4))
        self._filter_var = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self._filter_var, width=32)
        ent.pack(side="left", padx=(0, 6))
        ent.bind("<Return>", lambda _e: self._apply_filter())
        ttk.Button(bar, text="Applica", command=self._apply_filter).pack(side="left")

        self._meta_var = tk.StringVar(value="Nessun workbook in data/.")
        ttk.Label(parent, textvariable=self._meta_var, wraplength=1100).pack(
            anchor="w", padx=10, pady=(0, 4)
        )

        tree_frame = ttk.Frame(parent)
        tree_frame.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        scroll_y = ttk.Scrollbar(tree_frame, orient="vertical")
        scroll_x = ttk.Scrollbar(tree_frame, orient="horizontal")
        self._tree = ttk.Treeview(
            tree_frame,
            show="headings",
            yscrollcommand=scroll_y.set,
            xscrollcommand=scroll_x.set,
        )
        scroll_y.config(command=self._tree.yview)
        scroll_x.config(command=self._tree.xview)
        self._tree.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

        self._foot_var = tk.StringVar(value="")
        ttk.Label(parent, textvariable=self._foot_var, font=("Segoe UI", 8)).pack(
            anchor="w", padx=10, pady=(0, 6)
        )

    def _on_sheet_list_select(self, _event: object | None = None) -> None:
        if self._sheet_list is None:
            return
        sel = self._sheet_list.curselection()
        if not sel:
            return
        name = self._sheet_list.get(sel[0])
        if self._sheet_var.get() == name and not self._loading:
            return
        self._sheet_var.set(name)
        self._load_current_sheet()

    def _open_excel(self) -> None:
        if self._on_open_excel:
            self._on_open_excel()
            return
        p = self._get_xlsx_path()
        if p is None:
            messagebox.showinfo("SuperNova", "Nessun workbook in data/.")
            return
        try:
            os.startfile(str(p))  # type: ignore[attr-defined]
        except Exception as exc:
            messagebox.showerror("SuperNova", str(exc))

    def reload_paths_only(self) -> None:
        """Aggiorna elenco fogli e metadati senza bloccare la UI."""
        p = self._get_xlsx_path()
        if p is None:
            self._meta_var.set("Nessun file Excel trovato (data/). Esegui la Pipeline.")
            if self._sheet_combo is not None:
                self._sheet_combo.configure(values=())
            if self._sheet_list is not None:
                self._sheet_list.delete(0, tk.END)
            self._sheet_var.set("")
            self._clear_tree()
            return
        try:
            mtime = datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
        except Exception:
            mtime = "?"
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
            names = list(wb.sheetnames)
            wb.close()
        except Exception as exc:
            self._meta_var.set(f"Errore lettura {p.name}: {exc}")
            return

        self._meta_var.set(f"{p.name}  ·  {len(names)} fogli  ·  aggiornato {mtime}")
        if self._sheet_combo is not None:
            self._sheet_combo.configure(values=names)
        if self._sheet_list is not None:
            self._sheet_list.delete(0, tk.END)
            for n in names:
                self._sheet_list.insert(tk.END, n)

        cur = self._sheet_var.get()
        if cur not in names and names:
            self._sheet_var.set(names[0])
            if self._sheet_list is not None:
                self._sheet_list.selection_clear(0, tk.END)
                self._sheet_list.selection_set(0)
                self._sheet_list.see(0)
        elif not names:
            self._sheet_var.set("")
        elif self._sheet_list is not None and cur in names:
            idx = names.index(cur)
            self._sheet_list.selection_clear(0, tk.END)
            self._sheet_list.selection_set(idx)
            self._sheet_list.see(idx)

        self._load_current_sheet()

    def reload_after_pipeline(self) -> None:
        """Chiamare al termine di un run orchestrator con successo."""
        self.after(0, self.reload_paths_only)

    def _clear_tree(self) -> None:
        self._tree.delete(*self._tree.get_children())
        self._tree["columns"] = ()

    def _load_current_sheet(self) -> None:
        if self._loading:
            return
        p = self._get_xlsx_path()
        sn = self._sheet_var.get().strip()
        if p is None or not sn:
            return
        self._loading = True
        self._foot_var.set("Caricamento in corso…")
        threading.Thread(
            target=self._thread_load,
            args=(str(p), sn),
            daemon=True,
        ).start()

    def _thread_load(self, path: str, sheet_name: str) -> None:
        err: str | None = None
        headers: tuple[str, ...] = ()
        header_display: list[str] = []
        rows: list[tuple[str, ...]] = []
        truncated_rows = False
        truncated_cols = False
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[sheet_name]
                it = ws.iter_rows(values_only=True)
                header_row = next(it, None)
                if header_row is None:
                    headers = ()
                    header_display = []
                else:
                    raw_h = [_cell_str(x) or f"col_{i}" for i, x in enumerate(header_row)]
                    if len(raw_h) > _MAX_COLS:
                        raw_h = raw_h[:_MAX_COLS]
                        truncated_cols = True
                    header_display = list(raw_h)
                    headers = tuple(f"c{i}" for i in range(len(header_display)))
                nh = len(headers)
                for idx, row in enumerate(it):
                    if idx >= _MAX_ROWS:
                        truncated_rows = True
                        break
                    cells = list(row) if row else []
                    if len(cells) < nh:
                        cells.extend([None] * (nh - len(cells)))
                    elif len(cells) > nh:
                        cells = cells[:nh]
                    rows.append(tuple(_cell_str(c) for c in cells))
            finally:
                wb.close()
        except Exception as exc:
            err = str(exc)

        def _done(
            _err: str | None = err,
            _headers: tuple[str, ...] = headers,
            _labels: list[str] = header_display,
            _rows: list[tuple[str, ...]] = rows,
            _tr: bool = truncated_rows,
            _tc: bool = truncated_cols,
            _sn: str = sheet_name,
        ) -> None:
            self._loading = False
            if _err:
                self._foot_var.set(f"Errore: {_err}")
                messagebox.showerror("SuperNova", f"Foglio «{_sn}»:\n{_err}")
                return
            self._headers = _headers
            self._header_display = list(_labels)
            self._row_buffer = _rows
            msg = f"{len(_rows)} righe"
            if _tr:
                msg += f" (limite visualizzazione {_MAX_ROWS}; file completo in Excel)"
            if _tc:
                msg += f" · prime {_MAX_COLS} colonne"
            self._foot_var.set(msg)
            self._apply_filter()

        self.after(0, _done)

    def _apply_filter(self) -> None:
        self._filter_norm = (self._filter_var.get() or "").strip().lower()
        self._populate_tree()

    def _populate_tree(self) -> None:
        self._tree.delete(*self._tree.get_children())
        headers = self._headers
        if not headers:
            return
        labels = self._header_display or [f"Col {i}" for i in range(len(headers))]
        self._tree["columns"] = list(headers)
        for i, cid in enumerate(headers):
            lbl = labels[i] if i < len(labels) else cid
            self._tree.heading(cid, text=lbl[:48] + ("…" if len(lbl) > 48 else ""))
            self._tree.column(cid, width=min(220, max(72, len(lbl) * 7)), stretch=True)

        fn = self._filter_norm
        shown = 0
        for idx, row in enumerate(self._row_buffer):
            if fn and not any(fn in (c or "").lower() for c in row):
                continue
            self._tree.insert("", "end", iid=str(idx), values=row)
            shown += 1
            if shown > 12000:
                break
        if fn:
            self._foot_var.set(
                f"Filtro «{fn}»: {shown} righe mostrate (su {len(self._row_buffer)} caricate)."
            )
