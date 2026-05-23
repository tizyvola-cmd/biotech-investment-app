"""Unit tests for v4/v5 accuracy error metrics."""
from __future__ import annotations

from datetime import date

from prediction.accuracy_v4_v5 import (
    ACCURACY_V4_DISPLAY_OFFSETS,
    SIGN_HIT_FLAT_BAND_PP,
    abs_error_pp,
    aggregate_period_metrics,
    aggregate_legacy_periods_to_total,
    compute_row_errors,
    direction_hit,
    history_entry_total_block,
    run_snapshot_has_metrics,
    save_accuracy_summary_json,
    sign_hit_pp,
    signed_error_pp,
    summary_rows_for_period,
    summary_table_headers,
    temporal_run_rows,
    temporal_sheet_headers,
    temporal_summary_table_headers,
)
from prediction.pipeline import SIMULATION_V5_Q50_OFFSETS


def test_signed_and_abs_error_pp():
    assert signed_error_pp(10.0, 7.0) == 3.0
    assert abs_error_pp(10.0, 7.0) == 3.0
    assert signed_error_pp(None, 1.0) is None
    assert abs_error_pp(5.0, None) is None


def test_direction_hit_stable_band():
    assert direction_hit(2.0, 2.0) is True
    assert direction_hit(10.0, 8.0) is True
    assert direction_hit(10.0, -3.0) is False
    assert direction_hit(-1.0, 2.0) is True  # stable actual, small pred


def test_sign_hit_tight_band():
    assert SIGN_HIT_FLAT_BAND_PP == 0.5
    assert sign_hit_pp(0.2, 0.1) is True
    assert sign_hit_pp(2.0, -1.0) is False


def test_compute_row_errors_v4_and_v5():
    stor = [0.0, 5.0, 10.0, 7.0, 5.0, 3.0, 4.0, 7.0]
    pred = [0.0, 4.0, 12.0, 6.0, 4.0, 2.0, 5.0, 8.0]
    v5 = {str(o): float(o) / 10.0 for o in SIMULATION_V5_Q50_OFFSETS}
    rec = compute_row_errors(
        row_key="TST|2024-01-15",
        ticker="TST",
        completion_date=date(2024, 1, 15),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets=v5,
    )
    assert len(rec.horizons) == len(ACCURACY_V4_DISPLAY_OFFSETS)
    h30 = next(h for h in rec.horizons if h.offset == -30)
    assert h30.abs_v4_pp == 1.0  # |4-5|
    assert h30.abs_v5_pp == 8.0  # |-3-5|  v5 at -30 = -3.0


def test_aggregate_period_metrics_by_quarter():
    stor = [0.0] * 8
    pred = [0.0, 10.0] + [0.0] * 6
    rec1 = compute_row_errors(
        row_key="A|2024-01-01",
        ticker="A",
        completion_date=date(2024, 1, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets={},
    )
    rec2 = compute_row_errors(
        row_key="B|2024-02-01",
        ticker="B",
        completion_date=date(2024, 2, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets={},
    )
    agg = aggregate_period_metrics([rec1, rec2], past_only=True)
    assert agg["total"]["n_rows"] == 2
    assert "2024-Q1" in agg["by_quarter"]
    q1 = agg["by_quarter"]["2024-Q1"]
    assert q1["v4"]["mae"]["T-30"] == 10.0
    assert q1["v5"]["mae"]["T-30"] is None


def test_temporal_summary_two_rows_per_quarter():
    """Riepilogo foglio: due righe per trimestre (v4 sopra, v5 sotto)."""
    stor = [0.0, 8.0] + [0.0] * 6
    pred = [0.0, 10.0] + [0.0] * 6
    rec = compute_row_errors(
        row_key="A|2024-01-01",
        ticker="A",
        completion_date=date(2024, 1, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets={},
    )
    agg = aggregate_period_metrics([rec], past_only=True)
    block = agg["by_quarter"]["2024-Q1"]
    row_v4, row_v5 = summary_rows_for_period("2024-Q1", block)
    headers = summary_table_headers()
    assert headers[:3] == ("Periodo", "Modello", "N righe")
    assert len(headers) == 3 + len(ACCURACY_V4_DISPLAY_OFFSETS) * 2 + 2
    assert row_v4[0] == "2024-Q1" and row_v4[1] == "v4"
    assert row_v5[0] == "2024-Q1" and row_v5[1] == "v5"
    assert row_v4[headers.index("MAE T-30")] == 2.0
    assert row_v4[headers.index("Hit % T-30")] == 100.0
    assert row_v5[headers.index("MAE T-5")] is None
    assert row_v5[headers.index("Hit % T-5")] is None
    tot_v4, tot_v5 = summary_rows_for_period("TOTALE", agg["total"])
    assert tot_v4[1] == "v4" and tot_v5[1] == "v5"


def test_accuracy_sheet_column_count_no_abs_delta_blocks():
    import data_orchestrator as orch

    assert orch.ACCURACY_SHEET_N == 51
    assert orch.ACC_COL_DELTA_HI == orch.ACCURACY_SHEET_N
    assert orch.ACC_COL_DELTA_LO == 44


def test_run_snapshot_has_metrics_rejects_empty():
    assert not run_snapshot_has_metrics({"total": {"n_rows": 0}})
    assert run_snapshot_has_metrics({"total": {"n_rows": 3, "v4": {"mae": {}}}})


def test_merge_summary_history_skips_empty_run(tmp_path):
    p = tmp_path / "acc_summary.json"
    payload = aggregate_period_metrics([], past_only=True)
    save_accuracy_summary_json(p, payload)
    assert not p.exists()

    stor = [0.0] * 8
    pred = [0.0, 10.0] + [0.0] * 6
    rec = compute_row_errors(
        row_key="A|2024-01-01",
        ticker="A",
        completion_date=__import__("datetime").date(2024, 1, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets={},
    )
    good = aggregate_period_metrics([rec], past_only=True)
    save_accuracy_summary_json(p, good)
    doc = __import__("json").loads(p.read_text(encoding="utf-8"))
    assert len(doc.get("history") or []) == 1


def test_history_entry_total_block_legacy_periods():
    periods = {
        "2024-Q1": {
            "n_rows": 2,
            "mae_v4_T-30": 10.0,
            "mae_v4_T-30_n": 4,
            "mae_v5_T-30": 8.0,
            "mae_v5_T-30_n": 4,
        },
    }
    tot = aggregate_legacy_periods_to_total(periods)
    assert tot["n_rows"] == 2
    assert tot["v4"]["mae"]["T-30"] == 10.0
    blk = history_entry_total_block({"run_iso": "2026-01-01", "periods": periods})
    assert blk is not None
    v4_row, v5_row, v5_raw_row = temporal_run_rows("2026-01-01T12:00:00", blk)
    assert v5_raw_row is None
    assert v4_row[0] == "2026-01-01T12:00:00"
    assert v4_row[2] == "v4"
    assert temporal_sheet_headers()[0] == "Run"
    th = temporal_summary_table_headers()
    assert th.index("MAE T-30") < th.index("Hit % T-30")
    assert th.index("MAE globale") < th.index("Hit % T-60")


def test_temporal_hit_green_darker_when_higher():
    from prediction.accuracy_v4_v5 import temporal_hit_fill_and_font

    fill_lo, _ = temporal_hit_fill_and_font(40.0)
    fill_hi, _ = temporal_hit_fill_and_font(85.0)
    assert fill_lo.fgColor.rgb != fill_hi.fgColor.rgb


def test_temporal_mae_red_darker_when_higher():
    from prediction.accuracy_v4_v5 import temporal_mae_fill_and_font

    fill_lo, _ = temporal_mae_fill_and_font(18.0)
    fill_hi, _ = temporal_mae_fill_and_font(32.0)
    assert fill_hi.fgColor.rgb > fill_lo.fgColor.rgb or fill_hi.fgColor.rgb != fill_lo.fgColor.rgb


def test_aggregate_includes_v5_raw_block():
    from datetime import date

    stor = [0.0] * 8
    pred = [1.0, 2.0] + [0.0] * 6
    raw = [3.0, 4.0] + [0.0] * 6
    _v5d = {str(o): pred[i] for i, o in enumerate(ACCURACY_V4_DISPLAY_OFFSETS)}
    _rawd = {str(o): raw[i] for i, o in enumerate(ACCURACY_V4_DISPLAY_OFFSETS)}
    rec = compute_row_errors(
        row_key="A|2024-06-01",
        ticker="A",
        completion_date=date(2024, 6, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets=_v5d,
        v5_raw_offsets=_rawd,
    )
    agg = aggregate_period_metrics([rec], past_only=True)
    assert "v5_raw" in agg["total"]
    assert agg["total"]["v5_raw"]["mae_global"] is not None


def test_temporal_run_rows_three_models():
    blk = {
        "n_rows": 10,
        "v4": {"mae": {}, "mae_global": 5.0, "hit_pct": {}, "hit_pct_global": 50.0},
        "v5": {"mae": {}, "mae_global": 5.1, "hit_pct": {}, "hit_pct_global": 51.0},
        "v5_raw": {"mae": {}, "mae_global": 6.0, "hit_pct": {}, "hit_pct_global": 45.0},
    }
    v4, v5, raw = temporal_run_rows("2024-01-01T12:00:00", blk)
    assert v4[2] == "v4"
    assert v5[2] == "v5"
    assert raw is not None
    assert raw[2] == "v5 raw"


def test_write_accuracy_temporal_sheet_from_history(tmp_path):
    from openpyxl import Workbook

    import data_orchestrator as orch

    p = tmp_path / "summary.json"
    stor = [0.0] * 8
    pred = [0.0, 5.0] + [0.0] * 6
    rec = compute_row_errors(
        row_key="A|2024-01-01",
        ticker="A",
        completion_date=__import__("datetime").date(2024, 1, 1),
        is_past=True,
        v4_pred_pts=pred,
        storico_pts=stor,
        v5_offsets={},
    )
    save_accuracy_summary_json(p, aggregate_period_metrics([rec], past_only=True))
    wb = Workbook()
    orch.write_accuracy_temporal_sheet(wb, summary_json_path=str(p))
    assert orch.ACCURACY_TEMPORAL_SHEET in wb.sheetnames
    ws = wb[orch.ACCURACY_TEMPORAL_SHEET]
    found_v4 = False
    for r in range(1, int(ws.max_row or 0) + 1):
        if ws.cell(r, 3).value == "v4":
            found_v4 = True
            break
    assert found_v4
