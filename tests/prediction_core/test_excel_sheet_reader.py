"""Validazione workbook xlsx per excel_sheet_reader."""
from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

import pytest

from excel_sheet_reader import WorkbookReadError, _load_workbook_readonly, _preflight_xlsx


def test_preflight_rejects_empty_file(tmp_path: Path) -> None:
    p = tmp_path / "empty.xlsx"
    p.write_bytes(b"")
    with pytest.raises(WorkbookReadError) as exc:
        _preflight_xlsx(p)
    assert exc.value.status_code == 503
    assert exc.value.cause == "empty_file"


def test_preflight_rejects_non_zip(tmp_path: Path) -> None:
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"NOTAZIPFILE")
    with pytest.raises(WorkbookReadError) as exc:
        _preflight_xlsx(p)
    assert exc.value.cause == "not_xlsx_zip"


def test_load_workbook_retries_bad_zip(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Primo BadZipFile → attesa 0.5s → secondo tentativo OK."""
    p = tmp_path / "book.xlsx"
    p.write_bytes(b"PK\x03\x04fake")

    calls = {"n": 0}

    class FakeWb:
        sheetnames = ["Simulation"]

        def close(self) -> None:
            pass

    def fake_load(path, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            raise BadZipFile("not zip")
        return FakeWb()

    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", fake_load)
    monkeypatch.setattr("excel_sheet_reader.time.sleep", lambda _: None)

    wb = _load_workbook_readonly(p)
    assert calls["n"] == 2
    wb.close()


def test_sheet_api_returns_503_on_bad_zip(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient

    import supernova_config as sn_cfg
    import supernova_api
    import importlib

    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"PK\x03\x04truncated")
    monkeypatch.setattr("excel_sheet_reader.FINAL_XLSX", str(bad))

    def boom(*args, **kwargs):
        raise BadZipFile("corrupt")

    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", boom)
    monkeypatch.setattr("excel_sheet_reader.time.sleep", lambda _: None)

    sn_cfg.reset_supernova_config()
    importlib.reload(supernova_api)
    client = TestClient(supernova_api.build_app(sn_cfg.SupernovaConfig.from_env()))
    r = client.get("/api/sheets/simulation")
    assert r.status_code == 503
    detail = r.json()["detail"]
    assert "Chiudi Excel" in detail["message_it"]
    assert str(bad) in detail["path"]

    r2 = client.get("/api/sheets/financial")
    assert r2.status_code == 503
    assert "Chiudi Excel" in r2.json()["detail"]["message_it"]


def test_read_sec_k8_preserves_edgar_hyperlink(tmp_path: Path) -> None:
    """Colonna «Elenco 8‑K (SEC EDGAR)» esporta {text, href} da cell.hyperlink."""
    from openpyxl import Workbook

    from excel_sheet_reader import read_sec_k8_table

    wb = Workbook()
    ws = wb.active
    ws.title = "SEC K-8"
    ws.cell(row=3, column=1, value="Ticker")
    ws.cell(row=3, column=7, value="Elenco 8‑K\n(SEC EDGAR)")
    ws.cell(row=4, column=1, value="MRNA")
    link_cell = ws.cell(row=4, column=7, value="Elenco 8‑K (SEC)")
    link_cell.hyperlink = "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=1682852&type=8-K"
    xlsx = tmp_path / "sec_k8_links.xlsx"
    wb.save(xlsx)

    payload = read_sec_k8_table(xlsx_path=xlsx)
    assert payload.get("error") is None
    assert payload["row_count"] == 1
    col = "Elenco 8‑K\n(SEC EDGAR)"
    assert col in payload["columns"]
    cell_val = payload["rows"][0][col]
    assert cell_val == {
        "text": "Elenco 8‑K (SEC)",
        "href": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=1682852&type=8-K",
    }


def test_read_sec_k8_elenco_fallback_from_cik_plain_text(tmp_path: Path) -> None:
    """Senza hyperlink Excel: URL EDGAR browse da CIK nella stessa riga."""
    from openpyxl import Workbook

    from excel_sheet_reader import read_sec_k8_table

    wb = Workbook()
    ws = wb.active
    ws.title = "SEC K-8"
    ws.cell(row=3, column=1, value="Ticker")
    ws.cell(row=3, column=2, value="CIK (SEC)")
    ws.cell(row=3, column=7, value="Elenco 8‑K\n(SEC EDGAR)")
    ws.cell(row=4, column=1, value="MRNA")
    ws.cell(row=4, column=2, value="0001682852")
    ws.cell(row=4, column=7, value="Elenco 8‑K (SEC)")
    xlsx = tmp_path / "sec_k8_plain.xlsx"
    wb.save(xlsx)

    payload = read_sec_k8_table(xlsx_path=xlsx)
    col = "Elenco 8‑K\n(SEC EDGAR)"
    cell_val = payload["rows"][0][col]
    assert cell_val["text"] == "Elenco 8‑K (SEC)"
    assert "CIK=1682852" in cell_val["href"]
    assert cell_val["href"].startswith("https://www.sec.gov/cgi-bin/browse-edgar")


def test_read_simulation_preserves_study_hyperlink(tmp_path: Path) -> None:
    """Colonna «Link studio»: hyperlink Excel + fallback CT.gov da NCT."""
    from openpyxl import Workbook

    from data_orchestrator import SIMULATION_36_HEADERS
    from excel_sheet_reader import read_simulation_table

    wb = Workbook()
    ws = wb.active
    ws.title = "Simulation"
    link_ci = SIMULATION_36_HEADERS.index("Link studio") + 1
    nct_ci = SIMULATION_36_HEADERS.index("NCT") + 1
    ws.cell(row=4, column=1, value="MRNA")
    link_cell = ws.cell(row=4, column=link_ci, value="Studio CT.gov")
    link_cell.hyperlink = "https://clinicaltrials.gov/study/NCT00000001"
    ws.cell(row=4, column=nct_ci, value="NCT00000001")
    xlsx = tmp_path / "sim_links.xlsx"
    wb.save(xlsx)

    payload = read_simulation_table(xlsx_path=xlsx)
    assert payload.get("error") is None
    assert payload["row_count"] == 1
    link_val = payload["rows"][0]["Link studio"]
    assert link_val == {
        "text": "Studio CT.gov",
        "href": "https://clinicaltrials.gov/study/NCT00000001",
    }
    nct_val = payload["rows"][0]["NCT"]
    assert nct_val["href"] == "https://clinicaltrials.gov/study/NCT00000001"


def test_read_simulation_study_link_fallback_from_nct(tmp_path: Path) -> None:
    """Senza hyperlink Excel: «Link studio» da NCT nella stessa riga."""
    from openpyxl import Workbook

    from data_orchestrator import SIMULATION_36_HEADERS
    from excel_sheet_reader import read_simulation_table

    wb = Workbook()
    ws = wb.active
    ws.title = "Simulation"
    link_ci = SIMULATION_36_HEADERS.index("Link studio") + 1
    nct_ci = SIMULATION_36_HEADERS.index("NCT") + 1
    ws.cell(row=4, column=1, value="MRNA")
    ws.cell(row=4, column=link_ci, value="Apri studio")
    ws.cell(row=4, column=nct_ci, value="NCT00000002")
    xlsx = tmp_path / "sim_plain.xlsx"
    wb.save(xlsx)

    payload = read_simulation_table(xlsx_path=xlsx)
    link_val = payload["rows"][0]["Link studio"]
    assert link_val["text"] == "Apri studio"
    assert link_val["href"] == "https://clinicaltrials.gov/study/NCT00000002"


def test_load_clinical_nct_hyperlink_and_plain_fallback(tmp_path: Path) -> None:
    """Colonna ``nct_id``: hyperlink Excel oppure URL CT.gov da testo ``NCT…``."""
    from openpyxl import Workbook

    from excel_sheet_reader import _load_clinical_from_xlsx

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="ticker")
    ws.cell(row=1, column=2, value="nct_id")
    ws.cell(row=2, column=1, value="MRNA")
    plain = ws.cell(row=2, column=2, value="NCT00000001")
    plain.hyperlink = "https://clinicaltrials.gov/study/NCT00000001"
    ws.cell(row=3, column=1, value="VRTX")
    ws.cell(row=3, column=2, value="NCT04932343")
    xlsx = tmp_path / "clinical_nct.xlsx"
    wb.save(xlsx)

    rows, cols, err = _load_clinical_from_xlsx(xlsx)
    assert err is None
    assert "nct_id" in cols
    assert rows[0]["nct_id"] == {
        "text": "NCT00000001",
        "href": "https://clinicaltrials.gov/study/NCT00000001",
    }
    assert rows[1]["nct_id"] == {
        "text": "NCT04932343",
        "href": "https://clinicaltrials.gov/study/NCT04932343",
    }


def test_build_simulation_cd_catalog_picks_nearest_future_within_horizon() -> None:
    from datetime import date, timedelta

    from excel_sheet_reader import CLINICAL_MAX_CD_DAYS, _build_simulation_cd_catalog

    near = (date.today() + timedelta(days=14)).isoformat()
    far = (date.today() + timedelta(days=CLINICAL_MAX_CD_DAYS + 30)).isoformat()
    past = (date.today() - timedelta(days=10)).isoformat()

    catalog = _build_simulation_cd_catalog(
        [
            {"Ticker": "AAA", "Completion Date": far, "NCT": "NCT11111111"},
            {"Ticker": "AAA", "Completion Date": near, "NCT": "NCT22222222"},
            {"Ticker": "BBB", "Completion Date": past, "NCT": "NCT33333333"},
        ]
    )
    assert catalog["AAA"]["nct"] == "NCT22222222"
    assert catalog["AAA"]["completion_date"] == near
    assert "BBB" not in catalog


def test_clinical_row_matches_catalyst_nct_or_date() -> None:
    from excel_sheet_reader import CLINICAL_MAX_CD_DAYS, _clinical_row_matches_catalyst
    from datetime import date, timedelta

    near = (date.today() + timedelta(days=14)).isoformat()
    entry = {
        "nct": "NCT04932343",
        "completion_date": near,
        "days_to_cd": 14,
    }
    cd_cols = ["primary_completion_date", "completion_date"]
    assert _clinical_row_matches_catalyst(
        {"nct_id": "NCT04932343", "primary_completion_date": near, "completion_date": "2028-01-01"},
        entry,
        nct_col="nct_id",
        cd_cols=cd_cols,
    )
    # Solo completion_date finale diversa → no match (serve primary)
    assert not _clinical_row_matches_catalyst(
        {"nct_id": "NCT04932343", "primary_completion_date": "2020-01-01", "completion_date": near},
        entry,
        nct_col="nct_id",
        cd_cols=cd_cols,
    )
    # NCT ok ma primary CD diversa → no match
    assert not _clinical_row_matches_catalyst(
        {"nct_id": "NCT04932343", "primary_completion_date": "2020-01-01"},
        entry,
        nct_col="nct_id",
        cd_cols=cd_cols,
    )
    assert not _clinical_row_matches_catalyst(
        {"nct_id": "NCT00000000", "primary_completion_date": near},
        entry,
        nct_col="nct_id",
        cd_cols=cd_cols,
    )
    far = (date.today() + timedelta(days=CLINICAL_MAX_CD_DAYS + 5)).isoformat()
    assert not _clinical_row_matches_catalyst(
        {"nct_id": "NCT04932343", "estimated_completion_date": far},
        {**entry, "completion_date": far, "days_to_cd": CLINICAL_MAX_CD_DAYS + 5},
        nct_col="nct_id",
        cd_cols=["estimated_completion_date"],
    )
