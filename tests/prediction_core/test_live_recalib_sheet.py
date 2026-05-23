"""Curva ricalibrata live + colonne K-8 su Simulation / Accuracy."""
from __future__ import annotations

from datetime import date, timedelta

import data_orchestrator as orch


def test_simulation_headers_include_k8_columns():
    assert "N° K-8\n[CD−finestra SEC]" in orch.SIMULATION_36_HEADERS
    assert "Date K-8\n(+1/+2/+3)" in orch.SIMULATION_36_HEADERS
    assert orch.SIMULATION_36_N == len(orch.SIMULATION_36_HEADERS)
    assert orch.SIM_COL_K8_COUNT < orch.SIM_COL_AFFIDABILITA


def test_accuracy_sheet_meta_includes_k8():
    assert orch.ACCURACY_SHEET_META_N == 10
    assert orch.ACCURACY_SHEET_N == 51


def test_linear_interp_pct_at_offsets():
    from prediction.live_recalib_sheet import _linear_interp_pct_at_offsets

    knots = [(-60, 0.0), (-10, 10.0), (4, 20.0)]
    out = _linear_interp_pct_at_offsets((-60, -30, -10, 4, 7), knots)
    assert out[0] == 0.0
    assert out[2] == 10.0
    assert out[3] == 20.0
    assert out[1] == 6.0  # -30 tra -60 e -10 (interp. lineare)


def test_k8_interp_between_knots_not_flat_seq_future():
    """Tra nodi K-8 usa interp. lineare + modello; seq piatta futura non appiattisce."""
    from prediction.live_recalib_sheet import k8_interp_pct_points_in_knot_range

    cd = date(2026, 12, 1)
    today = date(2026, 5, 20)
    offsets = (-60, -30, -10, -3, 4, 7)
    model = [0.0, 2.0, 4.0, 6.0, 8.0, 10.0]
    flat_seq = [0.0, 3.0, 3.0, 3.0, 3.0, 3.0]
    knots = [(-60, 0.0), (-30, 12.0), (-10, 24.0)]
    out = k8_interp_pct_points_in_knot_range(
        knots,
        offsets=offsets,
        fallback_pts=model,
        prefer_base=flat_seq,
        completion_date=cd,
        today=today,
    )
    assert out[0] == 0.0
    assert out[1] == 12.0
    assert out[2] == 24.0
    assert out[3] == 6.0
    assert out[4] == 8.0
    assert out[5] == 10.0


def test_is_current_catalyst_row():
    from prediction.live_recalib_sheet import is_current_catalyst_row

    today = date(2026, 5, 20)
    assert is_current_catalyst_row(today, today) is True
    assert is_current_catalyst_row(today + timedelta(days=30), today) is True
    assert is_current_catalyst_row(today - timedelta(days=1), today) is False


def test_fill_pred_pct_pts_model_fallback_fills_gaps():
    row = {
        "price_at_cd": 10.0,
        "model_dm60_pct": 0.0,
        "model_dm30_pct": 10.0,
        "model_dm10_pct": 20.0,
        "model_dm7_pct": 18.0,
        "model_dm5_pct": 22.0,
        "model_dm3_pct": 24.0,
        "model_d4_pct": 30.0,
        "model_d7_pct": 35.0,
    }
    pts = [0.0, 10.0, None, None, 22.0, 24.0, 30.0, 35.0]
    out = orch._fill_pred_pct_pts_model_fallback(pts, row)
    assert out[0] == 0.0
    assert out[1] == 10.0
    assert out[2] is not None
    assert out[3] is not None
    assert out[2] == 20.0
    assert out[3] == 18.0


def test_accuracy_sheet_inserts_past_separator_row():
    """Riga viola tra coorte Simulation (CD≥oggi) e storico CD passate."""
    from openpyxl import Workbook
    from unittest.mock import patch

    today = date.today()
    cd_fut = today + timedelta(days=90)
    cd_past = today - timedelta(days=400)
    pk_fut = f"AAA|{cd_fut.isoformat()}"
    pk_past = f"BBB|{cd_past.isoformat()}"
    _filtered = {
        pk_fut: {
            "ticker": "AAA",
            "completion_date": cd_fut,
            "sponsor_match": "exact",
            "company_name_full": "Future Co",
            "nct_id": "NCT00000001",
            "model_dm60_pct": 0.0,
            "close_m60": 10.0,
            "price_at_cd": 10.0,
        },
        pk_past: {
            "ticker": "BBB",
            "completion_date": cd_past,
            "sponsor_match": "partial",
            "company_name_full": "Past Co",
            "nct_id": "NCT00000002",
            "model_dm60_pct": 0.0,
            "close_m60": 8.0,
            "price_at_cd": 8.0,
        },
    }
    _sim_rows = [
        {
            "ticker": "AAA",
            "completion_date": cd_fut,
            "sponsor_match": "exact",
            "name": "Future Co",
            "nct_id": "NCT00000001",
        }
    ]
    wb = Workbook()
    with patch.object(orch, "_accuracy_sim_histlib_backfill_session_closes", return_value=0):
        with patch.object(orch, "_accuracy_sim_merge_live_pred"):
            with patch.object(orch, "_accuracy_sim_capture_audit_by_row_key", return_value={}):
                with patch.object(orch, "_accuracy_capture_preserved_trailing_rows", return_value=[]):
                    with patch.object(orch, "_sec_k8_workbook_index", return_value={}):
                        orch._write_accuracy_simulation_sheet(
                            wb,
                            dict(_filtered),
                            sim_rows=_sim_rows,
                            sim_pred_data={},
                        )
    ws = wb["Accuracy"]
    sep_found = False
    for r in range(4, min(ws.max_row, 30) + 1):
        val = ws.cell(row=r, column=1).value
        if val and "CD passate" in str(val):
            sep_found = True
            break
    assert sep_found, "manca la riga separatore viola sopra lo storico"
