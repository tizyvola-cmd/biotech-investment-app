"""Tests for prediction refresh coordinator (mocked enrich / cohort)."""
from __future__ import annotations

import json
import os
from datetime import date
from unittest.mock import patch

import pytest

from past_pred_io import load_past_pred_document, save_predictions
from prediction.nct_cohort import select_nct_cohort
from prediction.refresh_coordinator import RefreshReport, run_prediction_refresh


def test_select_nct_cohort_empty_guida_strict_warning():
    pre = {
        "AAA|2024-01-01": {
            "ticker": "AAA",
            "completion_date": date(2024, 1, 1),
            "nct_relation_type": "N/D",
        },
    }
    with patch("prediction.nct_cohort.pred_nct_strict", return_value=True):
        sel = select_nct_cohort(pre, refresh_kind="guida")
    assert sel.cohort_empty is True
    assert sel.cohort_size == 0
    assert any("0 cohort" in w for w in sel.warnings)


def test_select_nct_cohort_accuracy_fallback_when_not_strict():
    pre = {
        "AAA|2024-01-01": {
            "ticker": "AAA",
            "completion_date": date(2024, 1, 1),
            "nct_relation_type": "indirect connections",
        },
    }
    with patch("prediction.nct_cohort.pred_nct_strict", return_value=False):
        sel = select_nct_cohort(pre, refresh_kind="accuracy")
    assert sel.cohort_size == 1
    assert any("fallback" in w.lower() for w in sel.warnings)


def test_save_predictions_refresh_meta(tmp_path):
    jpath = tmp_path / "past.json"
    doc = {"schema_version": 1, "rows": {"X|2024-06-01": {"ticker": "X"}}}
    meta = {
        "last_refresh_kind": "guida",
        "timestamp": "2026-05-17T12:00:00+00:00",
        "orchestrator_version": "test@1",
        "enrich_applied": True,
    }
    save_predictions(doc, meta=meta, json_path=str(jpath))
    loaded = json.loads(jpath.read_text(encoding="utf-8"))
    assert loaded["refresh_meta"]["last_refresh_kind"] == "guida"
    assert loaded["refresh_meta"]["enrich_applied"] is True


@patch("prediction.refresh_coordinator.enrich_past_pred_accuracy_metadata")
@patch("prediction.refresh_coordinator._refresh_guida")
@patch("os.path.isfile", return_value=True)
def test_run_guida_empty_cohort_skips_excel(_mock_isfile, mock_guida, mock_enrich, tmp_path):
    jpath = tmp_path / "past.json"
    doc = {
        "schema_version": 1,
        "rows": {
            "BBB|2025-01-01": {
                "ticker": "BBB",
                "completion_date": "2025-01-01",
                "nct_relation_type": "N/D",
            },
        },
    }
    jpath.write_text(json.dumps(doc), encoding="utf-8")
    mock_enrich.return_value = True
    mock_guida.side_effect = lambda r, *a, **k: r

    report = run_prediction_refresh(
        "guida",
        workbook_path=str(tmp_path / "missing.xlsx"),
        json_path=str(jpath),
        save_json=True,
        project_root=str(tmp_path),
    )
    assert report.records_in == 1
    assert report.cohort_empty is True
    assert report.enrich_applied is True
    mock_enrich.assert_called_once()
    loaded = load_past_pred_document(str(jpath))
    assert loaded.get("refresh_meta", {}).get("last_refresh_kind") == "guida"


def test_check_json_excel_alignment_keys(tmp_path):
    from prediction.io_sync import check_json_excel_alignment

    jpath = tmp_path / "past.json"
    jpath.write_text(
        json.dumps({
            "rows": {
                "AAA|2024-01-15": {"ticker": "AAA", "completion_date": "2024-01-15"},
            },
        }),
        encoding="utf-8",
    )
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Simulation"
        ws.cell(4, 1, "AAA")
        ws.cell(4, 3, "2024-01-15")
        xlsx = tmp_path / "book.xlsx"
        wb.save(xlsx)
    except ImportError:
        pytest.skip("openpyxl required")
    result = check_json_excel_alignment(str(jpath), str(xlsx))
    assert result["aligned_count"] == 1
    assert result["json_only_total"] == 0
    assert result["excel_only_total"] == 0


def test_refresh_report_exit_code_guida_empty():
    r = RefreshReport(
        refresh_kind="guida",
        records_in=1,
        records_out=1,
        cohort_size=0,
        nct_filter_applied=True,
        cohort_empty=True,
    )
    assert r.exit_code() == 2
